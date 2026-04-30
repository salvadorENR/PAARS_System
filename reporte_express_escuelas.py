import os
import sys
import glob
import pandas as pd
import numpy as np
import re
import matplotlib.pyplot as plt
import matplotlib.ticker as mtick
import textwrap
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# ==========================================
# 1. CONFIGURACIÓN GENERAL Y RUTAS
# ==========================================
PATH_GEISER_R1 = r"G:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\01_Resultados_Febrero\Interim_CSVs\Resultados"
PATH_GEISER_M1 = r"G:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\01_PROGRESO_Marzo\Interim_CSVs\Resultados"
PATH_GEISER_M2 = r"G:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\02_PROGRESO_Abril\Interim_CSVs\Resultados"

RUTA_MATRICULA_M1 = r"G:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\00_Metadata\MatriculaProgresoMes1.csv"
RUTA_MATRICULA_M2 = r"G:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\00_Metadata\MatriculaProgresoMes2.csv"

DIR_SALIDA = r"G:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\02_PROGRESO_Abril\Final_Reports"
os.makedirs(DIR_SALIDA, exist_ok=True)

ESCUELAS_TARGET = ['70097', '70046', '11976', '11453', '86394', '68097', '11679', '12124', '11336', '70042', '11992', '70001', '12053']

COLORES_CAT = {
    "Crítico": {"bg": "991B1B", "font": "FFFFFF", "hex": "#991B1B", "text_chart": "white"}, 
    "Bajo": {"bg": "FF8C2E", "font": "FFFFFF", "hex": "#FF8C2E", "text_chart": "white"},
    "Medio": {"bg": "FACC15", "font": "000000", "hex": "#FACC15", "text_chart": "black"}, 
    "Bueno": {"bg": "84CC16", "font": "000000", "hex": "#84CC16", "text_chart": "black"},
    "Excelente": {"bg": "065F46", "font": "FFFFFF", "hex": "#065F46", "text_chart": "white"}, 
    "Sin Datos": {"bg": "E0E0E0", "font": "000000", "hex": "#E0E0E0", "text_chart": "black"}
}

# ==========================================
# 2. FUNCIONES DE LECTURA Y TRADUCCIÓN
# ==========================================
def extraer_grado_geiser(grado_str):
    if pd.isna(grado_str): return None
    match = re.search(r'\d+', str(grado_str))
    return int(match.group()) if match else None

def normalizar_grado_matricula(grado_str):
    if pd.isna(grado_str): return None
    texto = str(grado_str).lower().strip()
    if "primer año" in texto or "1er año" in texto: return 10
    if "segundo año" in texto or "2do año" in texto: return 11
    if "segundo" in texto: return 2
    if "tercer" in texto: return 3
    if "cuarto" in texto: return 4
    if "quinto" in texto: return 5
    if "sexto" in texto: return 6
    if "séptimo" in texto or "septimo" in texto: return 7
    if "octavo" in texto: return 8
    if "noveno" in texto: return 9
    match = re.search(r'\d+', texto)
    return int(match.group()) if match else None

def clasificar_puntaje(val):
    if pd.isna(val) or val == "": return "Sin Datos"
    val = float(val)
    if val <= 35: return "Crítico"
    elif val <= 45: return "Bajo"
    elif val <= 55: return "Medio"
    elif val <= 65: return "Bueno"
    else: return "Excelente"

def es_vacio(x):
    if pd.isna(x): return True
    val = str(x).strip().lower()
    if val in ['', 'nan', 'null', '-', 'nr', 'n/a', 'na']: return True
    return False

def cargar_matricula(ruta):
    print(f"   -> Cargando Matrícula: {os.path.basename(ruta)}")
    if not os.path.exists(ruta): return pd.DataFrame()
    try:
        df = pd.read_csv(ruta, sep=';', dtype=str, encoding_errors='ignore')
        df.columns = df.columns.astype(str).str.strip()
        col_cod = next((c for c in df.columns if 'codigo' in c.lower() or 'centro' in c.lower()), None)
        col_nie = next((c for c in df.columns if 'nie' in c.lower() or 'documento' in c.lower()), None)
        col_grado = next((c for c in df.columns if 'grado' in c.lower()), None)
        if not (col_cod and col_nie and col_grado): return pd.DataFrame()
        df['codigo'] = df[col_cod].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
        df['nie'] = df[col_nie].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
        df['grado_num'] = df[col_grado].apply(normalizar_grado_matricula)
        df = df.dropna(subset=['grado_num'])
        df = df[(df['grado_num'] >= 2) & (df['grado_num'] <= 11)]
        return df[['codigo', 'nie', 'grado_num']].drop_duplicates()
    except Exception: return pd.DataFrame()

# =========================================================================
# MODIFICADO: Extrae TODA la data cruda. NO BORRA alumnos sin grado o NIE
# para que la Media Consolidada cuadre EXACTAMENTE con generador_rankings.py
# =========================================================================
def cargar_geiser_universal(ruta, mes_label):
    archivos = glob.glob(os.path.join(ruta, "*.xlsx")) + glob.glob(os.path.join(ruta, "*.xls")) + glob.glob(os.path.join(ruta, "*.csv"))
    lista_dfs = []
    if not archivos: return pd.DataFrame()
    for f in archivos:
        nombre_archivo = os.path.basename(f)
        if 'legend' in f.lower() or nombre_archivo.startswith('~$'): continue
        try:
            if f.lower().endswith('.csv'): df = pd.read_csv(f, dtype=str, encoding_errors='ignore')
            else: df = pd.read_excel(f, dtype=str)
            df.columns = df.columns.astype(str).str.strip()
            materia = 'Mat' if 'MAT' in nombre_archivo.upper() else 'Len'
            
            col_cod = next((c for c in df.columns if 'centro' in c.lower() or 'codigo' in c.lower()), None)
            col_centro = next((c for c in df.columns if 'centro' in c.lower() and 'nro' not in c.lower()), None)
            col_nie = next((c for c in df.columns if 'documento' in c.lower() or 'nie' in c.lower()), None)
            col_nombre = next((c for c in df.columns if 'nombre' in c.lower()), None)
            col_apellido = next((c for c in df.columns if 'apellido' in c.lower()), None)
            col_grado = next((c for c in df.columns if 'grado' in c.lower()), None)
            col_grupo = next((c for c in df.columns if 'grupo' in c.lower() or 'sección' in c.lower()), None)
            col_theta = next((c for c in df.columns if '0-100' in c.lower()), None)
            col_anular = next((c for c in df.columns if 'anular_prueba' in c.lower()), None)
            
            if not col_cod: continue
            
            if materia == 'Mat': item_cols = [c for c in df.columns if str(c).upper().strip().startswith('MAT') and any(char.isdigit() for char in str(c))]
            else: item_cols = [c for c in df.columns if str(c).upper().strip().startswith('LEC') and any(char.isdigit() for char in str(c))]
            
            if item_cols: df['tiene_items_vacios'] = df[item_cols].apply(lambda col: col.apply(es_vacio)).any(axis=1)
            else: df['tiene_items_vacios'] = False

            if col_anular: df['Anulado'] = df[col_anular].apply(lambda x: False if es_vacio(x) or str(x).strip().lower() in ['0', 'false', 'falso', 'no'] else True)
            else: df['Anulado'] = False
            
            df['codigo'] = df[col_cod].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
            df['Centro_Nom'] = df[col_centro].str.strip() if col_centro else "Desconocido"
            
            df['nie'] = df[col_nie].astype(str).str.replace(r'\.0$', '', regex=True).str.strip() if col_nie else ""
            df['grado_num'] = df[col_grado].apply(extraer_grado_geiser) if col_grado else np.nan
            df['Grupo'] = df[col_grupo].astype(str).str.strip() if col_grupo else ""
            
            nom = df[col_nombre].fillna('').str.strip() if col_nombre else ""
            ape = df[col_apellido].fillna('').str.strip() if col_apellido else ""
            df['Apellido+Nombre'] = ape + " " + nom
            df['Puntaje'] = pd.to_numeric(df[col_theta].str.replace(',', '.'), errors='coerce') if col_theta else np.nan
            df['Materia'] = materia
            df['Mes'] = mes_label
            
            cols_mantener = ['codigo', 'Centro_Nom', 'nie', 'Apellido+Nombre', 'grado_num', 'Grupo', 'Puntaje', 'Materia', 'Mes', 'Anulado', 'tiene_items_vacios']
            lista_dfs.append(df[cols_mantener].copy())
        except Exception: pass
    return pd.concat(lista_dfs, ignore_index=True) if lista_dfs else pd.DataFrame()

def generar_matriz_basica(df_data, values_col):
    if df_data.empty: return pd.DataFrame()
    pivot = df_data.pivot_table(index='codigo', columns='grado_num', values=values_col, aggfunc='count', fill_value=0)
    pivot['TOTAL ESCUELA'] = pivot.sum(axis=1)
    pivot.loc['TOTAL GENERAL'] = pivot.sum(axis=0)
    pivot.columns = [f"{int(c)}°" if isinstance(c, (int, float)) else c for c in pivot.columns]
    pivot.index.name = "Centro Escolar"
    return pivot.reset_index()

def calcular_matriz_con_porcentajes(df_master, mes, materia, condicion, base_tipo):
    df_base = df_master[(df_master['Mes'] == mes) & (df_master['Materia'] == materia)].copy()
    if df_base.empty: return pd.DataFrame()
    pivot_den = df_base.pivot_table(index='codigo', columns='grado_num', values='nie', aggfunc='count', fill_value=0)

    if condicion == 'C1':   df_num = df_base[(df_base['Anulado'] == True) & (df_base['tiene_items_vacios'] == False)]
    elif condicion == 'C2': df_num = df_base[(df_base['Anulado'] == False) & (df_base['tiene_items_vacios'] == True)]
    elif condicion == 'C3': df_num = df_base[(df_base['Anulado'] == True) | (df_base['tiene_items_vacios'] == True)]
    elif condicion == 'C4': df_num = df_base[(df_base['Anulado'] == True) & (df_base['tiene_items_vacios'] == True)]
    else: df_num = pd.DataFrame()

    if df_num.empty: pivot_num = pd.DataFrame(0, index=pivot_den.index, columns=pivot_den.columns)
    else:
        pivot_num = df_num.pivot_table(index='codigo', columns='grado_num', values='nie', aggfunc='count', fill_value=0)
        pivot_num = pivot_num.reindex_like(pivot_den).fillna(0)

    pivot_num['TOTAL ESCUELA'] = pivot_num.sum(axis=1); pivot_num.loc['TOTAL GENERAL'] = pivot_num.sum(axis=0)
    pivot_den['TOTAL ESCUELA'] = pivot_den.sum(axis=1); pivot_den.loc['TOTAL GENERAL'] = pivot_den.sum(axis=0)
    pivot_str = pd.DataFrame(index=pivot_num.index, columns=pivot_num.columns)
    total_prueba_den = pivot_den.loc['TOTAL GENERAL', 'TOTAL ESCUELA'] if 'TOTAL GENERAL' in pivot_den.index and 'TOTAL ESCUELA' in pivot_den.columns else 0

    for col in pivot_num.columns:
        for idx in pivot_num.index:
            n = pivot_num.loc[idx, col]; d = None
            if base_tipo == 'B1':   d = pivot_den.loc[idx, col]
            elif base_tipo == 'B2': d = total_prueba_den
            elif base_tipo == 'B3': d = None if col == 'TOTAL ESCUELA' else pivot_den.loc['TOTAL GENERAL', col]
            elif base_tipo == 'B4': d = None if idx == 'TOTAL GENERAL' else pivot_den.loc[idx, 'TOTAL ESCUELA']
            if d is None: pivot_str.loc[idx, col] = f"{int(n)}"
            elif d > 0: pivot_str.loc[idx, col] = f"{int(n)} ({round((n / d) * 100, 1)}%)"
            else: pivot_str.loc[idx, col] = "0 (0.0%)"
    pivot_str.index.name = "Centro Escolar"
    pivot_str.columns = [f"{int(c)}°" if isinstance(c, (int, float)) else c for c in pivot_str.columns]
    return pivot_str.reset_index()


# ==========================================
# 4. EJECUCIÓN CENTRALIZADA
# ==========================================
print("\n[*] INICIANDO SÚPER-SCRIPT DE REQUERIMIENTOS ESPECIALES...")

print("\n[1/4] Extrayendo base de datos universal Geiser...")
df_g_r1 = cargar_geiser_universal(PATH_GEISER_R1, "R1")
df_g1 = cargar_geiser_universal(PATH_GEISER_M1, "M1")
df_g2 = cargar_geiser_universal(PATH_GEISER_M2, "M2")
if df_g1.empty and df_g2.empty and df_g_r1.empty:
    print("\n[!] ERROR CRÍTICO: No se cargaron datos de Geiser.")
    sys.exit()

# ======= ARQUITECTURA DE DOBLE CAPA (LIMPIO VS CRUDO) =======
df_geiser_crudo = pd.concat([df_g_r1, df_g1, df_g2], ignore_index=True)
df_geiser_crudo = df_geiser_crudo[df_geiser_crudo['Puntaje'].notna()].copy() 

# Data limpia: SOLO para las listas individuales de niños (para no ver "fantasmas" en la hoja principal)
df_geiser_validos = df_geiser_crudo[
    (df_geiser_crudo['Anulado'] == False) & 
    (df_geiser_crudo['grado_num'].notna()) &
    (df_geiser_crudo['grado_num'] >= 2) & 
    (df_geiser_crudo['grado_num'] <= 11) &
    (df_geiser_crudo['nie'] != "")
].copy() 

print("\n[2/4] Cargando registros de Matrícula...")
mat_m1 = cargar_matricula(RUTA_MATRICULA_M1)
mat_m2 = cargar_matricula(RUTA_MATRICULA_M2)

# ================= REQ 1: MATRÍCULA =================
print("\n[3/4] Generando Reportes Excel (Matrícula, 13 Escuelas, Todas las Escuelas, Ausentismo y Gráficos)...")
ruta_matricula = os.path.join(DIR_SALIDA, "Reporte_Matricula.xlsx")
with pd.ExcelWriter(ruta_matricula) as writer:
    if not mat_m1.empty: generar_matriz_basica(mat_m1, 'nie').to_excel(writer, sheet_name='Matrícula Mes 1', index=False)
    if not mat_m2.empty: generar_matriz_basica(mat_m2, 'nie').to_excel(writer, sheet_name='Matrícula Mes 2', index=False)


# =========================================================================================
# LA MAGIA: EL VIEJO ALGORITMO EXACTO INYECTADO PARA REPLICAR generador_rankings.py
# =========================================================================================
def calcular_promedios_agrupados(df_raw, index_cols):
    if df_raw.empty: return pd.DataFrame()
    
    # 1. Media individual por Materia (Matemática y Lengua)
    df_matlen = df_raw.pivot_table(index=index_cols, columns=['Mes', 'Materia'], values='Puntaje', aggfunc='mean').reset_index()
    df_matlen.columns = [' '.join(col).strip() if isinstance(col, tuple) else col for col in df_matlen.columns.values]
    
    # 2. EL VIEJO ALGORITMO: Agrupar TODO junto sin separar por materia y redondear a 1 decimal
    df_cons = df_raw.groupby(index_cols + ['Mes'])['Puntaje'].mean().reset_index()
    df_cons['Puntaje'] = df_cons['Puntaje'].round(1)  # ¡El redondeo exacto de tu script original!
    
    # Lo pivotamos para que los meses queden como columnas
    df_cons_pivot = df_cons.pivot_table(index=index_cols, columns=['Mes'], values='Puntaje', aggfunc='first').reset_index()
    df_cons_pivot.columns = [str(c) + ' Consolidado' if c in ['R1', 'M1', 'M2'] else c for c in df_cons_pivot.columns.values]
    
    # Unimos ambos cálculos
    df_res = pd.merge(df_matlen, df_cons_pivot, on=index_cols, how='left')
    
    for c in ['R1 Mat', 'M1 Mat', 'M2 Mat', 'R1 Len', 'M1 Len', 'M2 Len', 'R1 Consolidado', 'M1 Consolidado', 'M2 Consolidado']:
        if c not in df_res.columns: df_res[c] = np.nan
            
    rep = pd.DataFrame()
    for col in index_cols:
        if col == 'codigo': rep['Código de infraestructura'] = df_res[col]
        elif col == 'Centro_Nom': rep['Centro'] = df_res[col]
        elif col == 'grado_num': rep['Grado'] = df_res[col].apply(lambda x: f"{int(x)}°" if pd.notna(x) else "")
        elif col == 'Grupo': rep['Sección'] = df_res[col]

    for prefix in ['R1', '1', '2']:
        m_prefix = prefix if prefix == 'R1' else 'M' + prefix
        rep[f'Media {prefix} Mat'] = df_res[f'{m_prefix} Mat'].round(2); rep[f'Categoría {prefix} Mat'] = rep[f'Media {prefix} Mat'].apply(clasificar_puntaje)
        rep[f'Media {prefix} Len'] = df_res[f'{m_prefix} Len'].round(2); rep[f'Categoría {prefix} Len'] = rep[f'Media {prefix} Len'].apply(clasificar_puntaje)
        
        # INYECCIÓN DIRECTA: Insertamos la nota Consolidada pre-redondeada por el viejo algoritmo
        rep[f'Media {prefix} Cons'] = df_res[f'{m_prefix} Consolidado']
        rep[f'Categoría {prefix} Cons'] = rep[f'Media {prefix} Cons'].apply(clasificar_puntaje)
        
    for c in rep.columns:
        if 'Media' in c: rep[c] = rep[c].fillna("")
            
    return rep

# ================= FUNCIÓN PARA ARMAR EL REPORTE GENERAL =================
def generar_reporte_escuelas(df_estudiantes, df_crudo_escuela, nombre_archivo, incluir_agrupados=False):
    if df_estudiantes.empty: return False
    
    # 1. Hoja individual (con datos súper limpios)
    df_p = df_estudiantes.pivot_table(index=['codigo', 'Centro_Nom', 'grado_num', 'Grupo', 'nie', 'Apellido+Nombre'], columns=['Mes', 'Materia'], values='Puntaje', aggfunc='first').reset_index()
    df_p.columns = [' '.join(col).strip() for col in df_p.columns.values]
    for col in ['R1 Mat', 'M1 Mat', 'M2 Mat', 'R1 Len', 'M1 Len', 'M2 Len']:
        if col not in df_p.columns: df_p[col] = np.nan

    rep1 = pd.DataFrame()
    rep1['Código de infraestructura'] = df_p['codigo']; rep1['Centro'] = df_p['Centro_Nom']
    rep1['Grado'] = df_p['grado_num'].apply(lambda x: f"{int(x)}°" if pd.notna(x) else "")
    rep1['Sección'] = df_p['Grupo']; rep1['NIE'] = df_p['nie']; rep1['Apellido+Nombre'] = df_p['Apellido+Nombre']
    
    rep1['Puntaje R1 Mat'] = df_p['R1 Mat'].round(2); rep1['Categoría R1 Mat'] = rep1['Puntaje R1 Mat'].apply(clasificar_puntaje)
    rep1['Puntaje 1 Mat'] = df_p['M1 Mat'].round(2); rep1['Categoría 1 Mat'] = rep1['Puntaje 1 Mat'].apply(clasificar_puntaje)
    rep1['Puntaje 2 Mat'] = df_p['M2 Mat'].round(2); rep1['Categoría 2 Mat'] = rep1['Puntaje 2 Mat'].apply(clasificar_puntaje)
    rep1['Puntaje R1 Len'] = df_p['R1 Len'].round(2); rep1['Categoría R1 Len'] = rep1['Puntaje R1 Len'].apply(clasificar_puntaje)
    rep1['Puntaje 1 Len'] = df_p['M1 Len'].round(2); rep1['Categoría 1 Len'] = rep1['Puntaje 1 Len'].apply(clasificar_puntaje)
    rep1['Puntaje 2 Len'] = df_p['M2 Len'].round(2); rep1['Categoría 2 Len'] = rep1['Puntaje 2 Len'].apply(clasificar_puntaje)
    
    for col in ['Puntaje R1 Mat', 'Puntaje 1 Mat', 'Puntaje 2 Mat', 'Puntaje R1 Len', 'Puntaje 1 Len', 'Puntaje 2 Len']: rep1[col] = rep1[col].fillna("")
    rep1['grado_num'] = df_p['grado_num'] 

    rep1_excel = rep1.drop(columns=['grado_num']) 
    
    with pd.ExcelWriter(nombre_archivo, engine='openpyxl') as writer:
        rep1_excel.to_excel(writer, index=False, sheet_name='Resultados')
        ws = writer.book['Resultados']
        cols_cat = [rep1_excel.columns.get_loc(c) + 1 for c in ['Categoría R1 Mat', 'Categoría 1 Mat', 'Categoría 2 Mat', 'Categoría R1 Len', 'Categoría 1 Len', 'Categoría 2 Len']]
        for row in range(2, len(rep1_excel) + 2):
            for col_idx in cols_cat:
                val = ws.cell(row=row, column=col_idx).value
                if val in COLORES_CAT:
                    ws.cell(row=row, column=col_idx).fill = PatternFill(start_color=COLORES_CAT[val]['bg'], end_color=COLORES_CAT[val]['bg'], fill_type="solid")
                    ws.cell(row=row, column=col_idx).font = Font(color=COLORES_CAT[val]['font'], bold=True)
                    
        # 2. Hojas de promedios (Calculados con datos CRUDOS para cuadrar con el ranking histórico)
        if incluir_agrupados:
            rep_escuela = calcular_promedios_agrupados(df_crudo_escuela, ['codigo', 'Centro_Nom'])
            if not rep_escuela.empty:
                rep_escuela.to_excel(writer, index=False, sheet_name='Promedios por Centro')
                ws_esc = writer.book['Promedios por Centro']
                cols_cat_esc = [i+1 for i, c in enumerate(rep_escuela.columns) if 'Categoría' in c]
                for row in range(2, len(rep_escuela) + 2):
                    for col_idx in cols_cat_esc:
                        val = ws_esc.cell(row=row, column=col_idx).value
                        if val in COLORES_CAT:
                            ws_esc.cell(row=row, column=col_idx).fill = PatternFill(start_color=COLORES_CAT[val]['bg'], end_color=COLORES_CAT[val]['bg'], fill_type="solid")
                            ws_esc.cell(row=row, column=col_idx).font = Font(color=COLORES_CAT[val]['font'], bold=True)
            
            rep_secc = calcular_promedios_agrupados(df_crudo_escuela, ['codigo', 'Centro_Nom', 'grado_num', 'Grupo'])
            if not rep_secc.empty:
                rep_secc.to_excel(writer, index=False, sheet_name='Promedios por Sección')
                ws_sec = writer.book['Promedios por Sección']
                cols_cat_sec = [i+1 for i, c in enumerate(rep_secc.columns) if 'Categoría' in c]
                for row in range(2, len(rep_secc) + 2):
                    for col_idx in cols_cat_sec:
                        val = ws_sec.cell(row=row, column=col_idx).value
                        if val in COLORES_CAT:
                            ws_sec.cell(row=row, column=col_idx).fill = PatternFill(start_color=COLORES_CAT[val]['bg'], end_color=COLORES_CAT[val]['bg'], fill_type="solid")
                            ws_sec.cell(row=row, column=col_idx).font = Font(color=COLORES_CAT[val]['font'], bold=True)
    return rep1

# ================= REQ 2: 13 ESCUELAS (R1 + M1 + M2) =================
df_req1_validos = df_geiser_validos[df_geiser_validos['codigo'].isin(ESCUELAS_TARGET)].copy()
df_req1_crudos = df_geiser_crudo[df_geiser_crudo['codigo'].isin(ESCUELAS_TARGET)].copy()

ruta_req1 = os.path.join(DIR_SALIDA, "Reporte_Especial_13_Escuelas.xlsx")
rep1 = generar_reporte_escuelas(df_req1_validos, df_req1_crudos, ruta_req1, incluir_agrupados=False) 
if rep1 is not False: print(f"  [OK] Creado: {ruta_req1}")

# ================= REQ 2.5: TODAS LAS ESCUELAS DE RESULTADOS (FEB) =================
escuelas_r1 = df_geiser_crudo[df_geiser_crudo['Mes'] == 'R1']['codigo'].unique()
num_escuelas_r1 = len(escuelas_r1)

if num_escuelas_r1 > 0:
    df_req_todas_validos = df_geiser_validos[df_geiser_validos['codigo'].isin(escuelas_r1)].copy()
    df_req_todas_crudos = df_geiser_crudo[df_geiser_crudo['codigo'].isin(escuelas_r1)].copy()
    
    ruta_req_todas = os.path.join(DIR_SALIDA, f"Reporte_General_{num_escuelas_r1}_Escuelas.xlsx")
    generar_reporte_escuelas(df_req_todas_validos, df_req_todas_crudos, ruta_req_todas, incluir_agrupados=True)
    print(f"  [OK] Creado: {ruta_req_todas} (Universo total de R1 con Promedio Histórico restaurado)")
else:
    print("  [!] No se encontraron escuelas participantes en la prueba de Resultados (Febrero).")

if rep1 is not False:
    # >>> GRÁFICOS INDIVIDUALES <<<
    print("\n[+] Generando Gráficos de Distribución (Para 13 Escuelas)...")
    pruebas_a_graficar = [('1 Mat', 'Prueba Progreso Mes 1 (Marzo)'), ('2 Mat', 'Prueba Progreso Mes 2 (Abril)'), ('1 Len', 'Prueba Progreso Mes 1 (Marzo)'), ('2 Len', 'Prueba Progreso Mes 2 (Abril)')]
    orden_cats = ['Crítico', 'Bajo', 'Medio', 'Bueno', 'Excelente']
    colores_hex = [COLORES_CAT[c]['hex'] for c in orden_cats]
    text_colors = [COLORES_CAT[c]['text_chart'] for c in orden_cats]

    for col_suffix, title in pruebas_a_graficar:
        col_cat = f'Categoría {col_suffix}'
        df_valid = rep1[rep1[col_cat].isin(orden_cats)].copy()
        if not df_valid.empty:
            ct_centro = pd.crosstab(df_valid['Centro'], df_valid[col_cat], normalize='index') * 100
            for c in orden_cats:
                if c not in ct_centro.columns: ct_centro[c] = 0
            ct_centro = ct_centro[orden_cats].sort_index(ascending=False)
            ct_centro.index = [textwrap.fill(str(x), width=35) for x in ct_centro.index]
            fig, ax = plt.subplots(figsize=(14, 8))
            ct_centro.plot(kind='barh', stacked=True, color=colores_hex, ax=ax, width=0.75, edgecolor='white', linewidth=0.5)
            ax.set_title(f'Porcentaje de Estudiantes por Nivel de Logro en\n{title} (Por Centro)', fontsize=16, fontweight='bold', pad=20)
            ax.set_xlabel('Porcentaje (%)', fontsize=12, fontweight='bold', labelpad=10); ax.set_ylabel('')
            ax.xaxis.set_major_formatter(mtick.PercentFormatter()); ax.set_xlim(0, 100)
            handles, labels = ax.get_legend_handles_labels()
            ax.legend(reversed(handles), reversed(labels), title='Nivel de Logro', bbox_to_anchor=(1.02, 1), loc='upper left', fontsize=11)
            for j, c in enumerate(ax.containers):
                labels = [f"{v:.1f}%" if v > 4 else "" for v in c.datavalues]
                ax.bar_label(c, labels=labels, label_type='center', color=text_colors[j], fontsize=9, fontweight='bold')
            plt.tight_layout()
            plt.savefig(os.path.join(DIR_SALIDA, f"Grafico_Centro_{title.replace(' ', '_').replace('(', '').replace(')', '')}.png"), dpi=300, bbox_inches='tight')
            plt.close()

            ct_grado = pd.crosstab(df_valid['grado_num'], df_valid[col_cat], normalize='index') * 100
            for c in orden_cats:
                if c not in ct_grado.columns: ct_grado[c] = 0
            ct_grado = ct_grado[orden_cats].sort_index(ascending=False)
            ct_grado.index = [f"{int(x)}°" for x in ct_grado.index]
            fig, ax = plt.subplots(figsize=(12, 6))
            ct_grado.plot(kind='barh', stacked=True, color=colores_hex, ax=ax, width=0.75, edgecolor='white', linewidth=0.5)
            ax.set_title(f'Porcentaje de Estudiantes por Nivel de Logro en\n{title} (Por Grado)', fontsize=16, fontweight='bold', pad=20)
            ax.set_xlabel('Porcentaje (%)', fontsize=12, fontweight='bold', labelpad=10); ax.set_ylabel('Grado', fontsize=12, fontweight='bold', labelpad=10)
            ax.xaxis.set_major_formatter(mtick.PercentFormatter()); ax.set_xlim(0, 100)
            handles, labels = ax.get_legend_handles_labels()
            ax.legend(reversed(handles), reversed(labels), title='Nivel de Logro', bbox_to_anchor=(1.02, 1), loc='upper left', fontsize=11)
            for j, c in enumerate(ax.containers):
                labels = [f"{v:.1f}%" if v > 4 else "" for v in c.datavalues]
                ax.bar_label(c, labels=labels, label_type='center', color=text_colors[j], fontsize=9, fontweight='bold')
            plt.tight_layout()
            plt.savefig(os.path.join(DIR_SALIDA, f"Grafico_Grado_{title.replace(' ', '_').replace('(', '').replace(')', '')}.png"), dpi=300, bbox_inches='tight')
            plt.close()

    # >>> GRÁFICOS COMPARATIVOS <<<
    evaluaciones = {'Matemática': [('1 Mat', 'Prueba Progreso Mes 1 (Marzo)'), ('2 Mat', 'Prueba Progreso Mes 2 (Abril)')], 'Lengua': [('1 Len', 'Prueba Progreso Mes 1 (Marzo)'), ('2 Len', 'Prueba Progreso Mes 2 (Abril)')]}
    for materia, pruebas in evaluaciones.items():
        fig_c, axes_c = plt.subplots(1, 2, figsize=(18, 10), sharey=True)
        fig_c.suptitle(f'Evolución de Niveles de Logro - {materia} (Por Centro Escolar)', fontsize=20, fontweight='bold')
        for idx, (col_suffix, title) in enumerate(pruebas):
            ax = axes_c[idx]
            col_cat = f'Categoría {col_suffix}'
            df_valid = rep1[rep1[col_cat].isin(orden_cats)].copy()
            if not df_valid.empty:
                ct = pd.crosstab(df_valid['Centro'], df_valid[col_cat], normalize='index') * 100
                for c in orden_cats:
                    if c not in ct.columns: ct[c] = 0
                ct = ct[orden_cats].sort_index(ascending=False); ct.index = [textwrap.fill(str(x), width=30) for x in ct.index]
                ct.plot(kind='barh', stacked=True, color=colores_hex, ax=ax, width=0.8, edgecolor='white', linewidth=0.5, legend=False)
                ax.set_title(title, fontsize=16, fontweight='bold', pad=15); ax.set_xlabel('Porcentaje (%)', fontsize=12, fontweight='bold', labelpad=10); ax.set_ylabel('')
                ax.xaxis.set_major_formatter(mtick.PercentFormatter()); ax.set_xlim(0, 100)
                for j, container in enumerate(ax.containers):
                    labels = [f"{v:.1f}%" if v > 4 else "" for v in container.datavalues]
                    ax.bar_label(container, labels=labels, label_type='center', color=text_colors[j], fontsize=10, fontweight='bold')
            else:
                ax.set_title(f"{title}\n(Sin Datos Aún)", fontsize=16, fontweight='bold', pad=15); ax.axis('off')
        handles, labels = axes_c[0].get_legend_handles_labels()
        if handles: fig_c.legend(reversed(handles), reversed(labels), title='Nivel de Logro', bbox_to_anchor=(0.5, 0.03), loc='upper center', ncol=5, fontsize=13, title_fontsize=14)
        plt.tight_layout(rect=[0, 0.06, 1, 0.95])
        fig_c.savefig(os.path.join(DIR_SALIDA, f"Grafico_Comparativo_Centro_{materia}.png"), dpi=300, bbox_inches='tight'); plt.close(fig_c)

        fig_g, axes_g = plt.subplots(1, 2, figsize=(18, 8), sharey=True)
        fig_g.suptitle(f'Evolución de Niveles de Logro - {materia} (Por Grado)', fontsize=20, fontweight='bold')
        for idx, (col_suffix, title) in enumerate(pruebas):
            ax = axes_g[idx]
            col_cat = f'Categoría {col_suffix}'
            df_valid = rep1[rep1[col_cat].isin(orden_cats)].copy()
            if not df_valid.empty:
                ct = pd.crosstab(df_valid['grado_num'], df_valid[col_cat], normalize='index') * 100
                for c in orden_cats:
                    if c not in ct.columns: ct[c] = 0
                ct = ct[orden_cats].sort_index(ascending=False); ct.index = [f"{int(x)}°" for x in ct.index]
                ct.plot(kind='barh', stacked=True, color=colores_hex, ax=ax, width=0.8, edgecolor='white', linewidth=0.5, legend=False)
                ax.set_title(title, fontsize=16, fontweight='bold', pad=15); ax.set_xlabel('Porcentaje (%)', fontsize=12, fontweight='bold', labelpad=10); ax.set_ylabel('Grado' if idx == 0 else '', fontsize=12, fontweight='bold')
                ax.xaxis.set_major_formatter(mtick.PercentFormatter()); ax.set_xlim(0, 100)
                for j, container in enumerate(ax.containers):
                    labels = [f"{v:.1f}%" if v > 4 else "" for v in container.datavalues]
                    ax.bar_label(container, labels=labels, label_type='center', color=text_colors[j], fontsize=10, fontweight='bold')
            else:
                ax.set_title(f"{title}\n(Sin Datos Aún)", fontsize=16, fontweight='bold', pad=15); ax.axis('off')
        handles, labels = axes_g[0].get_legend_handles_labels()
        if handles: fig_g.legend(reversed(handles), reversed(labels), title='Nivel de Logro', bbox_to_anchor=(0.5, 0.02), loc='upper center', ncol=5, fontsize=13, title_fontsize=14)
        plt.tight_layout(rect=[0, 0.06, 1, 0.95])
        fig_g.savefig(os.path.join(DIR_SALIDA, f"Grafico_Comparativo_Grado_{materia}.png"), dpi=300, bbox_inches='tight'); plt.close(fig_g)


# ================= REQ 3: AUSENTISMO =================
ruta_ausentes = os.path.join(DIR_SALIDA, "Reporte_Ausentismo.xlsx")
with pd.ExcelWriter(ruta_ausentes) as writer:
    for hoja, m, mat, df_m in [('Matemática Mes 1', 'M1', 'Mat', mat_m1), ('Matemática Mes 2', 'M2', 'Mat', mat_m2), ('Lengua Mes 1', 'M1', 'Len', mat_m1), ('Lengua Mes 2', 'M2', 'Len', mat_m2)]:
        if df_m.empty: pd.DataFrame({'Info': ['Sin datos']}).to_excel(writer, sheet_name=hoja, index=False)
        else:
            df_g = df_geiser_validos[(df_geiser_validos['Mes'] == m) & (df_geiser_validos['Materia'] == mat)]
            df_a = df_m[~df_m['nie'].isin(set(df_g['nie'].unique()))].copy()
            generar_matriz_basica(df_a, 'nie').to_excel(writer, sheet_name=hoja, index=False)

# ================= REQ 4: DASHBOARD HTML =================
print("\n[4/4] Generando Dashboard Interactivo de Pruebas No Finalizadas...")
html_template = """<!DOCTYPE html>
<html lang="es">
<head>
    <meta charset="UTF-8">
    <title>Dashboard - Pruebas No Finalizadas</title>
    <script src="https://cdnjs.cloudflare.com/ajax/libs/xlsx/0.18.5/xlsx.full.min.js"></script>
    <style>
        body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; background-color: #f4f7f6; margin: 0; padding: 20px; }
        .header { background-color: #2c3e50; color: white; padding: 20px; border-radius: 8px; margin-bottom: 20px; text-align: center; }
        .controls { display: flex; justify-content: space-between; align-items: center; background: white; padding: 15px 20px; border-radius: 8px; box-shadow: 0 2px 5px rgba(0,0,0,0.1); margin-bottom: 20px; flex-wrap: wrap; gap: 15px; }
        select { padding: 10px; font-size: 16px; border: 1px solid #ccc; border-radius: 5px; font-weight: bold; cursor: pointer; }
        .btn-group { display: flex; gap: 10px; flex-wrap: wrap; }
        button { padding: 10px 15px; font-size: 14px; border: none; border-radius: 5px; cursor: pointer; font-weight: bold; transition: 0.3s; background-color: #e0e0e0; color: #333; }
        button.active { background-color: #3498db; color: white; box-shadow: 0 2px 5px rgba(0,0,0,0.2); }
        button.download-btn { background-color: #27ae60; color: white; margin-left: auto; }
        button.download-btn:hover { background-color: #219653; }
        .table-container { background: white; padding: 20px; border-radius: 8px; box-shadow: 0 2px 5px rgba(0,0,0,0.1); overflow-x: auto; }
        table { width: 100%; border-collapse: collapse; margin-top: 10px; font-size: 14px; }
        th, td { border: 1px solid #ddd; padding: 10px; text-align: center; white-space: nowrap; }
        th { background-color: #34495e; color: white; position: sticky; top: 0; }
        tr:nth-child(even) { background-color: #f9f9f9; }
        tr:hover { background-color: #f1f1f1; }
        .info-text { margin-top: 10px; font-size: 13px; color: #666; }
    </style>
</head>
<body>
    <div class="header">
        <h1 style="margin:0;">Dashboard: Pruebas No Finalizadas</h1>
        <p style="margin:5px 0 0 0;">Análisis interactivo de anulación por tiempo e ítems en blanco</p>
    </div>
    <div class="controls">
        <div style="display: flex; gap: 15px; flex-wrap: wrap;">
            <div>
                <label style="font-weight: bold; margin-right: 5px;">Seleccione la Prueba:</label>
                <select id="selectPrueba" onchange="updateView()">
                    <option value="M1_Mat">Matemática - Mes 1</option>
                    <option value="M2_Mat">Matemática - Mes 2</option>
                    <option value="M1_Len">Lengua - Mes 1</option>
                    <option value="M2_Len">Lengua - Mes 2</option>
                </select>
            </div>
            <div>
                <label style="font-weight: bold; margin-right: 5px;">Base del Porcentaje:</label>
                <select id="selectBase" onchange="updateView()">
                    <option value="B1">Por Celda (Centro y Grado)</option>
                    <option value="B2">Por Total de la Prueba</option>
                    <option value="B3">Por Grado (Columna)</option>
                    <option value="B4">Por Centro Escolar (Fila)</option>
                </select>
            </div>
        </div>
        <div class="btn-group" style="margin-top: 10px; width: 100%;">
            <button id="btn_C1" onclick="setCriteria('C1')">Anulada por tiempo</button>
            <button id="btn_C2" onclick="setCriteria('C2')">ítems incompletos</button>
            <button id="btn_C3" class="active" onclick="setCriteria('C3')">Tiempo o ítems incompletos</button>
            <button id="btn_C4" onclick="setCriteria('C4')">Tiempo e ítems incompletos</button>
            <button class="download-btn" onclick="downloadExcel()">⬇ Descargar Tabla Actual (Excel)</button>
        </div>
    </div>
    <div class="table-container">
        <div id="dynamic-tables">{tables_html}</div>
        <p class="info-text">* Los porcentajes se calculan dinámicamente según la <b>Base del Porcentaje</b> seleccionada.</p>
    </div>
    <script>
        let currentCriteria = 'C3'; let currentBase = 'B1';
        function setCriteria(crit) {
            currentCriteria = crit;
            document.querySelectorAll('.btn-group button').forEach(btn => btn.classList.remove('active'));
            document.getElementById('btn_' + crit).classList.add('active'); updateView();
        }
        function updateView() {
            let prueba = document.getElementById('selectPrueba').value; currentBase = document.getElementById('selectBase').value;
            let targetId = 'table_' + prueba + '_' + currentCriteria + '_' + currentBase;
            document.querySelectorAll('.data-table').forEach(tbl => tbl.style.display = 'none');
            let activeTable = document.getElementById(targetId);
            if (activeTable) { activeTable.style.display = 'table'; }
        }
        function downloadExcel() {
            let prueba = document.getElementById('selectPrueba').value;
            let targetId = 'table_' + prueba + '_' + currentCriteria + '_' + currentBase;
            let visibleTable = document.getElementById(targetId);
            if (!visibleTable) { alert("No hay datos para descargar."); return; }
            let pruebaNombre = document.getElementById('selectPrueba').options[document.getElementById('selectPrueba').selectedIndex].text;
            let baseNombre = document.getElementById('selectBase').options[document.getElementById('selectBase').selectedIndex].text;
            let btnActivo = document.querySelector('.btn-group button.active').innerText;
            let filename = `NoFinalizadas_${pruebaNombre}_${btnActivo}_${baseNombre}.xlsx`.replace(/ /g, "_").replace(/[()]/g, "");
            let wb = XLSX.utils.table_to_book(visibleTable, {sheet: "Resultados"});
            XLSX.writeFile(wb, filename);
        }
        updateView();
    </script>
</body>
</html>
"""

todas_las_tablas_html = ""
combinaciones = [('M1', 'Mat', 'M1_Mat'), ('M2', 'Mat', 'M2_Mat'), ('M1', 'Len', 'M1_Len'), ('M2', 'Len', 'M2_Len')]
criterios = ['C1', 'C2', 'C3', 'C4']
bases = ['B1', 'B2', 'B3', 'B4']

for mes, mat, id_p in combinaciones:
    for c in criterios:
        for b in bases:
            df_html = calcular_matriz_con_porcentajes(df_geiser_crudo, mes, mat, c, b)
            if not df_html.empty:
                html_table = df_html.to_html(index=False, classes="data-table", border=0)
                html_table = html_table.replace('<table', f'<table id="table_{id_p}_{c}_{b}" style="display:none;"')
            else:
                html_table = f'<table id="table_{id_p}_{c}_{b}" class="data-table" style="display:none;"><tr><td>No hay datos para esta combinación</td></tr></table>'
            todas_las_tablas_html += html_table + "\n"

ruta_dashboard = os.path.join(DIR_SALIDA, "Dashboard_NoFinalizadas.html")
with open(ruta_dashboard, "w", encoding="utf-8") as file:
    file.write(html_template.replace("{tables_html}", todas_las_tablas_html))


# ================= REQ 5: QUINTILES DE RESULTADOS (FEBRERO) (REVERTIDO AL CRUDO) =================
print("\n[+] Generando Reporte de Quintiles (Prueba de Resultados - Febrero)...")

def generar_hoja_quintiles(df_r1_crudo, materias, grados, col_media):
    if materias: df_f = df_r1_crudo[(df_r1_crudo['Materia'].isin(materias)) & (df_r1_crudo['grado_num'].isin(grados))].copy()
    else: df_f = df_r1_crudo[df_r1_crudo['grado_num'].isin(grados)].copy()
        
    if df_f.empty: return pd.DataFrame(columns=['Código de infraestructura', 'Nombre del centro', col_media, 'Quintil'])
    
    # EL VIEJO ALGORITMO: Promedio global de la bolsa cruda y redondeo directo a 1 decimal
    res = df_f.groupby(['codigo', 'Centro_Nom'])['Puntaje'].mean().reset_index()
    res.rename(columns={'codigo': 'Código de infraestructura', 'Centro_Nom': 'Nombre del centro', 'Puntaje': col_media}, inplace=True)
    res[col_media] = res[col_media].round(1) 
    
    if len(res) >= 5: res['Quintil'] = pd.qcut(res[col_media].rank(method='first'), 5, labels=['1', '2', '3', '4', '5'])
    else: res['Quintil'] = "N/A"
        
    return res

df_r1_master_crudo = df_geiser_crudo[df_geiser_crudo['Mes'] == 'R1'].copy()

hojas_quintiles = [
    ('Lengua primer ciclo', ['Len'], [3], 'Media de Lengua'),
    ('Lengua segundo ciclo', ['Len'], [4, 5, 6], 'Media de Lengua'),
    ('Matemática primer ciclo', ['Mat'], [3], 'Media de Matemática'),
    ('Matemática segundo ciclo', ['Mat'], [4, 5, 6], 'Media de Matemática'),
    ('Consolidado primer ciclo', ['Len', 'Mat'], [3], 'Media Consolidada'),
    ('Consolidado segundo ciclo', ['Len', 'Mat'], [4, 5, 6], 'Media Consolidada')
]

ruta_quintiles = os.path.join(DIR_SALIDA, "Reporte_Quintiles_Resultados_Feb.xlsx")
with pd.ExcelWriter(ruta_quintiles) as writer:
    for nombre_hoja, mats, grados, col_name in hojas_quintiles:
        df_q = generar_hoja_quintiles(df_r1_master_crudo, mats, grados, col_name)
        if df_q.empty: pd.DataFrame({'Info': [f'Sin datos para {nombre_hoja}']}).to_excel(writer, sheet_name=nombre_hoja, index=False)
        else:
            df_q = df_q.sort_values(by=col_name, ascending=False)
            df_q.to_excel(writer, sheet_name=nombre_hoja, index=False)


# ================= REQ 6: DISTRIBUCIÓN POR RANGOS (PROGRESO MES 2) =================
print("\n[+] Generando Reporte de Rangos (Prueba Progreso Mes 2)...")
df_m2 = df_geiser_validos[df_geiser_validos['Mes'] == 'M2'].copy()

if not df_m2.empty:
    bins = [-np.inf, 20, 40, np.inf]
    etiquetas = ['0 a 20', '>20 a 40', '>40']
    
    df_m2_mat = df_m2[df_m2['Materia'] == 'Mat']['Puntaje']
    cat_mat = pd.cut(df_m2_mat, bins=bins, labels=etiquetas).value_counts().reindex(etiquetas)
    res_mat = pd.DataFrame({'Cantidad de Estudiantes': cat_mat, 'Porcentaje (%)': (cat_mat / len(df_m2_mat) * 100).round(2)})
    res_mat.index.name = 'Rango de Puntaje'
    
    df_m2_len = df_m2[df_m2['Materia'] == 'Len']['Puntaje']
    cat_len = pd.cut(df_m2_len, bins=bins, labels=etiquetas).value_counts().reindex(etiquetas)
    res_len = pd.DataFrame({'Cantidad de Estudiantes': cat_len, 'Porcentaje (%)': (cat_len / len(df_m2_len) * 100).round(2)})
    res_len.index.name = 'Rango de Puntaje'
    
    df_m2_pivot = df_m2.pivot_table(index='nie', columns='Materia', values='Puntaje', aggfunc='mean')
    
    if 'Mat' in df_m2_pivot.columns and 'Len' in df_m2_pivot.columns:
        df_m2_pivot = df_m2_pivot.dropna(subset=['Mat', 'Len'])
        df_m2_pivot['Promedio'] = df_m2_pivot[['Mat', 'Len']].mean(axis=1)
        cat_prom = pd.cut(df_m2_pivot['Promedio'], bins=bins, labels=etiquetas).value_counts().reindex(etiquetas)
        res_prom = pd.DataFrame({'Cantidad de Estudiantes': cat_prom, 'Porcentaje (%)': (cat_prom / len(df_m2_pivot) * 100).round(2)})
        res_prom.index.name = 'Rango de Puntaje'
    else:
        res_prom = pd.DataFrame({'Cantidad de Estudiantes': [], 'Porcentaje (%)': []})
        res_prom.index.name = 'Rango de Puntaje'

    print("\n-----------------------------------------------------")
    print("      RANGOS DE PUNTAJES - PROGRESO MES 2 (ABRIL)     ")
    print("-----------------------------------------------------")
    print("\n--- MATEMÁTICA ---")
    print(res_mat.to_string())
    print("\n--- LENGUA ---")
    print(res_len.to_string())
    print("\n--- PROMEDIO GENERAL (Matemática + Lengua) ---")
    if not res_prom.empty: print(res_prom.to_string())
    else: print("Datos insuficientes para cruzar ambas materias.")
    print("-----------------------------------------------------")

    ruta_rangos = os.path.join(DIR_SALIDA, "Reporte_Rangos_Progreso2.xlsx")
    with pd.ExcelWriter(ruta_rangos) as writer:
        res_mat.reset_index().to_excel(writer, sheet_name='Matemática', index=False)
        res_len.reset_index().to_excel(writer, sheet_name='Lengua', index=False)
        if not res_prom.empty: res_prom.reset_index().to_excel(writer, sheet_name='Promedio General', index=False)


# ================= REQ 7: ANÁLISIS DE COBERTURA DE GRADOS (2° A 6°) =================
print("\n[+] Analizando cobertura de grados a nivel nacional...")
all_schools = set(df_geiser_crudo['codigo'].unique())
schools_2_6 = set(df_geiser_crudo[df_geiser_crudo['grado_num'].isin([2, 3, 4, 5, 6])]['codigo'].unique())
schools_only_upper = all_schools - schools_2_6

print("\n-----------------------------------------------------")
print("      COBERTURA DE GRADOS (HISTÓRICO R1, M1, M2)      ")
print("-----------------------------------------------------")
print(f"Número total de escuelas evaluadas en el sistema: {len(all_schools)}")
print(f"Escuelas que NO reportaron grados de 2° a 6° (solo 7° en adelante): {len(schools_only_upper)}")
print("-----------------------------------------------------")

print("\n[*] ¡Misión Cumplida! Tu Súper-Script ha generado todo exitosamente.")