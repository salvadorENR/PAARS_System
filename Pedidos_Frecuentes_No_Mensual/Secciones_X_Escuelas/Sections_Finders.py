"""
Genera el listado de secciones y el conteo por escuela, PERO solo para
un conjunto seleccionado de escuelas (lista CODIGOS_SELECCIONADOS).

Si CODIGOS_SELECCIONADOS se deja vacío ([]), procesa TODAS las escuelas.
El script avisa en consola si alguno de los códigos solicitados no existe
en el archivo de entrada.

Requisitos:  pip install pandas openpyxl
"""

import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

# --- Configuración de rutas ---------------------------------------------------
CARPETA = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Secciones_X_Escuelas"
ARCHIVO_ENTRADA = "Recolección_de_datos_G1-G2_25_05_2026.xlsx"
ARCHIVO_SALIDA  = "Listado_Secciones_Escuelas_Seleccionadas.xlsx"
HOJA = "DATA SIGES"

# Escuelas a incluir (como texto). Dejar [] para incluir todas.
CODIGOS_SELECCIONADOS = [
    "10235", "11108", "11244", "11870", "11935", "11943", "11954", "11988",
    "12012", "12049", "74005", "11406", "11377", "11883", "11410", "11330",
]

COLUMNAS = ["GRUPO", "CODIGO", "NOMBRE", "GRADO", "NOMBRE_SECCION", "CODIGO_SECCION"]

ruta_entrada = os.path.join(CARPETA, ARCHIVO_ENTRADA)
ruta_salida  = os.path.join(CARPETA, ARCHIVO_SALIDA)

# --- Lectura ------------------------------------------------------------------
df = pd.read_excel(ruta_entrada, sheet_name=HOJA,
                   dtype={"CODIGO": str, "CODIGO_SECCION": str})

# --- Filtro por escuelas ------------------------------------------------------
if CODIGOS_SELECCIONADOS:
    presentes = set(df["CODIGO"].unique())
    faltantes = [c for c in CODIGOS_SELECCIONADOS if c not in presentes]
    if faltantes:
        print("ADVERTENCIA: códigos no encontrados en el archivo:", ", ".join(faltantes))
    df = df[df["CODIGO"].isin(CODIGOS_SELECCIONADOS)]

# --- Hoja 1: listado de secciones (una fila por CODIGO_SECCION) ---------------
listado = (df[COLUMNAS]
           .drop_duplicates(subset="CODIGO_SECCION")
           .sort_values(["CODIGO", "GRADO", "NOMBRE_SECCION"])
           .reset_index(drop=True))

# --- Hoja 2: conteo de secciones por escuela ----------------------------------
conteo = (df.groupby("CODIGO")
            .agg(NOMBRE=("NOMBRE", "first"),
                 NUMERO_SECCIONES=("CODIGO_SECCION", "nunique"))
            .reset_index()
            .sort_values("CODIGO"))

# --- Exportar -----------------------------------------------------------------
with pd.ExcelWriter(ruta_salida, engine="openpyxl") as xl:
    listado.to_excel(xl, sheet_name="Listado de Secciones", index=False)
    conteo.to_excel(xl, sheet_name="Conteo por Escuela", index=False)

# --- Formato ------------------------------------------------------------------
wb = load_workbook(ruta_salida)
enc_fuente = Font(name="Arial", bold=True, color="FFFFFF")
enc_fondo  = PatternFill("solid", start_color="305496")
anchos = {
    "Listado de Secciones": {"A": 10, "B": 10, "C": 55, "D": 18, "E": 16, "F": 16},
    "Conteo por Escuela":   {"A": 12, "B": 55, "C": 18},
}
for nombre, w in anchos.items():
    ws = wb[nombre]
    for celda in ws[1]:
        celda.font = enc_fuente
        celda.fill = enc_fondo
        celda.alignment = Alignment(horizontal="center")
    for fila in ws.iter_rows(min_row=2):
        for celda in fila:
            celda.font = Font(name="Arial")
    for col, ancho in w.items():
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
wb.save(ruta_salida)

# --- Resumen en consola -------------------------------------------------------
print(f"Escuelas incluidas:              {len(conteo)}")
print(f"Secciones (filas en el listado): {len(listado)}")
print(f"Archivo generado:                {ruta_salida}")