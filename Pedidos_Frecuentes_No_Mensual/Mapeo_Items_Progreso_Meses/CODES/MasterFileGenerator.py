import os
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# 1. Define Directory Paths
DIR_MAT = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Mapeo_Items_Progreso_Meses\Achivement_Indicators\MATEMÁTICA"
DIR_LEN = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Mapeo_Items_Progreso_Meses\Achivement_Indicators\LENGUA"
DIR_OUT = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Mapeo_Items_Progreso_Meses\Master_Files"

# Ensure the output directory exists
os.makedirs(DIR_OUT, exist_ok=True)

# 2. Helper Functions
def extract_grade_from_filename(filename):
    """Extracts and normalizes the grade strictly as a number + '°'."""
    lower_name = filename.lower()
    
    # Handle Bachillerato mapping to 10° and 11°
    if "bachillerato" in lower_name:
        if "1" in lower_name:
            return "10°"
        if "2" in lower_name:
            return "11°"
            
    # Handle specific cases where 10 or 11 might just be numbers in the file name
    if "10°" in lower_name or "10" in lower_name:
        return "10°"
    if "11°" in lower_name or "11" in lower_name:
        return "11°"
        
    # Extract the first number found for grades 1-9
    match = re.search(r'(\d+)', lower_name)
    if match:
        num = match.group(1)
        return f"{num}°"
        
    return "N/A"

def process_excel_directory(directory, subject_code):
    """Reads all '.xlsx' files, extracts target columns, cleans, and sorts."""
    master_data = []
    
    if not os.path.exists(directory):
        print(f"Directory not found: {directory}")
        return pd.DataFrame()
        
    for filename in os.listdir(directory):
        if not filename.endswith(".xlsx") or filename.startswith("~"):
            continue
            
        filepath = os.path.join(directory, filename)
        grade_clean = extract_grade_from_filename(filename)
        
        try:
            xls = pd.read_excel(filepath, sheet_name=None, header=None)
        except Exception as e:
            print(f"Error reading {filename}: {e}")
            continue
            
        for sheet_name, raw_df in xls.items():
            lower_sheet = sheet_name.lower()
            
            # Identify if it's a "Pedido" sheet
            if not re.search(r'(pedido|p\s*\d+)', lower_sheet):
                continue
                
            month_match = re.search(r'(\d+)', lower_sheet)
            month = month_match.group(1) if month_match else "N/A"
            
            header_idx = -1
            for i, row in raw_df.iterrows():
                row_str = " ".join([str(val).lower() for val in row if pd.notna(val)])
                if "indicador" in row_str:
                    header_idx = i
                    break
            
            if header_idx == -1:
                continue 
                
            df = raw_df.iloc[header_idx+1:].copy()
            df.columns = raw_df.iloc[header_idx]
            
            item_col = next((col for col in df.columns if pd.notna(col) and ('ítem' in str(col).lower() or 'item' in str(col).lower())), None)
            ind_col = next((col for col in df.columns if pd.notna(col) and ('indicador' in str(col).lower())), None)
            
            if item_col and ind_col:
                for _, row in df.iterrows():
                    item_val = row[item_col]
                    ind_val = row[ind_col]
                    
                    if pd.notna(item_val) and pd.notna(ind_val) and str(item_val).strip() != "":
                        master_data.append({
                            "Subject": subject_code,
                            "Grado": grade_clean,
                            "Mes": month,
                            "Código de ítem": str(item_val).strip(),
                            "Indicador de logro": str(ind_val).strip()
                        })
                        
    df_master = pd.DataFrame(master_data)
    
    # Apply numerical sorting by Grade and Month
    if not df_master.empty:
        df_master['Grade_Num'] = df_master['Grado'].str.replace('°', '').astype(int)
        df_master['Mes_Num'] = pd.to_numeric(df_master['Mes'], errors='coerce').fillna(0)
        
        df_master = df_master.sort_values(by=['Grade_Num', 'Mes_Num']).drop(columns=['Grade_Num', 'Mes_Num'])
        
    return df_master

def save_and_format_excel(df, output_path):
    """Saves the DataFrame to Excel and applies alternating colors by Grade."""
    if df.empty:
        print(f"No data to save for {output_path}")
        return
        
    # First, save the basic DataFrame to Excel
    df.to_excel(output_path, index=False)
    
    # Load the workbook to apply formatting
    wb = load_workbook(output_path)
    ws = wb.active
    
    # Define our two alternating background colors
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_blue = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid") # Light Blue
    
    current_fill = fill_white
    previous_grade = None
    
    # Column 2 is "Grado"
    grade_col_index = 2 
    
    # Iterate through rows starting from row 2 (skipping the header)
    for row in range(2, ws.max_row + 1):
        grade_cell_value = ws.cell(row=row, column=grade_col_index).value
        
        if previous_grade is None:
            previous_grade = grade_cell_value
            
        # Toggle the color if the grade changes
        if grade_cell_value != previous_grade:
            current_fill = fill_blue if current_fill == fill_white else fill_white
            previous_grade = grade_cell_value
            
        # Apply the selected color to the entire row
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).fill = current_fill
            
    wb.save(output_path)
    print(f"Success! Master file sorted, color-formatted, and saved at: {output_path}")

# 3. Execution and Export
print("Processing MATEMÁTICA files...")
df_mat = process_excel_directory(DIR_MAT, "MAT")
mat_output_path = os.path.join(DIR_OUT, "Master_MAT.xlsx")
save_and_format_excel(df_mat, mat_output_path)

print("\nProcessing LENGUA files...")
df_len = process_excel_directory(DIR_LEN, "LEN")
len_output_path = os.path.join(DIR_OUT, "Master_LEN.xlsx")
save_and_format_excel(df_len, len_output_path)