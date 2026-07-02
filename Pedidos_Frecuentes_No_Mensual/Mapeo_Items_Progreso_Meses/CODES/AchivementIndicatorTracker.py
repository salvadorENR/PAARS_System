import os
import re
import pandas as pd
import unicodedata
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# 1. Definir Rutas de Directorios
DIR_MASTER = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Mapeo_Items_Progreso_Meses\Master_Files"
DIR_CURRICULUM = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Mapeo_Items_Progreso_Meses\Priorized_MAT_LEC_Sequence"
DIR_OUT = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Mapeo_Items_Progreso_Meses\Tracking_Results"

os.makedirs(DIR_OUT, exist_ok=True)

# 2. Mapeo de Rangos Esperados
RANGOS_ESPERADOS = {
    1: (1, 20),
    2: (21, 40),
    3: (41, 60),
    4: (61, 80),
    5: (81, 100),
    6: (101, 120),
    7: (121, 140),
    8: (141, 160)
}

def get_expected_range_str(mes):
    try:
        mes = int(mes)
        if mes in RANGOS_ESPERADOS:
            return f"{RANGOS_ESPERADOS[mes][0]}-{RANGOS_ESPERADOS[mes][1]}"
    except:
        pass
    return "N/A"

def evaluar_estado(dias_list, mes):
    """
    Evalúa solo el PRIMER día de clase de la lista contra el rango esperado.
    Retorna el 'Estado' (Alerta 1, Alerta 2, En rango).
    """
    if not dias_list:
        return "Sin Día"

    try:
        mes_num = int(str(mes).strip())
        if mes_num not in RANGOS_ESPERADOS:
            return "N/A"
    except ValueError:
        return "N/A"

    inicio, fin = RANGOS_ESPERADOS[mes_num]

    # Se evalúa estrictamente el PRIMER día (el día que se introduce el tema)
    primer_dia = dias_list[0]

    if primer_dia < inicio:
        return "Alerta 1"
    elif primer_dia > fin:
        return "Alerta 2"
    else:
        return "En rango"

def normalize_math_chars(text):
    """Reemplaza superíndices y estandariza operadores matemáticos."""
    replacements = {
        '²': '2', '³': '3', '¹': '1', '⁰': '0',
        '⁴': '4', '⁵': '5', '⁶': '6', '⁷': '7', '⁸': '8', '⁹': '9',
        '⁺': '+', '⁻': '-', '×': 'x', '÷': '/', '–': '-', '—': '-', '−': '-'
    }
    for k, v in replacements.items():
        text = text.replace(k, v)
    return text

def clean_text_for_match(text, remove_all_spaces=False):
    """
    Limpia el texto eliminando (U1), acentos, puntuación y normalizando espacios.

    FIXES APPLIED:
    1. NFKD normalization (before NFD) converts Unicode math-italic letters such as
       𝑔 (U+1D454), 𝑥 (U+1D465), 𝑎 (U+1D44E), 𝑓 (U+1D453) — common in curriculum
       files exported from PDFs — into their plain ASCII equivalents (g, x, a, f).
       Without this, indicators like "4.2 Grafica g(x) = ax³" failed to match the
       curriculum row "4.2 Grafi ca 𝑔(𝑥) = 𝑎𝑥³" because the italic Unicode letters
       were invisible to the word-overlap and solid-string checks.

    2. The standalone high-overlap threshold (Logic 3) was raised from 0.80 to 0.95.
       NFKD also boosted the similarity between near-identical but semantically opposite
       indicators (e.g. "proporcionalidad directa y=ax" vs "proporcionalidad inversa
       y=a/x") from 0.65 to 0.81, which would have created false positives under the
       old 0.80 threshold. The higher threshold keeps Logic 3 as a last-resort safety
       net for near-identical texts only, while Logic 1 (code + overlap >= 0.40) and
       Logic 2 (solid-string containment) handle all real matches.
    """
    if pd.isna(text):
        return ""
    text = str(text).lower()
    text = normalize_math_chars(text)

    # FIX 1: NFKD maps Unicode math-italic/bold/script letters to plain ASCII.
    # Must run BEFORE NFD so the decomposed ASCII letters are then processed normally.
    text = unicodedata.normalize('NFKD', text)

    # Eliminar etiquetas como (U1), (U 12), (Unidad 3), etc.
    text = re.sub(r'\(\s*[uU](nidad)?\s*\d+\s*\)', '', text)

    # Eliminar acentos
    text = ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')

    # Eliminar signos de puntuación
    text = re.sub(r'[^\w\s]', '', text)

    if remove_all_spaces:
        return re.sub(r'\s+', '', text)
    else:
        return re.sub(r'\s+', ' ', text).strip()

def extract_indicator_code(text):
    """Extrae el código numérico/alfanumérico (ej. '1.5' o 'A.2') del indicador."""
    match = re.search(r'^([A-Za-z]\.\d+|\d+\.\d+)', str(text).strip())
    return match.group(1) if match else None

def normalize_sheet_grade(sheet_name):
    """Normaliza los nombres de las hojas del currículo para que coincidan con el Grado."""
    lower_name = str(sheet_name).lower()
    if any(x in lower_name for x in ["10", "primer", "1er", "1.er"]):
        return "10°"
    if any(x in lower_name for x in ["11", "segundo", "2do", "2.o"]):
        return "11°"

    match = re.search(r'(\d+)', lower_name)
    if match:
        return f"{match.group(1)}°"
    return "N/A"

def find_curriculum_file(subject_code):
    for file in os.listdir(DIR_CURRICULUM):
        if not file.endswith(".xlsx") or file.startswith("~"):
            continue
        if subject_code == "MAT" and "Math" in file:
            return os.path.join(DIR_CURRICULUM, file)
        if subject_code == "LEN" and "Language" in file:
            return os.path.join(DIR_CURRICULUM, file)
    return None

def process_tracking(master_path, subject_code):
    """Cruza el archivo Master con el Currículo con búsqueda dinámica de encabezados."""
    if not os.path.exists(master_path):
        print(f"Master file no encontrado: {master_path}")
        return pd.DataFrame()

    curriculum_path = find_curriculum_file(subject_code)
    if not curriculum_path:
        print(f"Archivo de currículo para {subject_code} no encontrado.")
        return pd.DataFrame()

    df_master = pd.read_excel(master_path)

    # IMPORTANTE: header=None permite leer desde la fila 0 para encontrar dinámicamente la tabla real
    curriculum_xls = pd.read_excel(curriculum_path, sheet_name=None, header=None)

    curriculum_sheets = {}
    for sheet_name, raw_df in curriculum_xls.items():
        grade_key = normalize_sheet_grade(sheet_name)

        # ESCÁNER DINÁMICO: Buscar la fila que contiene los encabezados
        header_idx = -1
        for i, row in raw_df.iterrows():
            row_str = " ".join([str(val).lower() for val in row if pd.notna(val)])
            if "indicador" in row_str or "contenido" in row_str:
                header_idx = i
                break

        # Si encuentra el encabezado, recorta la tabla a partir de esa fila
        if header_idx != -1:
            df_clean = raw_df.iloc[header_idx+1:].copy()
            df_clean.columns = raw_df.iloc[header_idx]
            curriculum_sheets[grade_key] = df_clean

    results = []

    for _, row in df_master.iterrows():
        grado = str(row.get("Grado", "")).strip()
        mes = str(row.get("Mes", "")).strip()
        codigo_item = str(row.get("Código de ítem", "")).strip()
        indicador_master = str(row.get("Indicador de logro", "")).strip()

        dias_encontrados_num = set()

        master_clean = clean_text_for_match(indicador_master, False)
        master_solid = clean_text_for_match(indicador_master, True)
        code_master = extract_indicator_code(indicador_master)

        if grado in curriculum_sheets:
            df_curr = curriculum_sheets[grado]

            col_dia = next((c for c in df_curr.columns if pd.notna(c) and ('dia' in str(c).lower() or 'día' in str(c).lower() or 'clase' in str(c).lower())), None)

            col_ind = next((c for c in df_curr.columns if pd.notna(c) and 'indicador' in str(c).lower()), None)
            if not col_ind:
                col_ind = next((c for c in df_curr.columns if pd.notna(c) and 'contenido' in str(c).lower()), None)

            if col_dia and col_ind:
                for _, curr_row in df_curr.iterrows():
                    curr_dia_val = curr_row[col_dia]
                    curr_ind_val = str(curr_row[col_ind]).strip()

                    if pd.isna(curr_dia_val) or not curr_ind_val or str(curr_dia_val).strip() == "":
                        continue

                    curr_clean = clean_text_for_match(curr_ind_val, False)
                    curr_solid = clean_text_for_match(curr_ind_val, True)
                    curr_code = extract_indicator_code(curr_ind_val)

                    is_match = False

                    if master_clean and curr_clean:
                        words_m = set(master_clean.split())
                        words_c = set(curr_clean.split())
                        overlap_ratio = 0
                        if words_m and words_c:
                            intersection = len(words_m.intersection(words_c))
                            union = len(words_m.union(words_c))
                            overlap_ratio = intersection / union

                        # Lógica 1: Coincidencia de código y superposición de palabras
                        if code_master and curr_code and (code_master == curr_code):
                            if overlap_ratio >= 0.40:
                                is_match = True

                        # Lógica 2: Match tipo "Rayos X" (Ignora espacios y símbolos)
                        if not is_match and master_solid and curr_solid:
                            if master_solid in curr_solid or curr_solid in master_solid:
                                is_match = True

                        # Lógica 3: Alta similitud sin código.
                        # FIX 2: Umbral elevado de 0.80 → 0.95 para evitar falsos positivos
                        # entre indicadores semánticamente distintos cuya similitud léxica
                        # fue aumentada artificialmente por la normalización NFKD del Fix 1.
                        # Ejemplo: "proporcionalidad directa y=ax" vs "proporcionalidad
                        # inversa y=a/x" pasaron de overlap=0.65 a 0.81 con NFKD, lo que
                        # habría cruzado el antiguo umbral de 0.80 generando una coincidencia
                        # errónea. Con 0.95 solo textos prácticamente idénticos califican.
                        if not is_match and overlap_ratio >= 0.95:
                            is_match = True

                    if is_match:
                        day_matches = re.findall(r'\d+\.?\d*', str(curr_dia_val))
                        for dm in day_matches:
                            try:
                                dia_num = float(dm)
                                if dia_num.is_integer():
                                    dias_encontrados_num.add(int(dia_num))
                                else:
                                    dias_encontrados_num.add(dia_num)
                            except ValueError:
                                pass

        dias_sorted = sorted(list(dias_encontrados_num))
        dia_str = ", ".join(map(str, dias_sorted)) if dias_sorted else "Sin Día"

        rango_esperado = get_expected_range_str(mes)
        estado = evaluar_estado(dias_sorted, mes)

        results.append({
            "Grado": grado,
            "Mes": mes,
            "Código de ítem": codigo_item,
            "Indicador de Logro": indicador_master,
            "Día(clase)": dia_str,
            "Rango de clases esperado": rango_esperado,
            "Estado": estado
        })

    return pd.DataFrame(results)

def save_and_format_tracking(df, output_path):
    """Guarda a Excel y aplica los formatos de color a la columna 'Estado'."""
    if df.empty:
        return

    df.to_excel(output_path, index=False)

    wb = load_workbook(output_path)
    ws = wb.active

    # Definición de Colores
    fill_alerta1 = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid") # Naranja para Alerta 1
    font_alerta1 = Font(color="000000", bold=True)

    fill_alerta2 = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid") # Rojo para Alerta 2
    font_alerta2 = Font(color="FFFFFF", bold=True)

    fill_ok = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid") # Verde para En Rango
    font_ok = Font(color="000000", bold=True)

    estado_col_idx = None
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=col).value == "Estado":
            estado_col_idx = col
            break

    if estado_col_idx:
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=estado_col_idx)
            val = cell.value

            if val == "Alerta 1":
                cell.fill = fill_alerta1
                cell.font = font_alerta1
            elif val == "Alerta 2":
                cell.fill = fill_alerta2
                cell.font = font_alerta2
            elif val == "En rango":
                cell.fill = fill_ok
                cell.font = font_ok

    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 60)

    wb.save(output_path)
    print(f"¡Éxito! Archivo de tracking guardado: {output_path}")

# 3. Ejecución
master_mat_path = os.path.join(DIR_MASTER, "Master_MAT.xlsx")
master_len_path = os.path.join(DIR_MASTER, "Master_LEN.xlsx")

print("Mapeando Tracking de MATEMÁTICA...")
df_track_mat = process_tracking(master_mat_path, "MAT")
save_and_format_tracking(df_track_mat, os.path.join(DIR_OUT, "Tracking_MAT.xlsx"))

print("Mapeando Tracking de LENGUA...")
df_track_len = process_tracking(master_len_path, "LEN")
save_and_format_tracking(df_track_len, os.path.join(DIR_OUT, "Tracking_LEN.xlsx"))

print("\n¡Operación completada exitosamente!")