import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy
import os

BASE_DIR = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Agregando_Columna_Sección"

B1_PATH  = os.path.join(BASE_DIR, "Centros_Escolares_B1.xlsx")
B2_PATH  = os.path.join(BASE_DIR, "Centros_Escolares_B2.xlsx")
MAT_PATH = os.path.join(BASE_DIR, "MatriculaProgresoMes3.csv")

OUT_B1   = os.path.join(BASE_DIR, "Centros_Escolares_B1_Actualizado.xlsx")
OUT_B2   = os.path.join(BASE_DIR, "Centros_Escolares_B2_Actualizado.xlsx")

# ── Load Matricula lookup ────────────────────────────────────────────────────
mat = pd.read_csv(MAT_PATH, sep=";", dtype=str)
mat["NIE_CLEAN"] = mat["NIE"].str.lstrip("'").str.strip()
nie_to_seccion = dict(zip(mat["NIE_CLEAN"], mat["NOMBRE_SECCIÓN"]))


def build_flat_headers(ws):
    """
    Reads the two-row merged header and returns a flat list of column header
    strings (one per column, 1-indexed dict: {col_index: header_string}).
    
    Logic:
      - Columns merged vertically (e.g. A1:A2): the row-1 value is the header.
      - Columns merged horizontally (e.g. M1:R1 for 'Matemática'): each
        sub-column gets  "<group> - <subheader from row 2>".
    """
    # Collect merge ranges
    vertical_merged = set()   # cols where row1 value spans row2 (single-col merges)
    horizontal_groups = {}    # col -> group name  (multi-col row-1 merges)

    for merged_range in ws.merged_cells.ranges:
        min_col, max_col = merged_range.min_col, merged_range.max_col
        min_row, max_row = merged_range.min_row, merged_range.max_row
        if min_col == max_col:          # vertical merge (single column)
            vertical_merged.add(min_col)
        else:                           # horizontal merge (group header)
            group_name = ws.cell(min_row, min_col).value
            for c in range(min_col, max_col + 1):
                horizontal_groups[c] = group_name

    flat = {}
    for c in range(1, ws.max_column + 1):
        if c in vertical_merged:
            flat[c] = ws.cell(1, c).value
        elif c in horizontal_groups:
            subheader = ws.cell(2, c).value or ""
            flat[c] = f"{horizontal_groups[c]} - {subheader}"
        else:
            # No merge at all — use row 1 value
            flat[c] = ws.cell(1, c).value
    return flat


def find_col(flat_headers, name_lower):
    for col, val in flat_headers.items():
        if val and str(val).strip().lower() == name_lower:
            return col
    return None


def copy_cell_style(src, dst):
    if src.has_style:
        dst.font        = copy(src.font)
        dst.fill        = copy(src.fill)
        dst.border      = copy(src.border)
        dst.alignment   = copy(src.alignment)
        dst.number_format = src.number_format


def process_file(input_path, output_path):
    wb = load_workbook(input_path)
    ws = wb.active

    # 1. Build flat header map BEFORE touching anything
    flat_headers = build_flat_headers(ws)

    grado_col = find_col(flat_headers, "grado")
    nie_col   = find_col(flat_headers, "nie")
    if grado_col is None or nie_col is None:
        raise ValueError(f"Could not find 'Grado' or 'NIE' in {input_path}")

    # 2. Unmerge all cells
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))

    # 3. Write flat headers into row 1
    for col, header in flat_headers.items():
        cell = ws.cell(1, col)
        cell.value = header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 4. Clear row 2 (old sub-headers) and shift data rows up by 1
    ws.delete_rows(2)

    # After deleting row 2, data starts at row 2 now.
    # nie_col and grado_col positions are unchanged (column indices didn't move).

    # 5. Insert new "Sección" column right after Grado
    insert_col = grado_col + 1
    ws.insert_cols(insert_col)

    # Write its header
    h_cell = ws.cell(1, insert_col)
    h_cell.value = "Sección"
    h_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # Copy style from adjacent header
    copy_cell_style(ws.cell(1, grado_col), h_cell)

    # 6. Fill Sección for every data row (row 2 onward after the delete)
    for row_idx in range(2, ws.max_row + 1):
        nie_val = ws.cell(row_idx, nie_col).value
        nie_str = str(nie_val).strip().lstrip("'") if nie_val is not None else ""
        ws.cell(row_idx, insert_col).value = nie_to_seccion.get(nie_str, "")

    # 7. Freeze pane on row 2 so header stays visible when scrolling
    ws.freeze_panes = "A2"

    # 8. Auto-width for all columns (approximate)
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    wb.save(output_path)
    print(f"Saved → {output_path}")


process_file(B1_PATH, OUT_B1)
process_file(B2_PATH, OUT_B2)
print("Done.")