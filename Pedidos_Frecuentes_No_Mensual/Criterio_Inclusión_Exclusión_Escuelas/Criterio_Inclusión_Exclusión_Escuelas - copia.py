import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# 1. LOAD & FILTER DATA
# ─────────────────────────────────────────────
INPUT_FILE  = r'C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Excel_Maestro_HojaUnica.csv'
OUTPUT_FILE = r'C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Centros_Escolares_B1.xlsx'

df = pd.read_csv(INPUT_FILE, sep=';')

# ── Determine the definitive group for each school ──────────────────
# The school's group is the MAJORITY value of GRUPO_Limpio across all
# its student rows. Any stray rows with a different value are data
# entry errors and are corrected to the school's definitive group.
school_group = (
    df.groupby('CODIGO')['GRUPO_Limpio']
    .agg(lambda x: x.value_counts().idxmax())   # majority vote
    .reset_index()
    .rename(columns={'GRUPO_Limpio': 'GRUPO_Escuela'})
)

# Merge the definitive school group back onto every student row
df = df.merge(school_group, on='CODIGO', how='left')

# Keep only students whose school definitively belongs to B1
valid_b1 = df[df['GRUPO_Escuela'] == 'B1'].copy()

# ─────────────────────────────────────────────
# 2. SCHOOL TYPE CLASSIFICATION (mutually exclusive)
# ─────────────────────────────────────────────
def classify_school(name):
    n = str(name).upper()
    if 'INSTITUTO' in n:
        return 'Instituto'
    elif 'COMPLEJO EDUCATIVO' in n:
        return 'Complejo Educativo'
    else:
        return 'Centro Escolar'

valid_b1['Tipo_Escuela'] = valid_b1['NOMBRE'].apply(classify_school)

# ─────────────────────────────────────────────
# 3. CONSTANTS
# ─────────────────────────────────────────────
NIVEL_COLS = {
    'mat': [
        'Nivel Matemática (Mes 1 (Marzo))',
        'Nivel Matemática (Mes 2 (Abril))',
        'Nivel Matemática (Mes 3 (Mayo))',
    ],
    'len': [
        'Nivel Lengua (Mes 1 (Marzo))',
        'Nivel Lengua (Mes 2 (Abril))',
        'Nivel Lengua (Mes 3 (Mayo))',
    ]
}

SCORE_COLS = {
    'mat': [
        'Puntaje Matemática (Mes 1 (Marzo))',
        'Puntaje Matemática (Mes 2 (Abril))',
        'Puntaje Matemática (Mes 3 (Mayo))',
    ],
    'len': [
        'Puntaje Lengua (Mes 1 (Marzo))',
        'Puntaje Lengua (Mes 2 (Abril))',
        'Puntaje Lengua (Mes 3 (Mayo))',
    ]
}

GRADE_ORDER = ['2º', '3º', '4º', '5º', '6º', '7°', '8°', '9°', '10°', '11°']
GRADE_LABELS = {
    '2º':  '2° Grado',
    '3º':  '3° Grado',
    '4º':  '4° Grado',
    '5º':  '5° Grado',
    '6º':  '6° Grado',
    '7°':  '7° Grado',
    '8°':  '8° Grado',
    '9°':  '9° Grado',
    '10°': '10° (1er Año Bach.)',
    '11°': '11° (2do Año Bach.)',
}

STATUS_ORDER = ['Alerta', 'Regular', 'Bueno', 'Excelente']

# ─────────────────────────────────────────────
# 4. CRITICAL % CALCULATIONS
# ─────────────────────────────────────────────
def pct_critico(series):
    """Percentage of 'Crítico' out of non-null values."""
    valid = series.dropna()
    if len(valid) == 0:
        return np.nan
    return (valid == 'Crítico').sum() / len(valid)

def compute_status(mat_wins, len_wins):
    """
    mat_wins, len_wins: int 0-3 (times % crítico <= universe %)
    Returns: 'Excelente', 'Bueno', 'Regular', 'Alerta'
    """
    if mat_wins >= 2 and len_wins >= 2:
        return 'Excelente'
    elif (mat_wins == 3 and len_wins == 1) or (mat_wins == 1 and len_wins == 3):
        return 'Bueno'
    elif (mat_wins == 2 and len_wins == 1) or (mat_wins == 1 and len_wins == 2):
        return 'Regular'
    else:
        return 'Alerta'

# ─────────────────────────────────────────────
# 5. BUILD UNIVERSE % CRÍTICO PER MONTH
# ─────────────────────────────────────────────
# Universe is split by school type
def build_universe(data, tipo):
    """Returns dict: {subject: [pct_m1, pct_m2, pct_m3]} for a given school type."""
    subset = data[data['Tipo_Escuela'] == tipo]
    universe = {}
    for subj, cols in NIVEL_COLS.items():
        universe[subj] = [pct_critico(subset[c]) for c in cols]
    return universe

universes = {
    tipo: build_universe(valid_b1, tipo)
    for tipo in ['Instituto', 'Complejo Educativo', 'Centro Escolar']
}

# ─────────────────────────────────────────────
# 6. COMPUTE SCHOOL-LEVEL STATUS
# ─────────────────────────────────────────────
school_records = []

for (codigo, nombre), grp in valid_b1.groupby(['CODIGO', 'NOMBRE']):
    tipo = grp['Tipo_Escuela'].iloc[0]
    univ = universes[tipo]

    # Count wins per subject across 3 months
    wins = {}
    for subj, cols in NIVEL_COLS.items():
        w = 0
        for i, col in enumerate(cols):
            school_pct = pct_critico(grp[col])
            univ_pct   = univ[subj][i]
            if not np.isnan(school_pct) and not np.isnan(univ_pct):
                if school_pct <= univ_pct:
                    w += 1
        wins[subj] = w

    status = compute_status(wins['mat'], wins['len'])

    # Grade-level status
    grade_statuses = {}
    for grade in GRADE_ORDER:
        grade_grp = grp[grp['Grado_Str'] == grade]
        if len(grade_grp) == 0:
            grade_statuses[grade] = ''
            continue
        # Build grade-level universe (same type, same grade)
        grade_univ_data = valid_b1[
            (valid_b1['Tipo_Escuela'] == tipo) &
            (valid_b1['Grado_Str'] == grade)
        ]
        g_wins = {}
        for subj, cols in NIVEL_COLS.items():
            w = 0
            for col in cols:
                s_pct = pct_critico(grade_grp[col])
                u_pct = pct_critico(grade_univ_data[col])
                if not np.isnan(s_pct) and not np.isnan(u_pct):
                    if s_pct <= u_pct:
                        w += 1
            g_wins[subj] = w
        grade_statuses[grade] = compute_status(g_wins['mat'], g_wins['len'])

    school_records.append({
        'Código':        codigo,
        'Centro Escolar': nombre,
        'Tipo':          tipo,
        'Estatus':       status,
        'mat_wins':      wins['mat'],
        'len_wins':      wins['len'],
        **{f'grado_{g}': grade_statuses[g] for g in GRADE_ORDER}
    })

schools_df = pd.DataFrame(school_records)

# Sort: Alerta → Regular → Bueno → Excelente
schools_df['status_sort'] = schools_df['Estatus'].map(
    {s: i for i, s in enumerate(STATUS_ORDER)}
)
schools_df.sort_values('status_sort', inplace=True)
schools_df.drop(columns='status_sort', inplace=True)
schools_df.reset_index(drop=True, inplace=True)


# ─────────────────────────────────────────────
# 6b. BUILD VERIFICATION DATA
# ─────────────────────────────────────────────
# For each school x subject x month: school %, universe %, win flag
verif_records = []

MONTH_LABELS = ['Mes 1 (Marzo)', 'Mes 2 (Abril)', 'Mes 3 (Mayo)']
SUBJ_LABELS  = {'mat': 'Matemática', 'len': 'Lenguaje'}

for (codigo, nombre), grp in valid_b1.groupby(['CODIGO', 'NOMBRE']):
    tipo  = grp['Tipo_Escuela'].iloc[0]
    univ  = universes[tipo]
    estatus = schools_df.loc[schools_df['Código'] == codigo, 'Estatus'].values[0]
    mat_w = schools_df.loc[schools_df['Código'] == codigo, 'mat_wins'].values[0]
    len_w = schools_df.loc[schools_df['Código'] == codigo, 'len_wins'].values[0]

    for subj, cols in NIVEL_COLS.items():
        for i, col in enumerate(cols):
            school_pct = pct_critico(grp[col])
            univ_pct   = univ[subj][i]
            if not np.isnan(school_pct) and not np.isnan(univ_pct):
                win = school_pct <= univ_pct
            else:
                win = None
            verif_records.append({
                'Código':             codigo,
                'Centro Escolar':     nombre,
                'Tipo':               tipo,
                'Estatus':            estatus,
                'Mat Wins':           f"{mat_w}/3",
                'Len Wins':           f"{len_w}/3",
                'Asignatura':         SUBJ_LABELS[subj],
                'Mes':                MONTH_LABELS[i],
                '% Crítico Escuela':  round(school_pct * 100, 2) if not np.isnan(school_pct) else '',
                '% Crítico Universo (Ref.)': round(univ_pct   * 100, 2) if not np.isnan(univ_pct)   else '',
                'Escuela ≤ Universo': ('✔ Sí' if win else ('✘ No' if win is not None else 'N/D')),
            })

verif_df = pd.DataFrame(verif_records)

# Sort same as schools sheet: Alerta → Excelente, then school name, then subject, then month
verif_df['status_sort'] = verif_df['Estatus'].map({s: i for i, s in enumerate(STATUS_ORDER)})
subj_sort = {'Matemática': 0, 'Lenguaje': 1}
mes_sort  = {m: i for i, m in enumerate(MONTH_LABELS)}
verif_df['subj_sort'] = verif_df['Asignatura'].map(subj_sort)
verif_df['mes_sort']  = verif_df['Mes'].map(mes_sort)
verif_df.sort_values(['status_sort', 'Centro Escolar', 'subj_sort', 'mes_sort'], inplace=True)
verif_df.drop(columns=['status_sort', 'subj_sort', 'mes_sort'], inplace=True)
verif_df.reset_index(drop=True, inplace=True)

# ─────────────────────────────────────────────
# 7. BUILD STUDENTS SHEET DATA
# ─────────────────────────────────────────────
students_df = valid_b1[[
    'CODIGO', 'NOMBRE', 'Tipo_Escuela',
    'NIE', 'PRIMER_NOMBRE', 'SEGUNDO_NOMBRE', 'PRIMER_APELLIDO', 'SEGUNDO_APELLIDO',
    'Grado_Str', 'CÓDIGO_SECCIÓN',
    'Puntaje Matemática (Mes 1 (Marzo))', 'Nivel Matemática (Mes 1 (Marzo))',
    'Puntaje Matemática (Mes 2 (Abril))', 'Nivel Matemática (Mes 2 (Abril))',
    'Puntaje Matemática (Mes 3 (Mayo))',  'Nivel Matemática (Mes 3 (Mayo))',
    'Puntaje Lengua (Mes 1 (Marzo))',     'Nivel Lengua (Mes 1 (Marzo))',
    'Puntaje Lengua (Mes 2 (Abril))',     'Nivel Lengua (Mes 2 (Abril))',
    'Puntaje Lengua (Mes 3 (Mayo))',      'Nivel Lengua (Mes 3 (Mayo))',
]].copy()

students_df.rename(columns={
    'CODIGO':          'Código',
    'NOMBRE':          'Centro Escolar',
    'Tipo_Escuela':    'Tipo',
    'NIE':             'NIE',
    'PRIMER_NOMBRE':   'Primer Nombre',
    'SEGUNDO_NOMBRE':  'Segundo Nombre',
    'PRIMER_APELLIDO': 'Primer Apellido',
    'SEGUNDO_APELLIDO':'Segundo Apellido',
    'Grado_Str':       'Grado',
    'CÓDIGO_SECCIÓN':  'Código Sección LXP',
}, inplace=True)

# Sort students by school status, then school name, then grade
# Also bring in school Estatus and per-grade statuses to derive per-student grade status
students_df = students_df.merge(
    schools_df[['Código', 'Estatus'] + [f'grado_{g}' for g in GRADE_ORDER]],
    on='Código', how='left'
)
students_df['status_sort'] = students_df['Estatus'].map(
    {s: i for i, s in enumerate(STATUS_ORDER)}
)
grade_sort_map = {g: i for i, g in enumerate(GRADE_ORDER)}
students_df['grade_sort'] = students_df['Grado'].map(grade_sort_map)
students_df.sort_values(['status_sort', 'Centro Escolar', 'grade_sort'], inplace=True)

# Build per-student grade status from their grade's column
def get_grade_status(row):
    col = f'grado_{row["Grado"]}'
    return row[col] if col in row.index else ''
students_df['Estatus Grado'] = students_df.apply(get_grade_status, axis=1)
students_df.rename(columns={'Estatus': 'Estatus Centro Escolar'}, inplace=True)

# Drop helper columns
students_df.drop(columns=['status_sort', 'grade_sort'] + [f'grado_{g}' for g in GRADE_ORDER], inplace=True)
students_df.reset_index(drop=True, inplace=True)

# ─────────────────────────────────────────────
# 8. STYLING HELPERS
# ─────────────────────────────────────────────
STATUS_COLORS = {
    'Excelente': '1E7145',  # dark green
    'Bueno':     '70AD47',  # light green
    'Regular':   'FFD966',  # yellow
    'Alerta':    'FF4B4B',  # red
}
STATUS_FONT_COLORS = {
    'Excelente': 'FFFFFF',
    'Bueno':     'FFFFFF',
    'Regular':   '000000',
    'Alerta':    'FFFFFF',
}
HEADER_FILL   = PatternFill('solid', start_color='2E4057', end_color='2E4057')
HEADER_FONT   = Font(name='Arial', bold=True, color='FFFFFF', size=10)
SUBHEADER_FILL= PatternFill('solid', start_color='4A6FA5', end_color='4A6FA5')
SUBHEADER_FONT= Font(name='Arial', bold=True, color='FFFFFF', size=9)
BODY_FONT     = Font(name='Arial', size=9)
ALT_FILL      = PatternFill('solid', start_color='F2F2F2', end_color='F2F2F2')
WHITE_FILL    = PatternFill('solid', start_color='FFFFFF', end_color='FFFFFF')
CENTER_ALIGN  = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN    = Alignment(horizontal='left',   vertical='center', wrap_text=True)

THIN = Side(style='thin', color='CCCCCC')
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def set_header(cell, value):
    cell.value = value
    cell.font  = HEADER_FONT
    cell.fill  = HEADER_FILL
    cell.alignment = CENTER_ALIGN
    cell.border = THIN_BORDER

def set_subheader(cell, value):
    cell.value = value
    cell.font  = SUBHEADER_FONT
    cell.fill  = SUBHEADER_FILL
    cell.alignment = CENTER_ALIGN
    cell.border = THIN_BORDER

def set_body(cell, value, align='left', row_idx=0):
    cell.value = value
    cell.font  = BODY_FONT
    cell.fill  = ALT_FILL if row_idx % 2 == 0 else WHITE_FILL
    cell.alignment = CENTER_ALIGN if align == 'center' else LEFT_ALIGN
    cell.border = THIN_BORDER

def status_cell(cell, value, row_idx=0):
    cell.value = value if value else ''
    if value and value in STATUS_COLORS:
        cell.fill = PatternFill('solid',
                                start_color=STATUS_COLORS[value],
                                end_color=STATUS_COLORS[value])
        cell.font = Font(name='Arial', size=9, bold=True,
                         color=STATUS_FONT_COLORS[value])
    else:
        cell.font = BODY_FONT
        cell.fill = ALT_FILL if row_idx % 2 == 0 else WHITE_FILL
    cell.alignment = CENTER_ALIGN
    cell.border = THIN_BORDER

# ─────────────────────────────────────────────
# 9. CREATE WORKBOOK
# ─────────────────────────────────────────────
wb = Workbook()

# ── SHEET 1: SCHOOLS ──────────────────────────
ws1 = wb.active
ws1.title = 'Centros Escolares'
ws1.freeze_panes = 'A3'

# Build column layout for schools sheet
# Fixed cols: Código, Centro Escolar, Tipo, Estatus, Mat Wins, Len Wins
# Then one column per grade
fixed_headers_r1 = ['Código', 'Centro Escolar', 'Tipo de Centro', 'Estatus General',
                     'Mat: veces ≤ universo (de 3)',
                     'Len: veces ≤ universo (de 3)']
grade_col_labels = [GRADE_LABELS[g] for g in GRADE_ORDER]

# Row 1: main headers + "Estatus por Grado" merged over grade columns
# Row 2: sub-labels for grade columns

n_fixed = len(fixed_headers_r1)
n_grades = len(GRADE_ORDER)
total_cols = n_fixed + n_grades

# Row 1
for ci, h in enumerate(fixed_headers_r1, start=1):
    set_header(ws1.cell(1, ci), h)
    ws1.merge_cells(start_row=1, start_column=ci, end_row=2, end_column=ci)

# "Estatus por Grado" spanning all grade columns
grade_start_col = n_fixed + 1
grade_end_col   = n_fixed + n_grades
ws1.merge_cells(start_row=1, start_column=grade_start_col,
                end_row=1,   end_column=grade_end_col)
set_header(ws1.cell(1, grade_start_col), 'Estatus por Grado')

# Row 2: individual grade labels
for gi, label in enumerate(grade_col_labels, start=grade_start_col):
    set_subheader(ws1.cell(2, gi), label)

# Row 1 merged cells also need border/fill fixed
for ci in range(1, n_fixed + 1):
    ws1.cell(2, ci).fill  = HEADER_FILL
    ws1.cell(2, ci).border = THIN_BORDER

# Data rows
for ri, row in schools_df.iterrows():
    excel_row = ri + 3  # 1-indexed, skip 2 header rows
    alt = ri

    set_body(ws1.cell(excel_row, 1), row['Código'],         'center', alt)
    set_body(ws1.cell(excel_row, 2), row['Centro Escolar'], 'left',   alt)
    set_body(ws1.cell(excel_row, 3), row['Tipo'],           'center', alt)
    status_cell(ws1.cell(excel_row, 4), row['Estatus'],     alt)
    set_body(ws1.cell(excel_row, 5), f"{row['mat_wins']}/3", 'center', alt)
    set_body(ws1.cell(excel_row, 6), f"{row['len_wins']}/3", 'center', alt)

    for gi, grade in enumerate(GRADE_ORDER, start=grade_start_col):
        status_cell(ws1.cell(excel_row, gi), row[f'grado_{grade}'], alt)

# Column widths
ws1.column_dimensions['A'].width = 10
ws1.column_dimensions['B'].width = 52
ws1.column_dimensions['C'].width = 20
ws1.column_dimensions['D'].width = 14
ws1.column_dimensions['E'].width = 14
ws1.column_dimensions['F'].width = 14
for gi in range(grade_start_col, grade_start_col + n_grades):
    ws1.column_dimensions[get_column_letter(gi)].width = 14

ws1.row_dimensions[1].height = 28
ws1.row_dimensions[2].height = 28

# ── SHEET 2: STUDENTS ─────────────────────────
ws2 = wb.create_sheet('Estudiantes')
ws2.freeze_panes = 'A3'

# Two header rows:
# Row 1: fixed cols merged + "Matemática" merged (3 pairs) + "Lenguaje" merged (3 pairs)
# Row 2: sub-labels (Mes 1 Puntaje/Nivel, etc.)

fixed_stu = ['Código', 'Centro Escolar', 'Tipo de Centro', 'NIE',
             'Primer Nombre', 'Segundo Nombre', 'Primer Apellido', 'Segundo Apellido',
             'Grado', 'Estatus Centro Escolar', 'Estatus Grado', 'Código Sección LXP']
n_fixed_stu = len(fixed_stu)

# Each subject has 3 months × 2 cols (Puntaje, Nivel) = 6 cols each
# Layout: fixed(12) | MAT(6) | LEN(6)
mat_start = n_fixed_stu + 1
mat_end   = n_fixed_stu + 6
len_start = n_fixed_stu + 7
len_end   = n_fixed_stu + 12

# Row 1
for ci, h in enumerate(fixed_stu, start=1):
    set_header(ws2.cell(1, ci), h)
    ws2.merge_cells(start_row=1, start_column=ci, end_row=2, end_column=ci)
    ws2.cell(2, ci).fill   = HEADER_FILL
    ws2.cell(2, ci).border = THIN_BORDER

ws2.merge_cells(start_row=1, start_column=mat_start, end_row=1, end_column=mat_end)
set_header(ws2.cell(1, mat_start), 'Matemática')

ws2.merge_cells(start_row=1, start_column=len_start, end_row=1, end_column=len_end)
set_header(ws2.cell(1, len_start), 'Lenguaje')

# Row 2: Puntaje / Nivel labels for each month
month_labels = ['Puntaje M1', 'Nivel M1', 'Puntaje M2', 'Nivel M2', 'Puntaje M3', 'Nivel M3']
for offset, label in enumerate(month_labels):
    set_subheader(ws2.cell(2, mat_start + offset), label)
    set_subheader(ws2.cell(2, len_start + offset), label)

# Score/level column names in the dataframe
stu_score_cols = [
    'Puntaje Matemática (Mes 1 (Marzo))', 'Nivel Matemática (Mes 1 (Marzo))',
    'Puntaje Matemática (Mes 2 (Abril))', 'Nivel Matemática (Mes 2 (Abril))',
    'Puntaje Matemática (Mes 3 (Mayo))',  'Nivel Matemática (Mes 3 (Mayo))',
    'Puntaje Lengua (Mes 1 (Marzo))',     'Nivel Lengua (Mes 1 (Marzo))',
    'Puntaje Lengua (Mes 2 (Abril))',     'Nivel Lengua (Mes 2 (Abril))',
    'Puntaje Lengua (Mes 3 (Mayo))',      'Nivel Lengua (Mes 3 (Mayo))',
]

NIVEL_COLORS = {
    'Crítico':   'FF4B4B',
    'Bajo':      'FF9966',
    'Medio':     'FFD966',
    'Bueno':     'A9D18E',
    'Excelente': '1E7145',
}
NIVEL_FONT_COLORS = {
    'Crítico':   'FFFFFF',
    'Bajo':      'FFFFFF',
    'Medio':     '000000',
    'Bueno':     '000000',
    'Excelente': 'FFFFFF',
}

for ri, row in students_df.iterrows():
    excel_row = ri + 3
    alt = ri
    set_body(ws2.cell(excel_row, 1), row['Código'],          'center', alt)
    set_body(ws2.cell(excel_row, 2), row['Centro Escolar'],  'left',   alt)
    set_body(ws2.cell(excel_row, 3), row['Tipo'],            'center', alt)
    set_body(ws2.cell(excel_row, 4), row['NIE'],             'center', alt)
    set_body(ws2.cell(excel_row, 5), row['Primer Nombre'],   'left',   alt)
    set_body(ws2.cell(excel_row, 6), row['Segundo Nombre'],  'left',   alt)
    set_body(ws2.cell(excel_row, 7), row['Primer Apellido'], 'left',   alt)
    set_body(ws2.cell(excel_row, 8), row['Segundo Apellido'],'left',   alt)
    set_body(ws2.cell(excel_row, 9),  row['Grado'],                    'center', alt)
    status_cell(ws2.cell(excel_row, 10), row['Estatus Centro Escolar'], alt)
    status_cell(ws2.cell(excel_row, 11), row['Estatus Grado'],          alt)
    set_body(ws2.cell(excel_row, 12), row['Código Sección LXP'],        'center', alt)

    for offset, col_name in enumerate(stu_score_cols):
        ci    = n_fixed_stu + 1 + offset
        val   = row[col_name]
        cell  = ws2.cell(excel_row, ci)
        is_nivel = 'Nivel' in col_name

        if is_nivel and pd.notna(val) and val in NIVEL_COLORS:
            cell.value = val
            cell.fill  = PatternFill('solid',
                                     start_color=NIVEL_COLORS[val],
                                     end_color=NIVEL_COLORS[val])
            cell.font  = Font(name='Arial', size=9,
                              color=NIVEL_FONT_COLORS[val])
            cell.alignment = CENTER_ALIGN
            cell.border    = THIN_BORDER
        else:
            display_val = round(val, 2) if pd.notna(val) and not is_nivel else (val if pd.notna(val) else '')
            set_body(cell, display_val, 'center', alt)

# Column widths for students sheet
ws2.column_dimensions['A'].width = 10
ws2.column_dimensions['B'].width = 50
ws2.column_dimensions['C'].width = 20
ws2.column_dimensions['D'].width = 12
ws2.column_dimensions['E'].width = 16
ws2.column_dimensions['F'].width = 16
ws2.column_dimensions['G'].width = 16
ws2.column_dimensions['H'].width = 16
ws2.column_dimensions['I'].width = 10
ws2.column_dimensions['J'].width = 20
ws2.column_dimensions['K'].width = 16
ws2.column_dimensions['L'].width = 18
for ci in range(n_fixed_stu + 1, n_fixed_stu + 13):
    ws2.column_dimensions[get_column_letter(ci)].width = 13

ws2.row_dimensions[1].height = 28
ws2.row_dimensions[2].height = 28


# ── SHEET 3: VERIFICATION ─────────────────────
ws3 = wb.create_sheet('Verificación de Cálculo')
ws3.freeze_panes = 'A2'

verif_headers = [
    'Código', 'Centro Escolar', 'Tipo de Centro', 'Estatus',
    'Mat: veces ≤ universo', 'Len: veces ≤ universo',
    'Asignatura', 'Mes',
    '% Crítico Escuela', '% Crítico Universo (Ref.)',
    'Escuela ≤ Universo'
]

# Single header row
for ci, h in enumerate(verif_headers, start=1):
    set_header(ws3.cell(1, ci), h)

WIN_YES_FILL  = PatternFill('solid', start_color='C6EFCE', end_color='C6EFCE')
WIN_YES_FONT  = Font(name='Arial', size=9, bold=True, color='276221')
WIN_NO_FILL   = PatternFill('solid', start_color='FFC7CE', end_color='FFC7CE')
WIN_NO_FONT   = Font(name='Arial', size=9, bold=True, color='9C0006')
WIN_NA_FILL   = PatternFill('solid', start_color='EEEEEE', end_color='EEEEEE')
WIN_NA_FONT   = Font(name='Arial', size=9, color='888888')

for ri, row in verif_df.iterrows():
    er  = ri + 2
    alt = ri
    set_body(ws3.cell(er,  1), row['Código'],              'center', alt)
    set_body(ws3.cell(er,  2), row['Centro Escolar'],      'left',   alt)
    set_body(ws3.cell(er,  3), row['Tipo'],                'center', alt)
    status_cell(ws3.cell(er, 4), row['Estatus'],           alt)
    set_body(ws3.cell(er,  5), row['Mat Wins'],            'center', alt)
    set_body(ws3.cell(er,  6), row['Len Wins'],            'center', alt)
    set_body(ws3.cell(er,  7), row['Asignatura'],          'center', alt)
    set_body(ws3.cell(er,  8), row['Mes'],                 'center', alt)

    # % Crítico Escuela
    esc_val = row['% Crítico Escuela']
    set_body(ws3.cell(er, 9), f"{esc_val}%" if esc_val != '' else '', 'center', alt)

    # % Crítico Universo
    univ_val = row['% Crítico Universo (Ref.)']
    set_body(ws3.cell(er, 10), f"{univ_val}%" if univ_val != '' else '', 'center', alt)

    # Win flag with dedicated color
    win_cell = ws3.cell(er, 11)
    win_val  = row['Escuela ≤ Universo']
    win_cell.value     = win_val
    win_cell.alignment = CENTER_ALIGN
    win_cell.border    = THIN_BORDER
    if win_val == '✔ Sí':
        win_cell.fill = WIN_YES_FILL
        win_cell.font = WIN_YES_FONT
    elif win_val == '✘ No':
        win_cell.fill = WIN_NO_FILL
        win_cell.font = WIN_NO_FONT
    else:
        win_cell.fill = WIN_NA_FILL
        win_cell.font = WIN_NA_FONT

# Column widths
ws3.column_dimensions['A'].width = 10
ws3.column_dimensions['B'].width = 52
ws3.column_dimensions['C'].width = 20
ws3.column_dimensions['D'].width = 14
ws3.column_dimensions['E'].width = 22
ws3.column_dimensions['F'].width = 22
ws3.column_dimensions['G'].width = 14
ws3.column_dimensions['H'].width = 18
ws3.column_dimensions['I'].width = 20
ws3.column_dimensions['J'].width = 24
ws3.column_dimensions['K'].width = 18
ws3.row_dimensions[1].height = 28

# ─────────────────────────────────────────────
# 10. SAVE
# ─────────────────────────────────────────────
wb.save(OUTPUT_FILE)
print(f'Saved: {OUTPUT_FILE}')
print(f'Schools sheet: {len(schools_df)} schools')
print(f'Students sheet: {len(students_df)} students')
print(f'Verification sheet: {len(verif_df)} rows ({len(verif_df)//6} schools x 6 month-subject combinations)')
print()
print('Status distribution:')
print(schools_df['Estatus'].value_counts())
print()
print('Type distribution:')
print(schools_df['Tipo'].value_counts())
print()
print('Schools with 0 wins in BOTH Math AND Language (0/3 and 0/3):')
zero_both = schools_df[(schools_df['mat_wins'] == 0) & (schools_df['len_wins'] == 0)]
print(f'  Total: {len(zero_both)}')
print(zero_both.groupby('Tipo').size().rename('count').to_string())
print()
print('  Detail:')
for _, r in zero_both[['Código','Centro Escolar','Tipo','Estatus']].iterrows():
    print(f'    [{r["Tipo"]}] {r["Código"]} - {r["Centro Escolar"]}')