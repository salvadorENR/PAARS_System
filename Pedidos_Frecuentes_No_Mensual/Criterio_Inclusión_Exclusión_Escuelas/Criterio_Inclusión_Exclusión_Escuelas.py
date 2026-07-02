import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# 1. LOAD & FILTER DATA
# ─────────────────────────────────────────────
INPUT_FILE   = r'C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Excel_Maestro_HojaUnica.csv'
OUTPUT_B1    = r'C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Centros_Escolares_B1.xlsx'
OUTPUT_B2    = r'C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Centros_Escolares_B2.xlsx'

df = pd.read_csv(INPUT_FILE, sep=';')

# ── Determine the definitive group for each school ──────────────────
# The school's group is the MAJORITY value of GRUPO_Limpio across all
# its student rows. Any stray rows with a different value are data
# entry errors and are corrected to the school's definitive group.
school_group = (
    df.groupby('CODIGO')['GRUPO_Limpio']
    .agg(lambda x: x.value_counts().idxmax())   # majority vote
    .reset_index()
    .rename(columns={'GRUPO_Limpio': 'GRUPO_Escuela'})
)

# Merge the definitive school group back onto every student row
df = df.merge(school_group, on='CODIGO', how='left')

# Separate B1 (3 applications) and B2 (1 application)
valid_b1 = df[df['GRUPO_Escuela'] == 'B1'].copy()
valid_b2 = df[df['GRUPO_Escuela'] == 'B2'].copy()

# ─────────────────────────────────────────────
# 2. SCHOOL TYPE CLASSIFICATION (mutually exclusive)
# ─────────────────────────────────────────────
def classify_school(name):
    n = str(name).upper()
    if 'INSTITUTO' in n:
        return 'Instituto'
    elif 'COMPLEJO EDUCATIVO' in n:
        return 'Complejo Educativo'
    else:
        return 'Centro Escolar'

valid_b1['Tipo_Escuela'] = valid_b1['NOMBRE'].apply(classify_school)
valid_b2['Tipo_Escuela'] = valid_b2['NOMBRE'].apply(classify_school)

# ─────────────────────────────────────────────
# 3. CONSTANTS
# ─────────────────────────────────────────────
# B1: 3 months of nivel/score columns
NIVEL_COLS_B1 = {
    'mat': [
        'Nivel Matemática (Mes 1 (Marzo))',
        'Nivel Matemática (Mes 2 (Abril))',
        'Nivel Matemática (Mes 3 (Mayo))',
    ],
    'len': [
        'Nivel Lengua (Mes 1 (Marzo))',
        'Nivel Lengua (Mes 2 (Abril))',
        'Nivel Lengua (Mes 3 (Mayo))',
    ]
}

SCORE_COLS_B1 = {
    'mat': [
        'Puntaje Matemática (Mes 1 (Marzo))',
        'Puntaje Matemática (Mes 2 (Abril))',
        'Puntaje Matemática (Mes 3 (Mayo))',
    ],
    'len': [
        'Puntaje Lengua (Mes 1 (Marzo))',
        'Puntaje Lengua (Mes 2 (Abril))',
        'Puntaje Lengua (Mes 3 (Mayo))',
    ]
}

# B2: only the most recent application (Mes 3 Mayo)
# Adjust these column names to whichever month B2 uses
NIVEL_COLS_B2 = {
    'mat': ['Nivel Matemática (Mes 3 (Mayo))'],
    'len': ['Nivel Lengua (Mes 3 (Mayo))'],
}

SCORE_COLS_B2 = {
    'mat': ['Puntaje Matemática (Mes 3 (Mayo))'],
    'len': ['Puntaje Lengua (Mes 3 (Mayo))'],
}

GRADE_ORDER = ['2º', '3º', '4º', '5º', '6º', '7°', '8°', '9°', '10°', '11°']
GRADE_LABELS = {
    '2º':  '2° Grado',
    '3º':  '3° Grado',
    '4º':  '4° Grado',
    '5º':  '5° Grado',
    '6º':  '6° Grado',
    '7°':  '7° Grado',
    '8°':  '8° Grado',
    '9°':  '9° Grado',
    '10°': '10° (1er Año Bach.)',
    '11°': '11° (2do Año Bach.)',
}

STATUS_ORDER = ['Alerta', 'Regular', 'Bueno', 'Excelente']

# ─────────────────────────────────────────────
# 4. CRITICAL % CALCULATIONS
# ─────────────────────────────────────────────
def pct_critico(series):
    """Percentage of 'Crítico' out of non-null values."""
    valid = series.dropna()
    if len(valid) == 0:
        return np.nan
    return (valid == 'Crítico').sum() / len(valid)


def compute_status_b1(mat_wins, len_wins):
    """
    B1 (Grupo 1): 3 applications → wins range 0-3 per subject.
    Returns: 'Excelente', 'Bueno', 'Regular', 'Alerta'
    """
    if mat_wins >= 2 and len_wins >= 2:
        return 'Excelente'
    elif (mat_wins == 3 and len_wins == 1) or (mat_wins == 1 and len_wins == 3):
        return 'Bueno'
    elif (mat_wins == 2 and len_wins == 1) or (mat_wins == 1 and len_wins == 2):
        return 'Regular'
    else:
        return 'Alerta'


def compute_status_b2(mat_wins, len_wins):
    """
    B2 (Grupo 2): 1 application → wins are 0 or 1 per subject.
    (1,1) → Excelente  — beats universe in BOTH subjects
    (1,0) or (0,1) → Regular  — beats universe in ONE subject
    (0,0) → Alerta  — does not beat universe in either subject
    Note: 'Bueno' is not applicable with a single application.
    """
    if mat_wins == 1 and len_wins == 1:
        return 'Excelente'
    elif mat_wins == 1 or len_wins == 1:
        return 'Regular'
    else:
        return 'Alerta'


# ─────────────────────────────────────────────
# 5. BUILD UNIVERSE % CRÍTICO
# ─────────────────────────────────────────────
def build_universe(data, tipo, nivel_cols):
    """Returns {subject: [pct_m1, pct_m2, ...]} for a given school type."""
    subset = data[data['Tipo_Escuela'] == tipo]
    universe = {}
    for subj, cols in nivel_cols.items():
        universe[subj] = [pct_critico(subset[c]) for c in cols]
    return universe

# B1 universes (split by school type, 3 months)
universes_b1 = {
    tipo: build_universe(valid_b1, tipo, NIVEL_COLS_B1)
    for tipo in ['Instituto', 'Complejo Educativo', 'Centro Escolar']
}

# B2 universes (split by school type, 1 month)
universes_b2 = {
    tipo: build_universe(valid_b2, tipo, NIVEL_COLS_B2)
    for tipo in ['Instituto', 'Complejo Educativo', 'Centro Escolar']
}

# ─────────────────────────────────────────────
# 6. COMPUTE SCHOOL-LEVEL STATUS — B1
# ─────────────────────────────────────────────
def count_participations(grp, nivel_cols):
    """
    Count how many months (applications) a school actually had data for.
    A month is considered 'participated' if at least one student has a
    non-null value in any of that month's nivel columns.
    B1 has 3 months so result is 0, 1, 2, or 3.
    """
    participated = 0
    # Zip together mat and len columns by month position
    mat_cols = nivel_cols['mat']
    len_cols = nivel_cols['len']
    for mat_col, len_col in zip(mat_cols, len_cols):
        has_mat = grp[mat_col].notna().any()
        has_len = grp[len_col].notna().any()
        if has_mat or has_len:
            participated += 1
    return participated


def build_school_records(valid_data, nivel_cols, universes, compute_status_fn,
                         include_participations=False):
    """
    Generic school record builder usable for both B1 and B2.
    Returns a list of dicts ready for pd.DataFrame().
    """
    records = []
    for (codigo, nombre), grp in valid_data.groupby(['CODIGO', 'NOMBRE']):
        tipo = grp['Tipo_Escuela'].iloc[0]
        univ = universes[tipo]

        # Count wins per subject
        wins = {}
        for subj, cols in nivel_cols.items():
            w = 0
            for i, col in enumerate(cols):
                school_pct = pct_critico(grp[col])
                univ_pct   = univ[subj][i]
                if not np.isnan(school_pct) and not np.isnan(univ_pct):
                    if school_pct <= univ_pct:
                        w += 1
            wins[subj] = w

        status = compute_status_fn(wins['mat'], wins['len'])

        # Grade-level status
        grade_statuses = {}
        for grade in GRADE_ORDER:
            grade_grp = grp[grp['Grado_Str'] == grade]
            if len(grade_grp) == 0:
                grade_statuses[grade] = ''
                continue
            grade_univ_data = valid_data[
                (valid_data['Tipo_Escuela'] == tipo) &
                (valid_data['Grado_Str'] == grade)
            ]
            g_wins = {}
            for subj, cols in nivel_cols.items():
                w = 0
                for col in cols:
                    s_pct = pct_critico(grade_grp[col])
                    u_pct = pct_critico(grade_univ_data[col])
                    if not np.isnan(s_pct) and not np.isnan(u_pct):
                        if s_pct <= u_pct:
                            w += 1
                g_wins[subj] = w
            grade_statuses[grade] = compute_status_fn(g_wins['mat'], g_wins['len'])

        record = {
            'Código':          codigo,
            'Centro Escolar':  nombre,
            'Tipo':            tipo,
            'Estatus':         status,
            'mat_wins':        wins['mat'],
            'len_wins':        wins['len'],
            **{f'grado_{g}': grade_statuses[g] for g in GRADE_ORDER}
        }

        # B1 only: add participation count
        if include_participations:
            record['N° Aplicaciones'] = count_participations(grp, nivel_cols)

        records.append(record)

    df_out = pd.DataFrame(records)
    df_out['status_sort'] = df_out['Estatus'].map(
        {s: i for i, s in enumerate(STATUS_ORDER)}
    )
    df_out.sort_values('status_sort', inplace=True)
    df_out.drop(columns='status_sort', inplace=True)
    df_out.reset_index(drop=True, inplace=True)
    return df_out


schools_b1 = build_school_records(
    valid_b1, NIVEL_COLS_B1, universes_b1,
    compute_status_b1, include_participations=True
)

schools_b2 = build_school_records(
    valid_b2, NIVEL_COLS_B2, universes_b2,
    compute_status_b2, include_participations=False
)

# ─────────────────────────────────────────────
# 6b. BUILD VERIFICATION DATA (generic)
# ─────────────────────────────────────────────
def build_verif_df(valid_data, nivel_cols, universes, schools_df,
                   month_labels, max_wins):
    """
    Build the verification DataFrame for either B1 or B2.
    max_wins: string like '3' or '1' used in the wins display column.
    """
    SUBJ_LABELS = {'mat': 'Matemática', 'len': 'Lenguaje'}
    verif_records = []

    for (codigo, nombre), grp in valid_data.groupby(['CODIGO', 'NOMBRE']):
        tipo    = grp['Tipo_Escuela'].iloc[0]
        univ    = universes[tipo]
        estatus = schools_df.loc[schools_df['Código'] == codigo, 'Estatus'].values[0]
        mat_w   = schools_df.loc[schools_df['Código'] == codigo, 'mat_wins'].values[0]
        len_w   = schools_df.loc[schools_df['Código'] == codigo, 'len_wins'].values[0]

        for subj, cols in nivel_cols.items():
            for i, col in enumerate(cols):
                school_pct = pct_critico(grp[col])
                univ_pct   = univ[subj][i]
                win = (school_pct <= univ_pct) if (
                    not np.isnan(school_pct) and not np.isnan(univ_pct)
                ) else None
                verif_records.append({
                    'Código':                    codigo,
                    'Centro Escolar':            nombre,
                    'Tipo':                      tipo,
                    'Estatus':                   estatus,
                    'Mat Wins':                  f"{mat_w}/{max_wins}",
                    'Len Wins':                  f"{len_w}/{max_wins}",
                    'Asignatura':                SUBJ_LABELS[subj],
                    'Mes':                       month_labels[i],
                    '% Crítico Escuela':         round(school_pct * 100, 2) if not np.isnan(school_pct) else '',
                    '% Crítico Universo (Ref.)': round(univ_pct   * 100, 2) if not np.isnan(univ_pct)   else '',
                    'Escuela ≤ Universo':        ('✔ Sí' if win else ('✘ No' if win is not None else 'N/D')),
                })

    verif_df = pd.DataFrame(verif_records)
    verif_df['status_sort'] = verif_df['Estatus'].map(
        {s: i for i, s in enumerate(STATUS_ORDER)}
    )
    subj_sort = {'Matemática': 0, 'Lenguaje': 1}
    mes_sort  = {m: i for i, m in enumerate(month_labels)}
    verif_df['subj_sort'] = verif_df['Asignatura'].map(subj_sort)
    verif_df['mes_sort']  = verif_df['Mes'].map(mes_sort)
    verif_df.sort_values(
        ['status_sort', 'Centro Escolar', 'subj_sort', 'mes_sort'], inplace=True
    )
    verif_df.drop(columns=['status_sort', 'subj_sort', 'mes_sort'], inplace=True)
    verif_df.reset_index(drop=True, inplace=True)
    return verif_df


MONTH_LABELS_B1 = ['Mes 1 (Marzo)', 'Mes 2 (Abril)', 'Mes 3 (Mayo)']
MONTH_LABELS_B2 = ['Mes 3 (Mayo)']

verif_b1 = build_verif_df(
    valid_b1, NIVEL_COLS_B1, universes_b1, schools_b1, MONTH_LABELS_B1, max_wins='3'
)
verif_b2 = build_verif_df(
    valid_b2, NIVEL_COLS_B2, universes_b2, schools_b2, MONTH_LABELS_B2, max_wins='1'
)

# ─────────────────────────────────────────────
# 7. BUILD STUDENTS SHEET DATA (generic)
# ─────────────────────────────────────────────
def build_students_df(valid_data, schools_df, score_cols_dict):
    """Build the students DataFrame for either B1 or B2."""
    # Flatten score columns in order: mat cols interleaved with nivel cols
    # We derive nivel from score by using the matching NIVEL column name
    all_score_cols = []
    for m_col in score_cols_dict['mat']:
        n_col = m_col.replace('Puntaje', 'Nivel')
        all_score_cols.extend([m_col, n_col])
    for l_col in score_cols_dict['len']:
        n_col = l_col.replace('Puntaje', 'Nivel')
        all_score_cols.extend([l_col, n_col])

    # Keep only columns that actually exist in the data
    all_score_cols = [c for c in all_score_cols if c in valid_data.columns]

    base_cols = [
        'CODIGO', 'NOMBRE', 'Tipo_Escuela',
        'NIE', 'PRIMER_NOMBRE', 'SEGUNDO_NOMBRE',
        'PRIMER_APELLIDO', 'SEGUNDO_APELLIDO',
        'Grado_Str', 'CÓDIGO_SECCIÓN',
    ]
    students = valid_data[base_cols + all_score_cols].copy()
    students.rename(columns={
        'CODIGO':          'Código',
        'NOMBRE':          'Centro Escolar',
        'Tipo_Escuela':    'Tipo',
        'NIE':             'NIE',
        'PRIMER_NOMBRE':   'Primer Nombre',
        'SEGUNDO_NOMBRE':  'Segundo Nombre',
        'PRIMER_APELLIDO': 'Primer Apellido',
        'SEGUNDO_APELLIDO':'Segundo Apellido',
        'Grado_Str':       'Grado',
        'CÓDIGO_SECCIÓN':  'Código Sección LXP',
    }, inplace=True)

    grade_cols = [f'grado_{g}' for g in GRADE_ORDER]
    students = students.merge(
        schools_df[['Código', 'Estatus'] + grade_cols],
        on='Código', how='left'
    )
    students['status_sort'] = students['Estatus'].map(
        {s: i for i, s in enumerate(STATUS_ORDER)}
    )
    grade_sort_map = {g: i for i, g in enumerate(GRADE_ORDER)}
    students['grade_sort'] = students['Grado'].map(grade_sort_map)
    students.sort_values(['status_sort', 'Centro Escolar', 'grade_sort'], inplace=True)

    def get_grade_status(row):
        col = f'grado_{row["Grado"]}'
        return row[col] if col in row.index else ''
    students['Estatus Grado'] = students.apply(get_grade_status, axis=1)
    students.rename(columns={'Estatus': 'Estatus Centro Escolar'}, inplace=True)
    students.drop(
        columns=['status_sort', 'grade_sort'] + grade_cols,
        inplace=True
    )
    students.reset_index(drop=True, inplace=True)
    return students, all_score_cols


students_b1, score_cols_b1 = build_students_df(valid_b1, schools_b1, SCORE_COLS_B1)
students_b2, score_cols_b2 = build_students_df(valid_b2, schools_b2, SCORE_COLS_B2)

# ─────────────────────────────────────────────
# 8. STYLING HELPERS
# ─────────────────────────────────────────────
STATUS_COLORS = {
    'Excelente': '1E7145',
    'Bueno':     '70AD47',
    'Regular':   'FFD966',
    'Alerta':    'FF4B4B',
}
STATUS_FONT_COLORS = {
    'Excelente': 'FFFFFF',
    'Bueno':     'FFFFFF',
    'Regular':   '000000',
    'Alerta':    'FFFFFF',
}
HEADER_FILL    = PatternFill('solid', start_color='2E4057', end_color='2E4057')
HEADER_FONT    = Font(name='Arial', bold=True, color='FFFFFF', size=10)
SUBHEADER_FILL = PatternFill('solid', start_color='4A6FA5', end_color='4A6FA5')
SUBHEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=9)
BODY_FONT      = Font(name='Arial', size=9)
ALT_FILL       = PatternFill('solid', start_color='F2F2F2', end_color='F2F2F2')
WHITE_FILL     = PatternFill('solid', start_color='FFFFFF', end_color='FFFFFF')
CENTER_ALIGN   = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN     = Alignment(horizontal='left',   vertical='center', wrap_text=True)
THIN           = Side(style='thin', color='CCCCCC')
THIN_BORDER    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NIVEL_COLORS = {
    'Crítico':   'FF4B4B',
    'Bajo':      'FF9966',
    'Medio':     'FFD966',
    'Bueno':     'A9D18E',
    'Excelente': '1E7145',
}
NIVEL_FONT_COLORS = {
    'Crítico':   'FFFFFF',
    'Bajo':      'FFFFFF',
    'Medio':     '000000',
    'Bueno':     '000000',
    'Excelente': 'FFFFFF',
}

WIN_YES_FILL = PatternFill('solid', start_color='C6EFCE', end_color='C6EFCE')
WIN_YES_FONT = Font(name='Arial', size=9, bold=True, color='276221')
WIN_NO_FILL  = PatternFill('solid', start_color='FFC7CE', end_color='FFC7CE')
WIN_NO_FONT  = Font(name='Arial', size=9, bold=True, color='9C0006')
WIN_NA_FILL  = PatternFill('solid', start_color='EEEEEE', end_color='EEEEEE')
WIN_NA_FONT  = Font(name='Arial', size=9, color='888888')


def set_header(cell, value):
    cell.value     = value
    cell.font      = HEADER_FONT
    cell.fill      = HEADER_FILL
    cell.alignment = CENTER_ALIGN
    cell.border    = THIN_BORDER

def set_subheader(cell, value):
    cell.value     = value
    cell.font      = SUBHEADER_FONT
    cell.fill      = SUBHEADER_FILL
    cell.alignment = CENTER_ALIGN
    cell.border    = THIN_BORDER

def set_body(cell, value, align='left', row_idx=0):
    cell.value     = value
    cell.font      = BODY_FONT
    cell.fill      = ALT_FILL if row_idx % 2 == 0 else WHITE_FILL
    cell.alignment = CENTER_ALIGN if align == 'center' else LEFT_ALIGN
    cell.border    = THIN_BORDER

def status_cell(cell, value, row_idx=0):
    cell.value = value if value else ''
    if value and value in STATUS_COLORS:
        cell.fill = PatternFill('solid',
                                start_color=STATUS_COLORS[value],
                                end_color=STATUS_COLORS[value])
        cell.font = Font(name='Arial', size=9, bold=True,
                         color=STATUS_FONT_COLORS[value])
    else:
        cell.font = BODY_FONT
        cell.fill = ALT_FILL if row_idx % 2 == 0 else WHITE_FILL
    cell.alignment = CENTER_ALIGN
    cell.border    = THIN_BORDER


# ─────────────────────────────────────────────
# 9. BUILD EXCEL WORKBOOK (generic)
# ─────────────────────────────────────────────
def build_workbook(schools_df, students_df, verif_df,
                   score_cols_list, month_labels,
                   max_wins_label, grupo_label,
                   include_participations=False):
    """
    Builds a complete workbook for either B1 or B2.

    include_participations: if True, adds the "N° Aplicaciones" column
                            to the schools sheet (B1 only).
    """
    wb = Workbook()

    # ── SHEET 1: SCHOOLS ──────────────────────
    ws1 = wb.active
    ws1.title = 'Centros Escolares'
    ws1.freeze_panes = 'A3'

    # Fixed headers — B1 adds "N° Aplicaciones" after wins columns
    fixed_headers_r1 = [
        'Código',
        'Centro Escolar',
        'Tipo de Centro',
        'Estatus General',
        f'Mat: veces ≤ universo (de {max_wins_label})',
        f'Len: veces ≤ universo (de {max_wins_label})',
    ]
    if include_participations:
        fixed_headers_r1.append('N° Aplicaciones\nParticipadas')

    grade_col_labels = [GRADE_LABELS[g] for g in GRADE_ORDER]
    n_fixed      = len(fixed_headers_r1)
    n_grades     = len(GRADE_ORDER)
    grade_start  = n_fixed + 1
    grade_end    = n_fixed + n_grades

    # Row 1: fixed headers (merged rows 1-2) + "Estatus por Grado" merged
    for ci, h in enumerate(fixed_headers_r1, start=1):
        set_header(ws1.cell(1, ci), h)
        ws1.merge_cells(start_row=1, start_column=ci,
                        end_row=2,   end_column=ci)

    ws1.merge_cells(start_row=1, start_column=grade_start,
                    end_row=1,   end_column=grade_end)
    set_header(ws1.cell(1, grade_start), 'Estatus por Grado')

    # Row 2: grade labels
    for gi, label in enumerate(grade_col_labels, start=grade_start):
        set_subheader(ws1.cell(2, gi), label)

    for ci in range(1, n_fixed + 1):
        ws1.cell(2, ci).fill   = HEADER_FILL
        ws1.cell(2, ci).border = THIN_BORDER

    # Data rows
    for ri, row in schools_df.iterrows():
        er  = ri + 3
        alt = ri
        set_body(ws1.cell(er, 1), row['Código'],         'center', alt)
        set_body(ws1.cell(er, 2), row['Centro Escolar'], 'left',   alt)
        set_body(ws1.cell(er, 3), row['Tipo'],           'center', alt)
        status_cell(ws1.cell(er, 4), row['Estatus'],     alt)
        set_body(ws1.cell(er, 5), f"{row['mat_wins']}/{max_wins_label}", 'center', alt)
        set_body(ws1.cell(er, 6), f"{row['len_wins']}/{max_wins_label}", 'center', alt)

        if include_participations:
            set_body(ws1.cell(er, 7), row['N° Aplicaciones'], 'center', alt)

        for gi, grade in enumerate(GRADE_ORDER, start=grade_start):
            status_cell(ws1.cell(er, gi), row[f'grado_{grade}'], alt)

    # Column widths
    ws1.column_dimensions['A'].width = 10
    ws1.column_dimensions['B'].width = 52
    ws1.column_dimensions['C'].width = 20
    ws1.column_dimensions['D'].width = 14
    ws1.column_dimensions['E'].width = 18
    ws1.column_dimensions['F'].width = 18
    if include_participations:
        ws1.column_dimensions['G'].width = 16
    for gi in range(grade_start, grade_start + n_grades):
        ws1.column_dimensions[get_column_letter(gi)].width = 14
    ws1.row_dimensions[1].height = 28
    ws1.row_dimensions[2].height = 28

    # ── SHEET 2: STUDENTS ─────────────────────
    ws2 = wb.create_sheet('Estudiantes')
    ws2.freeze_panes = 'A3'

    fixed_stu = [
        'Código', 'Centro Escolar', 'Tipo de Centro', 'NIE',
        'Primer Nombre', 'Segundo Nombre', 'Primer Apellido', 'Segundo Apellido',
        'Grado', 'Estatus Centro Escolar', 'Estatus Grado', 'Código Sección LXP'
    ]
    n_fixed_stu = len(fixed_stu)

    # Subject blocks: MAT (half of score_cols) then LEN (other half)
    n_score_per_subj = len(score_cols_list) // 2   # e.g. 6 for B1, 2 for B2
    mat_start = n_fixed_stu + 1
    mat_end   = n_fixed_stu + n_score_per_subj
    len_start = n_fixed_stu + n_score_per_subj + 1
    len_end   = n_fixed_stu + n_score_per_subj * 2

    # Row 1: fixed merged + subject headers
    for ci, h in enumerate(fixed_stu, start=1):
        set_header(ws2.cell(1, ci), h)
        ws2.merge_cells(start_row=1, start_column=ci,
                        end_row=2,   end_column=ci)
        ws2.cell(2, ci).fill   = HEADER_FILL
        ws2.cell(2, ci).border = THIN_BORDER

    ws2.merge_cells(start_row=1, start_column=mat_start,
                    end_row=1,   end_column=mat_end)
    set_header(ws2.cell(1, mat_start), 'Matemática')

    ws2.merge_cells(start_row=1, start_column=len_start,
                    end_row=1,   end_column=len_end)
    set_header(ws2.cell(1, len_start), 'Lenguaje')

    # Row 2: Puntaje/Nivel labels per month
    # For B1: 6 sub-cols = M1 Puntaje, M1 Nivel, M2 Puntaje, M2 Nivel, M3 P, M3 N
    # For B2: 2 sub-cols = M3 Puntaje, M3 Nivel
    month_sub_labels = []
    for ml in month_labels:
        short = ml.replace('Mes 1 (', 'M1 (').replace('Mes 2 (', 'M2 (').replace('Mes 3 (', 'M3 (')
        month_sub_labels.extend([f'Puntaje {short}', f'Nivel {short}'])

    for offset, label in enumerate(month_sub_labels):
        set_subheader(ws2.cell(2, mat_start + offset), label)
        set_subheader(ws2.cell(2, len_start + offset), label)

    # Score column names in the students DataFrame (renamed from original cols)
    stu_score_cols = [c for c in students_df.columns
                      if 'Puntaje' in c or 'Nivel' in c]

    for ri, row in students_df.iterrows():
        er  = ri + 3
        alt = ri
        set_body(ws2.cell(er,  1), row['Código'],                   'center', alt)
        set_body(ws2.cell(er,  2), row['Centro Escolar'],           'left',   alt)
        set_body(ws2.cell(er,  3), row['Tipo'],                     'center', alt)
        set_body(ws2.cell(er,  4), row['NIE'],                      'center', alt)
        set_body(ws2.cell(er,  5), row['Primer Nombre'],            'left',   alt)
        set_body(ws2.cell(er,  6), row['Segundo Nombre'],           'left',   alt)
        set_body(ws2.cell(er,  7), row['Primer Apellido'],          'left',   alt)
        set_body(ws2.cell(er,  8), row['Segundo Apellido'],         'left',   alt)
        set_body(ws2.cell(er,  9), row['Grado'],                    'center', alt)
        status_cell(ws2.cell(er, 10), row['Estatus Centro Escolar'], alt)
        status_cell(ws2.cell(er, 11), row['Estatus Grado'],          alt)
        set_body(ws2.cell(er, 12), row.get('Código Sección LXP', ''), 'center', alt)

        for offset, col_name in enumerate(stu_score_cols):
            ci   = n_fixed_stu + 1 + offset
            val  = row[col_name]
            cell = ws2.cell(er, ci)
            is_nivel = 'Nivel' in col_name

            if is_nivel and pd.notna(val) and val in NIVEL_COLORS:
                cell.value     = val
                cell.fill      = PatternFill('solid',
                                             start_color=NIVEL_COLORS[val],
                                             end_color=NIVEL_COLORS[val])
                cell.font      = Font(name='Arial', size=9,
                                      color=NIVEL_FONT_COLORS[val])
                cell.alignment = CENTER_ALIGN
                cell.border    = THIN_BORDER
            else:
                display = round(val, 2) if pd.notna(val) and not is_nivel else (
                          val if pd.notna(val) else '')
                set_body(cell, display, 'center', alt)

    # Column widths — students
    ws2.column_dimensions['A'].width = 10
    ws2.column_dimensions['B'].width = 50
    ws2.column_dimensions['C'].width = 20
    ws2.column_dimensions['D'].width = 12
    ws2.column_dimensions['E'].width = 16
    ws2.column_dimensions['F'].width = 16
    ws2.column_dimensions['G'].width = 16
    ws2.column_dimensions['H'].width = 16
    ws2.column_dimensions['I'].width = 10
    ws2.column_dimensions['J'].width = 20
    ws2.column_dimensions['K'].width = 16
    ws2.column_dimensions['L'].width = 18
    for ci in range(n_fixed_stu + 1, len_end + 1):
        ws2.column_dimensions[get_column_letter(ci)].width = 13
    ws2.row_dimensions[1].height = 28
    ws2.row_dimensions[2].height = 28

    # ── SHEET 3: VERIFICATION ─────────────────
    ws3 = wb.create_sheet('Verificación de Cálculo')
    ws3.freeze_panes = 'A2'

    verif_headers = [
        'Código', 'Centro Escolar', 'Tipo de Centro', 'Estatus',
        f'Mat: veces ≤ universo (/{max_wins_label})',
        f'Len: veces ≤ universo (/{max_wins_label})',
        'Asignatura', 'Mes',
        '% Crítico Escuela', '% Crítico Universo (Ref.)',
        'Escuela ≤ Universo'
    ]
    for ci, h in enumerate(verif_headers, start=1):
        set_header(ws3.cell(1, ci), h)

    for ri, row in verif_df.iterrows():
        er  = ri + 2
        alt = ri
        set_body(ws3.cell(er,  1), row['Código'],          'center', alt)
        set_body(ws3.cell(er,  2), row['Centro Escolar'],  'left',   alt)
        set_body(ws3.cell(er,  3), row['Tipo'],            'center', alt)
        status_cell(ws3.cell(er, 4), row['Estatus'],       alt)
        set_body(ws3.cell(er,  5), row['Mat Wins'],        'center', alt)
        set_body(ws3.cell(er,  6), row['Len Wins'],        'center', alt)
        set_body(ws3.cell(er,  7), row['Asignatura'],      'center', alt)
        set_body(ws3.cell(er,  8), row['Mes'],             'center', alt)

        esc_val  = row['% Crítico Escuela']
        univ_val = row['% Crítico Universo (Ref.)']
        set_body(ws3.cell(er,  9), f"{esc_val}%"  if esc_val  != '' else '', 'center', alt)
        set_body(ws3.cell(er, 10), f"{univ_val}%" if univ_val != '' else '', 'center', alt)

        win_cell       = ws3.cell(er, 11)
        win_val        = row['Escuela ≤ Universo']
        win_cell.value = win_val
        win_cell.alignment = CENTER_ALIGN
        win_cell.border    = THIN_BORDER
        if win_val == '✔ Sí':
            win_cell.fill = WIN_YES_FILL
            win_cell.font = WIN_YES_FONT
        elif win_val == '✘ No':
            win_cell.fill = WIN_NO_FILL
            win_cell.font = WIN_NO_FONT
        else:
            win_cell.fill = WIN_NA_FILL
            win_cell.font = WIN_NA_FONT

    ws3.column_dimensions['A'].width = 10
    ws3.column_dimensions['B'].width = 52
    ws3.column_dimensions['C'].width = 20
    ws3.column_dimensions['D'].width = 14
    ws3.column_dimensions['E'].width = 24
    ws3.column_dimensions['F'].width = 24
    ws3.column_dimensions['G'].width = 14
    ws3.column_dimensions['H'].width = 18
    ws3.column_dimensions['I'].width = 20
    ws3.column_dimensions['J'].width = 24
    ws3.column_dimensions['K'].width = 18
    ws3.row_dimensions[1].height = 28

    return wb


# ─────────────────────────────────────────────
# 10. SAVE BOTH WORKBOOKS
# ─────────────────────────────────────────────
wb_b1 = build_workbook(
    schools_df=schools_b1,
    students_df=students_b1,
    verif_df=verif_b1,
    score_cols_list=score_cols_b1,
    month_labels=MONTH_LABELS_B1,
    max_wins_label='3',
    grupo_label='B1',
    include_participations=True,       # ← B1: show N° Aplicaciones
)
wb_b1.save(OUTPUT_B1)

wb_b2 = build_workbook(
    schools_df=schools_b2,
    students_df=students_b2,
    verif_df=verif_b2,
    score_cols_list=score_cols_b2,
    month_labels=MONTH_LABELS_B2,
    max_wins_label='1',
    grupo_label='B2',
    include_participations=False,      # ← B2: single application, no count needed
)
wb_b2.save(OUTPUT_B2)

# ─────────────────────────────────────────────
# 11. CONSOLE SUMMARY
# ─────────────────────────────────────────────
for label, schools_df, verif_df, students_df in [
    ('B1 (3 aplicaciones)', schools_b1, verif_b1, students_b1),
    ('B2 (1 aplicación)',   schools_b2, verif_b2, students_b2),
]:
    n_months = 3 if 'B1' in label else 1
    print(f'\n{"="*60}')
    print(f'  GRUPO {label}')
    print(f'{"="*60}')
    print(f'  Escuelas: {len(schools_df)} | Estudiantes: {len(students_df)}')
    print(f'  Verificación: {len(verif_df)} filas '
          f'({len(verif_df)//(2*n_months)} escuelas × '
          f'{2*n_months} combinaciones asignatura-mes)')
    print(f'\n  Distribución de estatus:')
    print(schools_df['Estatus'].value_counts().to_string())
    print(f'\n  Por tipo de centro:')
    print(schools_df['Tipo'].value_counts().to_string())

    zero_both = schools_df[(schools_df['mat_wins'] == 0) & (schools_df['len_wins'] == 0)]
    print(f'\n  Escuelas con 0 wins en ambas asignaturas: {len(zero_both)}')
    for _, r in zero_both[['Código', 'Centro Escolar', 'Tipo', 'Estatus']].iterrows():
        print(f'    [{r["Tipo"]}] {r["Código"]} - {r["Centro Escolar"]}')

    if 'B1' in label:
        print(f'\n  N° Aplicaciones participadas (distribución):')
        print(schools_b1['N° Aplicaciones'].value_counts().sort_index().to_string())

print(f'\n{"="*60}')
print(f'  Archivos guardados:')
print(f'    {OUTPUT_B1}')
print(f'    {OUTPUT_B2}')
print(f'{"="*60}\n')