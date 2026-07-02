"""
agregar_fundamentos_ce.py
--------------------------
Adds per-grade Fundamentos columns to the Centros Escolares sheet
of both B1 and B2 files, reading data from Reporte_fundamentos_Final.xlsx.

SOURCE (Reporte_fundamentos_Final.xlsx → Escuelas sheet):
  col 1  Cohorte                      → determines which B file
  col 2  Código de Infraestructura    → matches Código in B1/B2 CE
  col 4  Grado                        → Segundo Grado / Tercer Grado / Cuarto Grado
  col 5  Matrícula Evaluada
  col 6  % Aprobado (Mat)
  col 7  % Aprobado Condicional (Mat)
  col 8  % No Aprobado (Mat)
  col 9  % Aprobado (Lec)
  col 10 % Aprobado Condicional (Lec)
  col 11 % No Aprobado (Lec)

TARGET: 21 new columns appended at end of Centros Escolares, in order:
  For each grade in [2°, 3°, 4°]:
    Matrícula Evaluada {grade} Fundamentos
    % Aprobado (Mat) {grade} Fundamentos
    % Aprobado Condicional (Mat) {grade} Fundamentos
    % No Aprobado (Mat) {grade} Fundamentos
    % Aprobado (Lec) {grade} Fundamentos
    % Aprobado Condicional (Lec) {grade} Fundamentos
    % No Aprobado (Lec) {grade} Fundamentos

Matching key:  Código de Infraestructura (Reporte)  ==  Código (B-file CE col 1)
Cohorte B1 → B1 file,  Cohorte B2 → B2 file

Output saved to OUTPUT_DIR (created automatically).

Requirements: pip install openpyxl
"""

import os
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from collections import defaultdict

# =============================================================================
# PATHS — edit BASE_DIR if needed
# =============================================================================
BASE_DIR = (
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System"
    r"\Pedidos_Frecuentes_No_Mensual\Todos_los_puntajes_a_la_Fecha"
    r"\Incluiding_PM4\Actualizados_Con_Fundamentos"
)

B1_FILE  = rf"{BASE_DIR}\Centros_Escolares_B1_Actualizado_M4_Fund.xlsx"
B2_FILE  = rf"{BASE_DIR}\Centros_Escolares_B2_Actualizado_M4_Fund.xlsx"
RPT_FILE = rf"{BASE_DIR}\Reporte_fundamentos_Final.xlsx"

OUTPUT_DIR = rf"{BASE_DIR}\Con_Fundamentos_Por_Grado"
B1_OUT = rf"{OUTPUT_DIR}\Centros_Escolares_B1_Actualizado_M4_Fund_GradoFund.xlsx"
B2_OUT = rf"{OUTPUT_DIR}\Centros_Escolares_B2_Actualizado_M4_Fund_GradoFund.xlsx"

# =============================================================================
# COLUMN DEFINITIONS
# =============================================================================
# Grade labels as they appear in the Reporte and as suffix in output headers
GRADE_ORDER = [
    ("Segundo Grado",  "2°"),
    ("Tercer Grado",   "3°"),
    ("Cuarto Grado",   "4°"),
]

# The 7 base metric names (order matches Reporte cols 5-11)
METRIC_NAMES = [
    "Matrícula Evaluada",
    "% Aprobado (Mat)",
    "% Aprobado Condicional (Mat)",
    "% No Aprobado (Mat)",
    "% Aprobado (Lec)",
    "% Aprobado Condicional (Lec)",
    "% No Aprobado (Lec)",
]

# New headers: 7 metrics × 3 grades = 21 columns
# Label format: "{metric} {grade_label} Fundamentos"
NEW_HEADERS = [
    f"{metric} {grade_label} Fundamentos"
    for _, grade_label in GRADE_ORDER
    for metric in METRIC_NAMES
]

# =============================================================================
# STYLES
# =============================================================================
_thin       = Side(style="thin", color="BFBFBF")
BORDER      = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER      = Alignment(horizontal="center", vertical="center", wrap_text=True)
HEADER_FILL = PatternFill("solid", fgColor="FF2E75B6")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=9)


def apply_header(cell, text):
    cell.value     = text
    cell.fill      = HEADER_FILL
    cell.font      = HEADER_FONT
    cell.alignment = CENTER
    cell.border    = BORDER


def write_value(ws, row, col, value):
    """Write a data value with border and center alignment."""
    c = ws.cell(row, col, value)
    c.alignment = CENTER
    c.border    = BORDER
    # Format percentages (floats 0-1) as percentage
    if isinstance(value, float):
        c.number_format = "0.0%"


def blank_cell(ws, row, col):
    c = ws.cell(row, col)
    c.border    = BORDER
    c.alignment = CENTER


def normalize_code(raw):
    """Return school code as plain integer string."""
    if raw is None:
        return None
    try:
        return str(int(float(str(raw).strip())))
    except (ValueError, TypeError):
        return str(raw).strip()


# =============================================================================
# STEP 1 — Build lookup from Reporte_fundamentos_Final.xlsx
# =============================================================================
def build_lookup(rpt_path):
    """
    Returns:
      lookup : dict {
          cohorte (B1|B2) -> {
              codigo_str -> {
                  grado_str -> [matricula, pct_apr_mat, pct_cond_mat,
                                pct_no_mat, pct_apr_lec, pct_cond_lec,
                                pct_no_lec]
              }
          }
      }
    """
    wb  = openpyxl.load_workbook(rpt_path, read_only=True)
    ws  = wb["Escuelas"]
    lkp = defaultdict(lambda: defaultdict(dict))

    for row in ws.iter_rows(min_row=2, values_only=True):
        cohorte = str(row[0]).strip()  if row[0] else ''
        codigo  = normalize_code(row[1])
        grado   = str(row[3]).strip()  if row[3] else ''
        vals    = list(row[4:11])      # 7 values: Matrícula + 6 pct cols

        if cohorte in ('B1', 'B2') and codigo and grado:
            lkp[cohorte][codigo][grado] = vals

    wb.close()
    print(f"  Lookup built: B1={len(lkp['B1'])} schools, B2={len(lkp['B2'])} schools")
    return lkp


# =============================================================================
# STEP 2 — Append columns to one B-file
# =============================================================================
def update_ce_sheet(src_path, dst_path, cohorte_data, label):
    """
    Appends 21 new columns to the Centros Escolares sheet of src_path,
    saves to dst_path.
    cohorte_data : dict  { codigo_str -> { grado_str -> [7 values] } }
    """
    wb = openpyxl.load_workbook(src_path)
    ws = wb["Centros Escolares"]

    start_col = ws.max_column + 1   # first new column index
    n_new     = len(NEW_HEADERS)    # 21

    print(f"\n  {label}: CE has {ws.max_column} cols, "
          f"appending {n_new} new cols starting at col {start_col}")

    # ── Write headers ─────────────────────────────────────────────────────────
    for offset, header_text in enumerate(NEW_HEADERS):
        col = start_col + offset
        apply_header(ws.cell(1, col), header_text)
        ws.column_dimensions[get_column_letter(col)].width = 22

    # ── Write data rows ────────────────────────────────────────────────────────
    matched = 0
    unmatched = 0

    for row_idx in range(2, ws.max_row + 1):
        codigo = normalize_code(ws.cell(row_idx, 1).value)
        if codigo is None:
            continue

        school_data = cohorte_data.get(codigo)   # None if school not in Reporte

        col = start_col
        for grado_full, _ in GRADE_ORDER:
            if school_data and grado_full in school_data:
                vals = school_data[grado_full]   # list of 7 values
                for v in vals:
                    write_value(ws, row_idx, col, v)
                    col += 1
                matched += 1
            else:
                # School has no Reporte data for this grade → blank cells
                for _ in range(len(METRIC_NAMES)):
                    blank_cell(ws, row_idx, col)
                    col += 1

    total_grade_slots = (ws.max_row - 1) * len(GRADE_ORDER)
    print(f"  Data rows: {ws.max_row - 1} schools | "
          f"Grade slots filled: {matched} | blank: {total_grade_slots - matched}")

    wb.save(dst_path)
    print(f"  Saved → {dst_path}")


# =============================================================================
# MAIN
# =============================================================================
def main():
    print("=" * 65)
    print("APPENDING FUNDAMENTOS POR GRADO COLUMNS TO B1 AND B2 CE SHEETS")
    print("=" * 65)
    print(f"\nNew columns per file: {len(NEW_HEADERS)} "
          f"({len(METRIC_NAMES)} metrics × {len(GRADE_ORDER)} grades)")
    print("Column order:")
    for i, h in enumerate(NEW_HEADERS, 1):
        print(f"  {i:2d}. {h}")

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # Step 1: build lookup from Reporte
    print("\n[1] Reading Reporte_fundamentos_Final.xlsx...")
    lookup = build_lookup(RPT_FILE)

    # Step 2: update B1
    print("\n[2] Updating B1...")
    update_ce_sheet(B1_FILE, B1_OUT, lookup['B1'], "B1")

    # Step 3: update B2
    print("\n[3] Updating B2...")
    update_ce_sheet(B2_FILE, B2_OUT, lookup['B2'], "B2")

    print("\n" + "=" * 65)
    print("DONE. Files saved to:", OUTPUT_DIR)
    print("=" * 65)


if __name__ == "__main__":
    main()