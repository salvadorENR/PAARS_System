"""
verificar_cobertura.py
----------------------
Checks which schools in MAT-resultados.xlsx and LEC-resultados.xlsx
are NOT present in B1, B2, or Desempeño_Conjunto_B2_Rezago.

For each missing school reports:
  - School code and name
  - Total student count across all result files/sheets
  - Breakdown: how many students per file and per grade sheet

Requirements: pip install openpyxl
"""

import os
import openpyxl
from collections import defaultdict

# =============================================================================
# PATHS — edit BASE_DIR if needed
# =============================================================================
BASE_DIR = (
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System"
    r"\Pedidos_Frecuentes_No_Mensual\Todos_los_puntajes_a_la_Fecha"
    r"\Incluiding_PM4\Actualizados_Con_Fundamentos"
)

MAT_FILE    = rf"{BASE_DIR}\MAT-resultados.xlsx"
LEC_FILE    = rf"{BASE_DIR}\LEC-resultados.xlsx"
B1_FILE     = rf"{BASE_DIR}\Centros_Escolares_B1_Actualizado_M4_Fund.xlsx"
B2_FILE     = rf"{BASE_DIR}\Centros_Escolares_B2_Actualizado_M4_Fund.xlsx"
REZAGO_FILE = rf"{BASE_DIR}\Desempeño_Conjunto_B2_Rezago.xlsx"

GRADE_SHEETS = ['2do', '3er', '4to', '5to', '6to',
                '7mo', '8vo', '9no', 'Bach-1er', 'Bach-2do']


# =============================================================================
# HELPERS
# =============================================================================
def normalize_code(raw):
    if raw is None:
        return None
    try:
        return str(int(float(str(raw).strip())))
    except (ValueError, TypeError):
        return str(raw).strip()


# =============================================================================
# STEP 1 — Collect schools AND student counts/locations from resultados files
# =============================================================================
def collect_resultados_data(*file_paths):
    """
    Reads every grade sheet of each resultados file.

    Returns:
      schools       : dict { code -> name }
      student_locs  : dict { code -> { (file_label, sheet) -> student_count } }

    Columns (0-based):
      col 1 = 'Nro de centro'
      col 2 = 'Centro'
    """
    schools      = {}
    student_locs = defaultdict(lambda: defaultdict(int))

    for fpath in file_paths:
        file_label = os.path.basename(fpath)
        wb         = openpyxl.load_workbook(fpath, read_only=True)

        for sname in GRADE_SHEETS:
            if sname not in wb.sheetnames:
                continue
            ws = wb[sname]
            for row in ws.iter_rows(min_row=2, values_only=True):
                code = normalize_code(row[1])   # Nro de centro
                name = str(row[2]).strip() if row[2] else ""  # Centro
                if not code:
                    continue
                # Store school name (first occurrence wins)
                if code not in schools:
                    schools[code] = name
                # Count this student
                student_locs[code][(file_label, sname)] += 1

        wb.close()
        print(f"  {file_label}: escuelas acumuladas = {len(schools)}")

    return schools, student_locs


# =============================================================================
# STEP 2 — Collect Código values from B1, B2, and Rezago files
# =============================================================================
def collect_b_file_codes(fpath, label):
    codes = set()
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True)
    except FileNotFoundError:
        print(f"  ADVERTENCIA: {label} no encontrado en {fpath}")
        return codes

    target = "Centros Escolares" if "Centros Escolares" in wb.sheetnames else wb.sheetnames[0]
    ws     = wb[target]

    header    = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    codigo_col = 1
    for i, h in enumerate(header, start=1):
        if h is not None and str(h).strip().lower() in ("código", "codigo"):
            codigo_col = i
            break

    for row in ws.iter_rows(min_row=2, values_only=True):
        code = normalize_code(row[codigo_col - 1])
        if code:
            codes.add(code)

    wb.close()
    print(f"  {label} ({target}): {len(codes)} códigos cargados")
    return codes


def collect_rezago_codes(fpath, label):
    codes = set()
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True)
    except FileNotFoundError:
        print(f"  ADVERTENCIA: {label} no encontrado en {fpath}")
        return codes

    for sname in wb.sheetnames:
        ws   = wb[sname]
        rows = list(ws.iter_rows(max_row=2, values_only=True))
        if not rows:
            continue
        header   = [str(v).strip().lower() if v else "" for v in rows[0]]
        code_col = None
        for i, h in enumerate(header):
            if h in ("código", "codigo", "nro de centro", "code",
                     "código de infraestructura"):
                code_col = i
                break
        if code_col is None:
            code_col = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            code = normalize_code(row[code_col])
            if code:
                codes.add(code)

    wb.close()
    print(f"  {label}: {len(codes)} códigos cargados en todos los sheets")
    return codes


# =============================================================================
# MAIN
# =============================================================================
def main():
    W = 70   # report width

    print("=" * W)
    print("VERIFICACIÓN DE COBERTURA DE ESCUELAS — Prueba de Progreso M4")
    print("=" * W)

    # ── Step 1: schools + student locations from resultados ──────────────────
    print("\n[1] Leyendo escuelas y estudiantes de archivos de resultados...")
    schools, student_locs = collect_resultados_data(MAT_FILE, LEC_FILE)
    total_schools = len(schools)
    total_students_all = sum(
        n for locs in student_locs.values() for n in locs.values()
    )
    print(f"\n  Escuelas únicas (MAT + LEC): {total_schools}")
    print(f"  Registros de estudiantes:    {total_students_all}")

    # ── Step 2: codes in coverage files ──────────────────────────────────────
    print("\n[2] Leyendo códigos de B1, B2 y Rezago...")
    b1_codes     = collect_b_file_codes(B1_FILE,     "B1")
    b2_codes     = collect_b_file_codes(B2_FILE,     "B2")
    rezago_codes = collect_rezago_codes(REZAGO_FILE, "B2_Rezago")
    covered      = b1_codes | b2_codes | rezago_codes

    # ── Step 3: find missing schools ─────────────────────────────────────────
    print("\n[3] Identificando escuelas NO cubiertas...")
    missing = {
        code: schools[code]
        for code in schools
        if code not in covered
    }

    # ── Step 4: detailed report for each missing school ───────────────────────
    print("\n" + "=" * W)
    if not missing:
        print("✓ TODAS las escuelas de los resultados están cubiertas en B1, B2 o B2_Rezago.")
    else:
        print(f"ESCUELAS EN RESULTADOS PERO AUSENTES EN B1, B2 Y B2_REZAGO ({len(missing)}):")
        print("=" * W)

        total_missing_students = 0

        for code in sorted(missing, key=lambda x: x.zfill(10)):
            name  = missing[code]
            locs  = student_locs[code]           # {(file, sheet) -> count}
            total = sum(locs.values())
            total_missing_students += total

            print(f"\n  Código : {code}")
            print(f"  Nombre : {name}")
            print(f"  Total estudiantes con resultados: {total}")
            print(f"  Ubicación en archivos de resultados:")

            # Group by file first, then sheet
            by_file = defaultdict(dict)
            for (fname, sname), cnt in sorted(locs.items()):
                by_file[fname][sname] = cnt

            for fname in sorted(by_file):
                file_total = sum(by_file[fname].values())
                print(f"    ┌─ {fname}  ({file_total} estudiantes)")
                for sname in sorted(by_file[fname],
                                    key=lambda s: GRADE_SHEETS.index(s)
                                    if s in GRADE_SHEETS else 99):
                    cnt = by_file[fname][sname]
                    print(f"    │  Hoja '{sname}': {cnt} estudiante{'s' if cnt != 1 else ''}")
                print(f"    └{'─'*50}")

    # ── Step 5: summary ───────────────────────────────────────────────────────
    missing_students = sum(
        sum(locs.values())
        for code, locs in student_locs.items()
        if code not in covered
    )

    print("\n" + "=" * W)
    print("RESUMEN FINAL")
    print("-" * W)
    print(f"  Escuelas únicas en MAT + LEC resultados:        {total_schools:>5}")
    print(f"  Registros de estudiantes en resultados:         {total_students_all:>5}")
    print(f"  ─────────────────────────────────────────────────────")
    print(f"  Escuelas cubiertas en B1:                       {len(b1_codes & schools.keys()):>5}")
    print(f"  Escuelas cubiertas en B2:                       {len(b2_codes & schools.keys()):>5}")
    print(f"  Escuelas cubiertas en B2_Rezago:                {len(rezago_codes & schools.keys()):>5}")
    print(f"  Escuelas cubiertas en al menos un archivo:      {len(covered & schools.keys()):>5}")
    print(f"  ─────────────────────────────────────────────────────")
    print(f"  Escuelas NO cubiertas en ningún archivo:        {len(missing):>5}  ← posibles errores")
    print(f"  Estudiantes de escuelas NO cubiertas:           {missing_students:>5}  ← posibles errores")
    print("=" * W)

    if missing:
        print("\n⚠  Las escuelas y estudiantes listados arriba tienen resultados en")
        print("   MAT/LEC pero NO están en B1, B2 ni Desempeño_Conjunto_B2_Rezago.")
        print("   Esto indica que esos resultados no fueron incluidos — revisar.")

    return missing


if __name__ == "__main__":
    main()