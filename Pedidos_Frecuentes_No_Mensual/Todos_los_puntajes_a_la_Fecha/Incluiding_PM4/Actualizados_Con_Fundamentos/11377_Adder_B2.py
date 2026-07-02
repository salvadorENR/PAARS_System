"""
agregar_11377.py
----------------
Adds school 11377 (COMPLEJO EDUCATIVO CANTÓN EL TULE) to:
  Centros_Escolares_B2_Actualizado_M4_Fund.xlsx

EXACT COLUMN TARGETS (confirmed from file inspection):

  Estudiantes sheet (26 cols):
    col  1  Código
    col  2  Centro Escolar
    col  3  Tipo de Centro
    col  4  NIE
    col  5  Primer Nombre
    col  6  Segundo Nombre
    col  7  Primer Apellido
    col  8  Segundo Apellido
    col  9  Grado
    cols 10-12  (Estatus / Sección — left blank)
    cols 13-18  (CML + Fundamentos — left blank, school has no prior data)
    cols 19-20  Mat Puntaje/Nivel M3 — left blank (no M3 data for this school)
    col 21  MAT_Puntaje_M4      ← theta.global (escala 0-100) from MAT file
    col 22  MAT_Nivel_M4        ← classified + coloured
    cols 23-24  Len Puntaje/Nivel M3 — left blank
    col 25  Len  Puntaje M4 (Mayo)  ← theta.global (escala 0-100) from LEC file
    col 26  Len  Nivel M4 (Mayo)    ← classified + coloured

  Centros Escolares sheet (15 cols):
    col  1  Código
    col  2  Centro_Escolar
    col  3  Tipo_de_Centro
    cols 4-9   (CML + M3 — left blank)
    col 10  MAT_Promedio_M4  ← mean of all students' MAT M4 scores
    col 11  MAT_Nivel_M4     ← classified + coloured
    cols 12-13 (LEN M3 — left blank)
    col 14  LEN_Promedio_M4  ← mean of all students' LEC M4 scores
    col 15  LEN_Nivel_M4     ← classified + coloured

Nivel thresholds:  ≤35 Crítico | ≤45 Bajo | ≤55 Medio | ≤65 Bueno | >65 Excelente

Output saved to: OUTPUT_DIR (created automatically)
"""

import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# =============================================================================
# PATHS — edit BASE_DIR if needed
# =============================================================================
BASE_DIR = (
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System"
    r"\Pedidos_Frecuentes_No_Mensual\Todos_los_puntajes_a_la_Fecha"
    r"\Incluiding_PM4\Actualizados_Con_Fundamentos"
)

B2_FUND_FILE = rf"{BASE_DIR}\Centros_Escolares_B2_Actualizado_M4_Fund.xlsx"
MAT_FILE     = rf"{BASE_DIR}\MAT-resultados.xlsx"
LEC_FILE     = rf"{BASE_DIR}\LEC-resultados.xlsx"

OUTPUT_DIR   = rf"{BASE_DIR}\Con_Escuela_11377"
OUTPUT_FILE  = rf"{OUTPUT_DIR}\Centros_Escolares_B2_Actualizado_M4_Fund_11377.xlsx"

SCHOOL_CODE  = 11377
SCHOOL_NAME  = "COMPLEJO EDUCATIVO CANTÓN EL TULE"
SCHOOL_TIPO  = "Complejo Educativo"

GRADE_SHEETS = ['2do','3er','4to','5to','6to',
                '7mo','8vo','9no','Bach-1er','Bach-2do']

# =============================================================================
# STYLES
# =============================================================================
_thin  = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

NIVEL_FILL = {
    "Crítico":   PatternFill("solid", fgColor="FFC00000"),
    "Bajo":      PatternFill("solid", fgColor="FFFF8C2E"),
    "Medio":     PatternFill("solid", fgColor="FFFFD966"),
    "Bueno":     PatternFill("solid", fgColor="FF92D050"),
    "Excelente": PatternFill("solid", fgColor="FF00B050"),
}
NIVEL_FONT = {
    "Crítico":   Font(color="FFFFFF", bold=True, size=9),
    "Bajo":      Font(color="FFFFFF", bold=True, size=9),
    "Medio":     Font(color="000000", bold=True, size=9),
    "Bueno":     Font(color="000000", bold=True, size=9),
    "Excelente": Font(color="FFFFFF", bold=True, size=9),
}

# =============================================================================
# HELPERS
# =============================================================================
def clasificar(val):
    if val is None:
        return ""
    try:
        v = float(val)
    except (ValueError, TypeError):
        return ""
    if   v <= 35: return "Crítico"
    elif v <= 45: return "Bajo"
    elif v <= 55: return "Medio"
    elif v <= 65: return "Bueno"
    else:         return "Excelente"


def set_cell(ws, row, col, value, is_nivel=False):
    """Write value with border and center alignment; apply colour if nivel."""
    c = ws.cell(row, col, value)
    c.alignment = CENTER
    c.border    = BORDER
    if is_nivel and value in NIVEL_FILL:
        c.fill = NIVEL_FILL[value]
        c.font = NIVEL_FONT[value]


def blank_cell(ws, row, col):
    ws.cell(row, col).border    = BORDER
    ws.cell(row, col).alignment = CENTER


def find_theta_col(header_row):
    """Return 0-based index of 'theta.global (escala 0-100)' column."""
    for i, h in enumerate(header_row):
        if h and '0-100' in str(h):
            return i
    return None


# =============================================================================
# STEP 1 — Read school 11377 students from MAT and LEC result files
# =============================================================================
def read_results():
    """
    Returns:
      students    : dict { nie (int) -> {nombre, apellido, grado,
                                         mat_score, lec_score} }
      mat_scores  : [float]
      lec_scores  : [float]
    """
    students = {}

    for fpath, subject in [(MAT_FILE, 'MAT'), (LEC_FILE, 'LEC')]:
        wb = openpyxl.load_workbook(fpath, read_only=True)

        for sname in GRADE_SHEETS:
            if sname not in wb.sheetnames:
                continue
            ws       = wb[sname]
            hdrs     = list(ws.iter_rows(max_row=1, values_only=True))[0]
            t_idx    = find_theta_col(hdrs)   # shifts between sheets — must be dynamic

            if t_idx is None:
                print(f"  WARNING: no 0-100 theta col in {subject}/{sname} — skipping")
                continue

            for row in ws.iter_rows(min_row=2, values_only=True):
                if str(row[1]).strip() != str(SCHOOL_CODE):
                    continue

                # Parse NIE (col 6, index 5)
                nie_raw = row[5]
                try:
                    nie = int(float(str(nie_raw).strip()))
                except (ValueError, TypeError):
                    continue

                # Parse theta (0-100 scale)
                try:
                    theta = float(row[t_idx]) if row[t_idx] is not None else None
                except (ValueError, TypeError):
                    theta = None

                # Build/update student entry
                if nie not in students:
                    nombre = str(row[6] or '').strip()
                    apell  = str(row[7] or '').strip()
                    grado  = str(row[3] or '').strip()
                    students[nie] = {
                        'nombre':    nombre,
                        'apellido':  apell,
                        'grado':     grado,
                        'mat_score': None,
                        'lec_score': None,
                    }

                if subject == 'MAT':
                    students[nie]['mat_score'] = theta
                else:
                    students[nie]['lec_score'] = theta

        wb.close()

    # Separate score lists for CE averages
    mat_scores = [s['mat_score'] for s in students.values()
                  if s['mat_score'] is not None]
    lec_scores = [s['lec_score'] for s in students.values()
                  if s['lec_score'] is not None]

    print(f"  Students loaded: {len(students)}")
    print(f"  MAT scores: {len(mat_scores)} | LEC scores: {len(lec_scores)}")
    return students, mat_scores, lec_scores


# =============================================================================
# STEP 2 — Append rows to Estudiantes sheet
# =============================================================================
def append_estudiantes(ws, students):
    """
    Appends one row per student.
    Cols 1-9   : identity (Código, Centro, Tipo, NIE, nombres, apellidos, Grado)
    Cols 10-20 : blank (no prior-month or CML data for this school)
    Col  21    : MAT_Puntaje_M4   ← from MAT theta 0-100
    Col  22    : MAT_Nivel_M4     ← classified + coloured
    Cols 23-24 : blank (no M3 LEN data)
    Col  25    : Len  Puntaje M4 (Mayo)  ← from LEC theta 0-100
    Col  26    : Len  Nivel M4 (Mayo)    ← classified + coloured
    """
    TOTAL_COLS = ws.max_column   # 26

    # Confirmed fixed positions (from file inspection):
    COL_CODIGO    = 1
    COL_CE_NAME   = 2
    COL_TIPO      = 3
    COL_NIE       = 4
    COL_PNOMBRE   = 5
    COL_SNOMBRE   = 6
    COL_PAPELL    = 7
    COL_SAPELL    = 8
    COL_GRADO     = 9
    COL_MAT_P4    = 21   # MAT_Puntaje_M4
    COL_MAT_N4    = 22   # MAT_Nivel_M4
    COL_LEN_P4    = 25   # Len  Puntaje M4 (Mayo)
    COL_LEN_N4    = 26   # Len  Nivel M4 (Mayo)

    added = 0
    for nie, s in sorted(students.items()):
        r = ws.max_row + 1

        # Apply border to every cell in the new row first
        for c in range(1, TOTAL_COLS + 1):
            blank_cell(ws, r, c)

        # Split Nombre → Primer / Segundo  (split on first space)
        nombre_parts = s['nombre'].split(' ', 1)
        apell_parts  = s['apellido'].split(' ', 1)
        primer_n = nombre_parts[0]
        segundo_n = nombre_parts[1] if len(nombre_parts) > 1 else ''
        primer_a = apell_parts[0]
        segundo_a = apell_parts[1] if len(apell_parts) > 1 else ''

        # Identity columns
        set_cell(ws, r, COL_CODIGO,  SCHOOL_CODE)
        set_cell(ws, r, COL_CE_NAME, SCHOOL_NAME)
        set_cell(ws, r, COL_TIPO,    SCHOOL_TIPO)
        set_cell(ws, r, COL_NIE,     nie)
        set_cell(ws, r, COL_PNOMBRE, primer_n)
        set_cell(ws, r, COL_SNOMBRE, segundo_n)
        set_cell(ws, r, COL_PAPELL,  primer_a)
        set_cell(ws, r, COL_SAPELL,  segundo_a)
        set_cell(ws, r, COL_GRADO,   s['grado'])

        # MAT M4
        if s['mat_score'] is not None:
            mat_score  = round(s['mat_score'], 2)
            mat_nivel  = clasificar(mat_score)
            set_cell(ws, r, COL_MAT_P4, mat_score)
            set_cell(ws, r, COL_MAT_N4, mat_nivel, is_nivel=True)

        # LEC M4
        if s['lec_score'] is not None:
            lec_score  = round(s['lec_score'], 2)
            lec_nivel  = clasificar(lec_score)
            set_cell(ws, r, COL_LEN_P4, lec_score)
            set_cell(ws, r, COL_LEN_N4, lec_nivel, is_nivel=True)

        added += 1

    print(f"  {added} student rows appended to Estudiantes")


# =============================================================================
# STEP 3 — Append school row to Centros Escolares sheet
# =============================================================================
def append_centros_escolares(ws, mat_scores, lec_scores):
    """
    Appends one school row.
    Cols 1-3   : Código, Centro_Escolar, Tipo_de_Centro
    Cols 4-9   : blank  (no CML / M3 data)
    Col  10    : MAT_Promedio_M4  ← mean of student MAT M4 scores
    Col  11    : MAT_Nivel_M4     ← classified + coloured
    Cols 12-13 : blank  (LEN M3)
    Col  14    : LEN_Promedio_M4  ← mean of student LEC M4 scores
    Col  15    : LEN_Nivel_M4     ← classified + coloured
    """
    TOTAL_COLS = ws.max_column   # 15

    COL_CODIGO   = 1
    COL_CE_NAME  = 2
    COL_TIPO     = 3
    COL_MAT_P4   = 10   # MAT_Promedio_M4
    COL_MAT_N4   = 11   # MAT_Nivel_M4
    COL_LEN_P4   = 14   # LEN_Promedio_M4
    COL_LEN_N4   = 15   # LEN_Nivel_M4

    avg_mat = round(sum(mat_scores) / len(mat_scores), 2) if mat_scores else None
    avg_lec = round(sum(lec_scores) / len(lec_scores), 2) if lec_scores else None

    r = ws.max_row + 1

    # Apply border to every cell first
    for c in range(1, TOTAL_COLS + 1):
        blank_cell(ws, r, c)

    # Identity
    set_cell(ws, r, COL_CODIGO,  SCHOOL_CODE)
    set_cell(ws, r, COL_CE_NAME, SCHOOL_NAME)
    set_cell(ws, r, COL_TIPO,    SCHOOL_TIPO)

    # MAT M4 average
    if avg_mat is not None:
        mat_nivel = clasificar(avg_mat)
        set_cell(ws, r, COL_MAT_P4, avg_mat)
        set_cell(ws, r, COL_MAT_N4, mat_nivel, is_nivel=True)
        print(f"  MAT_Promedio_M4 = {avg_mat}  →  {mat_nivel}")

    # LEC M4 average
    if avg_lec is not None:
        lec_nivel = clasificar(avg_lec)
        set_cell(ws, r, COL_LEN_P4, avg_lec)
        set_cell(ws, r, COL_LEN_N4, lec_nivel, is_nivel=True)
        print(f"  LEN_Promedio_M4 = {avg_lec}  →  {lec_nivel}")

    print(f"  School row appended to Centros Escolares (row {r})")


# =============================================================================
# MAIN
# =============================================================================
def main():
    print("=" * 60)
    print(f"ADDING SCHOOL {SCHOOL_CODE} — {SCHOOL_NAME}")
    print("=" * 60)

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # Step 1: collect student scores from MAT + LEC result files
    print("\n[1] Reading MAT and LEC results for school 11377...")
    students, mat_scores, lec_scores = read_results()

    if not students:
        print("  ERROR: no students found for school 11377. Check file paths.")
        return

    # Step 2: load B2 Fund workbook
    print("\n[2] Loading B2 Fund workbook...")
    wb     = openpyxl.load_workbook(B2_FUND_FILE)
    ws_est = wb["Estudiantes"]
    ws_ce  = wb["Centros Escolares"]
    print(f"  Estudiantes: {ws_est.max_row - 1} existing rows")
    print(f"  Centros Escolares: {ws_ce.max_row - 1} existing rows")

    # Step 3: append to Estudiantes
    print("\n[3] Appending students to Estudiantes...")
    append_estudiantes(ws_est, students)

    # Step 4: append to Centros Escolares
    print("\n[4] Appending school to Centros Escolares...")
    append_centros_escolares(ws_ce, mat_scores, lec_scores)

    # Step 5: save
    wb.save(OUTPUT_FILE)
    print(f"\n[5] Saved → {OUTPUT_FILE}")

    print("\n" + "=" * 60)
    print("SUMMARY")
    print("-" * 60)
    print(f"  Students added to Estudiantes : {len(students)}")
    print(f"  Schools added to CE           : 1  ({SCHOOL_CODE})")
    print(f"  MAT scores used for average   : {len(mat_scores)}")
    print(f"  LEC scores used for average   : {len(lec_scores)}")
    print(f"  Output file : {OUTPUT_FILE}")
    print("=" * 60)


if __name__ == "__main__":
    main()