import pandas as pd
import glob
import os
import statistics
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────
# 1. PATH CONFIGURATIONS
# ─────────────────────────────────────────────
FEB_DIR = r"H:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\01_Resultados_Febrero\Interim_CSVs\Resultados"
EXCEL_FILES = [
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Centros_Escolares_B1.xlsx",
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Criterio_Inclusión_Exclusión_Escuelas\Centros_Escolares_B2.xlsx"
]
FUNDAMENTOS_FILE = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Todos_los_puntajes_a_la_Fecha\Fundamentos_Consolidados.xlsx"
OUTPUT_DIR = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\Todos_los_puntajes_a_la_Fecha"

def clean_id(val):
    if pd.isna(val) or val == "": return ""
    try: return str(int(float(val)))
    except: return str(val).strip()

def get_nivel(score):
    """Calculates the Nivel category based on the standard scoring scale."""
    if pd.isna(score) or score == "": return ""
    try:
        s = float(score)
        if s < 35.0: return "Crítico"
        elif s < 45.0: return "Bajo"
        elif s < 55.0: return "Medio"
        elif s < 65.0: return "Bueno"
        else: return "Excelente"
    except:
        return ""

# ─────────────────────────────────────────────
# 2. LOAD EXTERNAL DATA 
# ─────────────────────────────────────────────
print("Loading February CSVs...")
csv_files = glob.glob(os.path.join(FEB_DIR, "*.csv"))
feb_scores = {}
for f in csv_files:
    try:
        df = pd.read_csv(f)
        # Check if a nivel column exists
        nivel_col = next((c for c in df.columns if 'nivel' in str(c).lower()), None)
        
        for _, row in df.iterrows():
            doc = clean_id(row.get('Documento'))
            subj = str(row.get('Area temática', '')).strip().lower()
            score = row.get('theta.global (escala 0-100)')
            
            # Fetch Nivel from CSV if it exists, otherwise CALCULATE it using get_nivel()
            if nivel_col and pd.notna(row.get(nivel_col)):
                nivel = str(row.get(nivel_col)).strip().capitalize()
            else:
                nivel = get_nivel(score)
            
            if pd.notna(score) and doc:
                if doc not in feb_scores: feb_scores[doc] = {}
                if 'mat' in subj:
                    feb_scores[doc]['mat_score'], feb_scores[doc]['mat_nivel'] = score, nivel
                elif 'lec' in subj or 'len' in subj:
                    feb_scores[doc]['len_score'], feb_scores[doc]['len_nivel'] = score, nivel
    except Exception as e: 
        print(f" -> Could not read {os.path.basename(f)}: {e}")

print(f" -> Loaded February scores for {len(feb_scores)} students.")

print("Loading Fundamentos Data directly from Excel...")
fund_stu_data, fund_sch_data = {}, {}
try:
    df_f_stu = pd.read_excel(FUNDAMENTOS_FILE, sheet_name='Estudiantes')
    for _, row in df_f_stu.iterrows():
        nie = clean_id(row.iloc[3])
        if nie:
            fund_stu_data[nie] = {
                'lec_score': row.iloc[4], 'lec_cat': str(row.iloc[5]).strip().capitalize() if pd.notna(row.iloc[5]) else get_nivel(row.iloc[4]),
                'mat_score': row.iloc[6], 'mat_cat': str(row.iloc[7]).strip().capitalize() if pd.notna(row.iloc[7]) else get_nivel(row.iloc[6])
            }
except Exception as e: print(f" -> Error Fundamentos Estudiantes: {e}")

try:
    df_f_sch = pd.read_excel(FUNDAMENTOS_FILE, sheet_name='Escuelas', skiprows=2, header=None)
    for _, row in df_f_sch.iterrows():
        cod = clean_id(row.iloc[0])
        if cod and pd.notna(row.iloc[9]): fund_sch_data[cod] = str(row.iloc[9]).strip()
except Exception as e: print(f" -> Error Fundamentos Escuelas: {e}")

# ─────────────────────────────────────────────
# 3. UPDATE THE EXCEL FILES
# ─────────────────────────────────────────────
os.makedirs(OUTPUT_DIR, exist_ok=True)
header_fill_main = PatternFill('solid', start_color='2E4057', end_color='2E4057')
header_fill_sub  = PatternFill('solid', start_color='4A6FA5', end_color='4A6FA5')
header_font_main = Font(name='Arial', bold=True, color='FFFFFF', size=10)
header_font_sub  = Font(name='Arial', bold=True, color='FFFFFF', size=9)
body_font        = Font(name='Arial', size=9)
alt_fill         = PatternFill('solid', start_color='F2F2F2', end_color='F2F2F2')
white_fill       = PatternFill('solid', start_color='FFFFFF', end_color='FFFFFF')
center_align     = Alignment(horizontal='center', vertical='center', wrap_text=True)
thin_border      = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'), 
                          top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))

STATUS_COLORS = {'Excelente': '1E7145', 'Bueno': '70AD47', 'Regular': 'FFD966', 'Alerta': 'FF4B4B'}
STATUS_FONT = {'Excelente': 'FFFFFF', 'Bueno': 'FFFFFF', 'Regular': '000000', 'Alerta': 'FFFFFF'}
NIVEL_COLORS = {'Crítico': 'FF4B4B', 'Bajo': 'FF9966', 'Medio': 'FFD966', 'Bueno': 'A9D18E', 'Excelente': '1E7145'}
NIVEL_FONT = {'Crítico': 'FFFFFF', 'Bajo': 'FFFFFF', 'Medio': '000000', 'Bueno': '000000', 'Excelente': 'FFFFFF'}

for excel_path in EXCEL_FILES:
    if not os.path.exists(excel_path): continue
        
    base_filename = os.path.basename(excel_path)
    print(f"\nUpdating: {base_filename}")
    wb = load_workbook(excel_path)
    school_scores = {}
    ordered_exams = []
    
    # ── A. UPDATE "ESTUDIANTES" SHEET ──
    if 'Estudiantes' in wb.sheetnames:
        ws_stu = wb['Estudiantes']
        nie_col_idx = next((c for c in range(1, ws_stu.max_column + 1) if ws_stu.cell(1, c).value == 'NIE' or ws_stu.cell(2, c).value == 'NIE'), None)
        cod_col_idx = next((c for c in range(1, ws_stu.max_column + 1) if ws_stu.cell(1, c).value == 'Código' or ws_stu.cell(2, c).value == 'Código'), None)
        
        if nie_col_idx and cod_col_idx:
            mc = ws_stu.max_column
            
            # Map existing score columns for aggregations
            score_cols_map = {}
            current_subj = ""
            for c in range(1, mc + 1):
                val1 = ws_stu.cell(1, c).value
                if val1 and isinstance(val1, str) and val1.strip() != "": current_subj = val1.strip()
                val2 = ws_stu.cell(2, c).value
                if val2 and isinstance(val2, str) and "Puntaje" in val2:
                    exam_name = f"{current_subj} - {val2.replace('Puntaje ', '')}"
                    score_cols_map[c] = exam_name
                    ordered_exams.append(exam_name)
                    
            # Define new columns
            new_cols = [
                ("Matemática (Feb - Conociendo Mis Logros)", "Puntaje Febrero", "Matemática - Febrero"), 
                ("Matemática (Feb - Conociendo Mis Logros)", "Nivel Febrero", None),
                ("Lenguaje (Feb - Conociendo Mis Logros)", "Puntaje Febrero", "Lenguaje - Febrero"), 
                ("Lenguaje (Feb - Conociendo Mis Logros)", "Nivel Febrero", None),
                ("Matemática Fundamentos", "Puntaje MAT Fund.", "Matemática - Fundamentos"), 
                ("Matemática Fundamentos", "Nivel MAT Fund.", None),
                ("Lenguaje Fundamentos", "Puntaje LEC Fund.", "Lenguaje - Fundamentos"), 
                ("Lenguaje Fundamentos", "Nivel LEC Fund.", None)
            ]
            
            for tup in new_cols:
                if tup[2]: ordered_exams.append(tup[2])
            
            # Write Headers
            for i, (main_h, sub_h, _) in enumerate(new_cols):
                col = mc + 1 + i
                ws_stu.cell(1, col, value=main_h).fill, ws_stu.cell(1, col).font = header_fill_main, header_font_main
                ws_stu.cell(2, col, value=sub_h).fill, ws_stu.cell(2, col).font = header_fill_sub, header_font_sub
                for r in [1, 2]: ws_stu.cell(r, col).alignment, ws_stu.cell(r, col).border = center_align, thin_border
                ws_stu.column_dimensions[get_column_letter(col)].width = 16

            # Write Data and Accumulate Scores
            for r in range(3, ws_stu.max_row + 1):
                nie_val = clean_id(ws_stu.cell(r, nie_col_idx).value)
                cod_val = clean_id(ws_stu.cell(r, cod_col_idx).value)
                fill_to_use = alt_fill if (r - 3) % 2 == 0 else white_fill
                
                if cod_val not in school_scores: school_scores[cod_val] = {}
                
                # Accumulate existing scores
                for c_idx, exam_name in score_cols_map.items():
                    val = ws_stu.cell(r, c_idx).value
                    if exam_name not in school_scores[cod_val]: school_scores[cod_val][exam_name] = []
                    if pd.notna(val) and val != "" and isinstance(val, (int, float)):
                        school_scores[cod_val][exam_name].append(float(val))
                
                # Fetch New Data
                m_feb_sc = round(feb_scores.get(nie_val, {}).get('mat_score', ""), 2) if isinstance(feb_scores.get(nie_val, {}).get('mat_score', ""), (int, float)) else ""
                m_feb_niv = feb_scores.get(nie_val, {}).get('mat_nivel', "")
                l_feb_sc = round(feb_scores.get(nie_val, {}).get('len_score', ""), 2) if isinstance(feb_scores.get(nie_val, {}).get('len_score', ""), (int, float)) else ""
                l_feb_niv = feb_scores.get(nie_val, {}).get('len_nivel', "")
                
                stu_f = fund_stu_data.get(nie_val, {})
                m_f_sc, m_f_cat = stu_f.get('mat_score', ""), stu_f.get('mat_cat', "")
                l_f_sc, l_f_cat = stu_f.get('lec_score', ""), stu_f.get('lec_cat', "")
                
                vals = [m_feb_sc, m_feb_niv, l_feb_sc, l_feb_niv, m_f_sc, m_f_cat, l_f_sc, l_f_cat]
                
                # Write New Columns & Accumulate
                for i, val in enumerate(vals):
                    c = ws_stu.cell(r, mc + 1 + i, value=val)
                    c.border, c.alignment = thin_border, center_align
                    
                    if 'Nivel' in new_cols[i][1] and pd.notna(val) and val in NIVEL_COLORS:
                        c.fill, c.font = PatternFill('solid', fgColor=NIVEL_COLORS[val]), Font(name='Arial', size=9, color=NIVEL_FONT[val])
                    else:
                        c.fill, c.font = fill_to_use, body_font
                    
                    # Accumulate new scores
                    exam_name_tracker = new_cols[i][2]
                    if exam_name_tracker:
                        if exam_name_tracker not in school_scores[cod_val]: school_scores[cod_val][exam_name_tracker] = []
                        if pd.notna(val) and val != "" and isinstance(val, (int, float)):
                            school_scores[cod_val][exam_name_tracker].append(float(val))

    # ── B. UPDATE "CENTROS ESCOLARES" SHEET ──
    if 'Centros Escolares' in wb.sheetnames:
        ws_sch = wb['Centros Escolares']
        sch_cod_idx = next((c for c in range(1, ws_sch.max_column + 1) if ws_sch.cell(1, c).value == 'Código' or ws_sch.cell(2, c).value == 'Código'), None)
        
        if sch_cod_idx:
            mc = ws_sch.max_column
            estatus_col = mc + 1
            
            # Estatus Escuela Fundamentos Column
            c1 = ws_sch.cell(1, estatus_col, value="Estatus Escuela\nFundamentos")
            c1.fill, c1.font, c1.alignment, c1.border = header_fill_main, header_font_main, center_align, thin_border
            ws_sch.merge_cells(start_row=1, start_column=estatus_col, end_row=2, end_column=estatus_col)
            ws_sch.cell(2, estatus_col).border = thin_border
            ws_sch.column_dimensions[get_column_letter(estatus_col)].width = 18
            
            # Generate Mean/Median Headers
            start_agg_col = estatus_col + 1
            curr_c = start_agg_col
            for exam_name in ordered_exams:
                ws_sch.cell(1, curr_c, value=exam_name)
                ws_sch.cell(1, curr_c).fill, ws_sch.cell(1, curr_c).font = header_fill_main, header_font_main
                ws_sch.cell(1, curr_c).alignment, ws_sch.cell(1, curr_c).border = center_align, thin_border
                ws_sch.merge_cells(start_row=1, start_column=curr_c, end_row=1, end_column=curr_c+3)
                
                sub_headers = ["Promedio", "Nivel Prom.", "Mediana", "Nivel Med."]
                for idx, sh in enumerate(sub_headers):
                    ws_sch.cell(2, curr_c + idx, value=sh)
                    ws_sch.cell(2, curr_c + idx).fill, ws_sch.cell(2, curr_c + idx).font = header_fill_sub, header_font_sub
                    ws_sch.cell(2, curr_c + idx).alignment, ws_sch.cell(2, curr_c + idx).border = center_align, thin_border
                    ws_sch.column_dimensions[get_column_letter(curr_c + idx)].width = 14 if 'Nivel' in sh else 11
                curr_c += 4

            # Write Data
            for r in range(3, ws_sch.max_row + 1):
                cod_val = clean_id(ws_sch.cell(r, sch_cod_idx).value)
                fill_to_use = alt_fill if (r - 3) % 2 == 0 else white_fill
                
                # Estatus Fundamentos
                estatus_val = fund_sch_data.get(cod_val, "")
                c = ws_sch.cell(r, estatus_col, value=estatus_val)
                c.border, c.alignment = thin_border, center_align
                if pd.notna(estatus_val) and estatus_val in STATUS_COLORS:
                    c.fill, c.font = PatternFill('solid', fgColor=STATUS_COLORS[estatus_val]), Font(name='Arial', size=9, bold=True, color=STATUS_FONT[estatus_val])
                else: c.fill, c.font = fill_to_use, body_font
                
                # Means and Medians
                curr_c = start_agg_col
                for exam_name in ordered_exams:
                    scores = school_scores.get(cod_val, {}).get(exam_name, [])
                    if len(scores) > 0:
                        prom, med = statistics.mean(scores), statistics.median(scores)
                        niv_prom, niv_med = get_nivel(prom), get_nivel(med)
                    else:
                        prom, med, niv_prom, niv_med = "", "", "", ""
                    
                    vals = [round(prom, 2) if prom != "" else "", niv_prom, round(med, 2) if med != "" else "", niv_med]
                    for idx, val in enumerate(vals):
                        c_agg = ws_sch.cell(r, curr_c + idx, value=val)
                        c_agg.border, c_agg.alignment = thin_border, center_align
                        if val in NIVEL_COLORS:
                            c_agg.fill, c_agg.font = PatternFill('solid', fgColor=NIVEL_COLORS[val]), Font(name='Arial', size=9, color=NIVEL_FONT[val])
                        else: c_agg.fill, c_agg.font = fill_to_use, body_font
                    
                    curr_c += 4

    out_path = os.path.join(OUTPUT_DIR, base_filename.replace(".xlsx", "_Actualizado.xlsx"))
    wb.save(out_path)
    print(f" -> Success! Saved to: {out_path}")