"""
Mejores_Escuelas_Sistema_Tasa_Victorias.py
===================================
Evalúa a las escuelas mediante un sistema de "Torneo" (1 punto si gana, 0 si pierde).
Se compite estrictamente Grado vs Grado dentro del mismo TIPO de institución 
(Institutos vs Institutos, Complejos vs Complejos, Centros Escolares vs Centros Escolares).

FASE 5 — MEDIA PONDERADA:
  El Win_Rate_Final de cada escuela se calcula como la media ponderada de los
  Win_Rate_Grado, usando Total_Examenes_Grado como peso. Esto evita que un grado
  pequeño (pocas observaciones, estimación ruidosa) tenga el mismo peso que un
  grado grande al momento de calcular el puntaje final de la institución.

  Fórmula:
      Win_Rate_Final = Σ(WR_g × n_g) / Σ(n_g)
  donde WR_g es el Win Rate del grado g y n_g es el total de exámenes en ese grado.
"""

import os
import glob
import unicodedata
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ══════════════════════════════════════════════════════════════════════════════
# 1. CONFIGURACIÓN Y RUTAS
# ══════════════════════════════════════════════════════════════════════════════
MESES = [
    {"mes": 1, "ruta": r"H:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\01_PROGRESO_Marzo\Interim_CSVs\Resultados"},
    {"mes": 2, "ruta": r"H:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\02_PROGRESO_Abril\Interim_CSVs\Resultados"},
    {"mes": 3, "ruta": r"H:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\03_PROGRESO_Mayo\Interim_CSVs\Resultados"},
    {"mes": 4, "ruta": r"H:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026\04_PROGRESO_Junio\Interim_CSVs\Resultados"}
]

PATH_MATRICULA = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\13. Matricula Progreso Mes 1-Mes 4\Matricula_P1.xlsx"
PATH_OUTPUT = r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System\Pedidos_Frecuentes_No_Mensual\13. Matricula Progreso Mes 1-Mes 4"

VIRTUAL_CODES = {"99999", "99998"}
COL_SCORE = "theta.global (escala 0-100)"

# ══════════════════════════════════════════════════════════════════════════════
# 2. FUNCIONES AUXILIARES
# ══════════════════════════════════════════════════════════════════════════════
def inferir_materia(nombre):
    n = nombre.upper()
    if "MAT" in n: return "MAT"
    if "LEC" in n or "LEN" in n: return "LEC"
    return "DESCONOCIDO"

def clasificar_puntaje(val):
    if pd.isna(val): return "Critico"
    try: val = float(val)
    except: return "Critico"
    if val <= 35: return "Critico"
    elif val <= 45: return "Bajo"
    elif val <= 55: return "Medio"
    elif val <= 65: return "Bueno"
    else: return "Excelente"

def quitar_tildes(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn")

def excluir_por_tiempo(val):
    if pd.isna(val): return False
    v = quitar_tildes(str(val).lower())
    return "5 min" in v or "demora" in v or "duracion" in v

def clasificar_tipo_escuela(nombre):
    n = str(nombre).upper()
    if "INSTITUTO" in n: return "Instituto"
    elif "COMPLEJO" in n: return "Complejo"
    else: return "Centro Escolar"

def normalize_dept(d):
    if not isinstance(d, str): return "DESCONOCIDO"
    d = d.replace("Á", "A").replace("É", "E").replace("Í", "I").replace("Ó", "O").replace("Ú", "U")
    return d.strip()

# ══════════════════════════════════════════════════════════════════════════════
# 3. CARGA DE DATOS E INFERENCIA DE TIPO DE ESCUELA
# ══════════════════════════════════════════════════════════════════════════════
print("[*] Cargando catálogo de departamentos...")
try:
    df_mat = pd.read_excel(PATH_MATRICULA, dtype=str)
    dep_col = next((c for c in df_mat.columns if "DEPART" in c.upper() and "CÓDIGO" not in c.upper() and "CODIGO" not in c.upper()), None)
    if dep_col and "CODIGO" in df_mat.columns:
        dept_map = df_mat.drop_duplicates("CODIGO").set_index("CODIGO")[dep_col].str.strip().str.upper().to_dict()
    else: dept_map = {}
except: dept_map = {}

all_data = []
for m in MESES:
    archivos = glob.glob(os.path.join(m["ruta"], "*.csv"))
    print(f"[*] Procesando Mes {m['mes']}: {len(archivos)} archivos encontrados.")
    for f in archivos:
        nombre_archivo = os.path.basename(f)
        if "legend" in nombre_archivo.lower(): continue
        df = None
        for enc in ["utf-8", "latin-1", "cp1252"]:
            try:
                df = pd.read_csv(f, dtype=str, encoding=enc, encoding_errors="ignore")
                df.columns = df.columns.str.strip()
                break
            except Exception:
                continue
        if df is None: continue

        if COL_SCORE not in df.columns or "Nro de centro" not in df.columns or "Grado" not in df.columns:
            continue

        if "Centro" not in df.columns: df["Centro"] = "Sin Nombre Registrado"

        df["Nro de centro"] = df["Nro de centro"].astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
        df["Puntaje"] = pd.to_numeric(df[COL_SCORE].astype(str).str.replace(",", "."), errors="coerce")
        df["Materia"] = inferir_materia(nombre_archivo)
        df["Es_Invalido_Tiempo"] = df["anular_prueba"].apply(excluir_por_tiempo) if "anular_prueba" in df.columns else False

        df = df.dropna(subset=["Puntaje"])
        df = df[~df["Nro de centro"].isin(VIRTUAL_CODES)]
        df = df[~df["Centro"].astype(str).str.contains("Centro Virtual", case=False, na=False)]

        if not df.empty:
            df["Nivel"] = df["Puntaje"].apply(clasificar_puntaje)
            df["Es_Bueno_Exc"] = df["Nivel"].isin(["Bueno", "Excelente"]).astype(int)
            df["Mes"] = m["mes"]
            cols_seguras = ["Nro de centro", "Centro", "Grado", "Materia", "Mes", "Es_Bueno_Exc", "Es_Invalido_Tiempo"]
            all_data.append(df[cols_seguras])

if not all_data:
    raise ValueError(
        "\n\n274c No se cargo ningun archivo CSV.\n"
        "   Causas mas comunes:\n"
        "   1. Las rutas en MESES no existen o estan mal escritas.\n"
        "   2. Los CSVs no tienen la columna: 'theta.global (escala 0-100)'\n"
        "   3. Los CSVs no tienen las columnas 'Nro de centro' o 'Grado'.\n"
        "   4. Todos los archivos se saltan por el filtro 'legend'.\n"
        "   -> Corre primero diagnostico_carga.py para identificar el problema exacto."
    )

df_total = pd.concat(all_data, ignore_index=True)
df_total["Departamento"] = df_total["Nro de centro"].map(dept_map).apply(normalize_dept)

df_total["Tipo_Institucion"] = df_total["Centro"].apply(clasificar_tipo_escuela)
df_validos = df_total[~df_total["Es_Invalido_Tiempo"]].copy()

# ══════════════════════════════════════════════════════════════════════════════
# 4. MOTOR DE TORNEO (BATALLAS 1 VS 0 SEGURAS)
# ══════════════════════════════════════════════════════════════════════════════
def calcular_tasa_victorias(df_subset):
    if df_subset.empty: return pd.DataFrame()

    # Exigir 4 meses
    escuelas_meses = df_subset.groupby("Nro de centro")["Mes"].nunique()
    escuelas_validas = escuelas_meses[escuelas_meses == 4].index
    df_work = df_subset[df_subset["Nro de centro"].isin(escuelas_validas)].copy()
    if df_work.empty: return pd.DataFrame()

    # 1. Porcentaje bruto por grado
    grade_stats = df_work.groupby(["Nro de centro", "Centro", "Tipo_Institucion", "Departamento", "Grado"]).agg(
        Total_Examenes_Grado=("Es_Bueno_Exc", "count"),
        Total_Bueno_Exc=("Es_Bueno_Exc", "sum")
    ).reset_index()
    grade_stats["Pct_Bueno_Exc"] = (grade_stats["Total_Bueno_Exc"] / grade_stats["Total_Examenes_Grado"]) * 100

    # 2. Torneo: Comparar 1 vs Todos dentro del mismo tipo e institución
    tasas_list = []

    for (tipo, grado), grupo in grade_stats.groupby(["Tipo_Institucion", "Grado"]):
        enemigos_totales = len(grupo) - 1

        for idx, fila in grupo.iterrows():
            valor = fila["Pct_Bueno_Exc"]
            ganadas  = sum(valor > grupo["Pct_Bueno_Exc"])
            empatadas = sum(valor == grupo["Pct_Bueno_Exc"]) - 1  # Restamos a sí misma

            if enemigos_totales > 0:
                tasa = ((ganadas * 1) + (empatadas * 0.5)) / enemigos_totales
            else:
                tasa = 1.0  # Sin contrincantes: supremacía total

            tasas_list.append((idx, tasa * 100))

    tasas_df = pd.DataFrame(tasas_list, columns=["index", "Win_Rate_Grado"]).set_index("index")
    grade_stats = grade_stats.join(tasas_df)

    # ── FASE 5: MEDIA PONDERADA ───────────────────────────────────────────────
    # Antes: Win_Rate_Final = mean(Win_Rate_Grado)  ← todos los grados pesan igual
    # Ahora: Win_Rate_Final = Σ(WR_g × n_g) / Σ(n_g)  ← grados más grandes pesan más
    #
    # Esto evita que un grado con 5 alumnos distorsione el puntaje final con la
    # misma fuerza que un grado con 80 alumnos.
    # ─────────────────────────────────────────────────────────────────────────
    grade_stats["WR_x_N"] = grade_stats["Win_Rate_Grado"] * grade_stats["Total_Examenes_Grado"]

    school_stats = grade_stats.groupby(["Nro de centro", "Centro", "Tipo_Institucion", "Departamento"]).agg(
        Sum_WR_x_N      =("WR_x_N",              "sum"),
        Sum_N           =("Total_Examenes_Grado", "sum"),
        Grados_Evaluados=("Grado",                "nunique")
    ).reset_index()

    school_stats["Win_Rate_Final"] = school_stats["Sum_WR_x_N"] / school_stats["Sum_N"]

    # Columnas auxiliares de trazabilidad (útiles para auditar el resultado)
    school_stats = school_stats.drop(columns=["Sum_WR_x_N"])
    school_stats = school_stats.rename(columns={
        "Sum_N":            "Total_Examenes_Escuela",
        "Grados_Evaluados": "Grados_Evaluados"
    })

    return school_stats.sort_values("Win_Rate_Final", ascending=False)

# ══════════════════════════════════════════════════════════════════════════════
# 5. EXPORTACIÓN DEL RESUMEN PARA COMPARTIR
# ══════════════════════════════════════════════════════════════════════════════
archivo_compartir = os.path.join(PATH_OUTPUT, "Top_10_Escuelas_Sistema_Victorias.xlsx")
print("\n[*] Calculando Top 10 por categoría (1 vs 0, media ponderada)...")

resultados_completos = calcular_tasa_victorias(df_validos)

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
alt_fill    = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
white_fill  = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
thin_border = Border(
    left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),  bottom=Side(style='thin', color='CCCCCC')
)

for tipo in ["Instituto", "Complejo", "Centro Escolar"]:
    df_tipo = resultados_completos[resultados_completos["Tipo_Institucion"] == tipo]

    top_nac = df_tipo.head(10).copy()
    top_nac.insert(0, 'Ranking', range(1, len(top_nac) + 1))

    top_ss = df_tipo[df_tipo['Departamento'].str.contains('SAN SALVADOR', case=False, na=False)].head(10).copy()
    top_ss.insert(0, 'Ranking', range(1, len(top_ss) + 1))

    for df_top, region in [(top_nac, "Nac"), (top_ss, "SS")]:
        if df_top.empty: continue

        nombre_hoja = f"{region}_{tipo[:4]}"
        ws = wb.create_sheet(nombre_hoja)

        df_top = df_top.rename(columns={
            "Nro de centro":         "Código",
            "Centro":                "Nombre de Institución",
            "Win_Rate_Final":        "Tasa de Victorias vs Pares (%)",
            "Total_Examenes_Escuela":"Total Exámenes",
            "Grados_Evaluados":      "Grados Evaluados"
        })
        df_top["Tasa de Victorias vs Pares (%)"] = (
            df_top["Tasa de Victorias vs Pares (%)"].round(2).astype(str) + "%"
        )
        df_top = df_top.drop(columns=["Tipo_Institucion"])

        # Encabezados
        ws.append(list(df_top.columns))
        for cell in ws[1]:
            cell.fill      = header_fill
            cell.font      = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Filas
        for r_idx, row in enumerate(df_top.values, 2):
            ws.append(list(row))
            for c_idx, cell in enumerate(ws[r_idx], 1):
                cell.fill      = alt_fill if r_idx % 2 == 0 else white_fill
                cell.border    = thin_border
                cell.alignment = Alignment(
                    horizontal="center" if c_idx == 1 or isinstance(cell.value, (int, float)) else "left"
                )

        # Anchos de columna
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 45)
        ws.row_dimensions[1].height = 25

        # Consola
        print(f"\n--- TOP 10 {region.upper()} ({tipo.upper()}) ---")
        for _, row in df_top.iterrows():
            print(
                f" {row['Ranking']:2d}. {str(row['Código']):<8} | "
                f"{str(row['Nombre de Institución'])[:45]:<45} | "
                f"WR: {row['Tasa de Victorias vs Pares (%)']} | "
                f"n={row['Total Exámenes']} | Grados={row['Grados Evaluados']}"
            )

wb.save(archivo_compartir)
print(f"\n[OK] Archivo generado exitosamente en:\n     {archivo_compartir}")