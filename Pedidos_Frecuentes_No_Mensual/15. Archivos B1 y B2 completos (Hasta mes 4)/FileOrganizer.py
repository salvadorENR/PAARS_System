"""
agregar_estatus_y_columnas_CE.py
=================================
Toma dos archivos Excel de B2 y realiza dos operaciones:

1. Hoja Estudiantes (destino = M4_Fund_GradoFund):
   - Rellena "Estatus Centro Escolar" y "Estatus Grado" (ya existen en cols 10-11)
     con los valores y colores de la hoja Estudiantes del archivo fuente,
     haciendo el join por NIE.

2. Hoja Centros Escolares (destino = M4_Fund_GradoFund):
   - Inserta 13 columnas ("Estatus General", "Mat: veces ≤ univ (/1)", ..., "11° (2do Año Bach.)")
     justo después de "Tipo_de_Centro" (col 3), copiando valores y colores
     desde la hoja Centros Escolares del archivo fuente, por Código.

El archivo resultante se guarda en la carpeta de salida indicada.
"""

import os
import shutil
from copy import copy
import openpyxl
from openpyxl.styles import PatternFill

# ── RUTAS ──────────────────────────────────────────────────────────────────────

DEST_FILE = (
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System"
    r"\Pedidos_Frecuentes_No_Mensual\15. Archivos B1 y B2 completos (Hasta mes 4)"
    r"\Scores_Fund_M4\Centros_Escolares_B2_Actualizado_M4_Fund_GradoFund.xlsx"
)

SRC_FILE = (
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System"
    r"\Pedidos_Frecuentes_No_Mensual\15. Archivos B1 y B2 completos (Hasta mes 4)"
    r"\Update_AllScores_M3\Centros_Escolares_B2_Actualizado.xlsx"
)

OUTPUT_DIR = (
    r"C:\Users\USUARIO.DESKTOP-KJ9DUHE.000\Documents\PAARS_System"
    r"\Pedidos_Frecuentes_No_Mensual\15. Archivos B1 y B2 completos (Hasta mes 4)"
    r"\B2_Completo_Con_Estatus"
)

OUTPUT_FILE = os.path.join(
    OUTPUT_DIR,
    "Centros_Escolares_B2_Actualizado_M4_Fund_GradoFund_ConEstatus.xlsx"
)

# ── COLUMNAS A TRANSFERIR ──────────────────────────────────────────────────────

# Hoja Estudiantes: columnas en el archivo FUENTE que queremos copiar
EST_SRC_COLS = ["Estatus Centro Escolar", "Estatus Grado"]

# Hoja Centros Escolares: columnas en el archivo FUENTE que queremos insertar
CE_SRC_COLS = [
    "Estatus General",
    "Mat: veces ≤ univ (/1)",
    "Len: veces ≤ univ (/1)",
    "2° Grado",
    "3° Grado",
    "4° Grado",
    "5° Grado",
    "6° Grado",
    "7° Grado",
    "8° Grado",
    "9° Grado",
    "10° (1er Año Bach.)",
    "11° (2do Año Bach.)",
]


# ── UTILIDADES ─────────────────────────────────────────────────────────────────

def copy_fill(src_cell):
    """Devuelve un PatternFill copiado de la celda fuente."""
    f = src_cell.fill
    if f and f.fill_type and f.fill_type != "none":
        return copy(f)
    return PatternFill()  # sin relleno


def get_col_index(ws, header_name):
    """Devuelve el índice (1-based) de la columna cuyo encabezado coincide."""
    for c in range(1, ws.max_column + 1):
        if ws.cell(1, c).value == header_name:
            return c
    return None


def build_lookup(ws, key_col_idx, value_col_idxs):
    """
    Construye un dict  {key_value: {col_idx: (cell_value, fill_copy)}}
    a partir de una hoja, usando key_col_idx como clave y value_col_idxs
    como columnas de interés.
    """
    lookup = {}
    for r in range(2, ws.max_row + 1):
        key = ws.cell(r, key_col_idx).value
        if key is None:
            continue
        lookup[key] = {}
        for c in value_col_idxs:
            cell = ws.cell(r, c)
            lookup[key][c] = (cell.value, copy_fill(cell))
    return lookup


# ── PASO 0: PREPARAR CARPETA DE SALIDA Y COPIAR DESTINO ───────────────────────

os.makedirs(OUTPUT_DIR, exist_ok=True)
shutil.copy2(DEST_FILE, OUTPUT_FILE)
print(f"Archivo copiado a: {OUTPUT_FILE}")

# ── CARGAR AMBOS LIBROS ────────────────────────────────────────────────────────

wb_dest = openpyxl.load_workbook(OUTPUT_FILE)
wb_src  = openpyxl.load_workbook(SRC_FILE)

# ══════════════════════════════════════════════════════════════════════════════
# PARTE 1 – HOJA ESTUDIANTES
# Rellenar cols "Estatus Centro Escolar" y "Estatus Grado" con datos del fuente
# Join por NIE
# ══════════════════════════════════════════════════════════════════════════════

ws_est_dest = wb_dest["Estudiantes"]
ws_est_src  = wb_src["Estudiantes"]

# Índices en fuente
src_nie_col   = get_col_index(ws_est_src, "NIE")
src_est_cols  = [get_col_index(ws_est_src, h) for h in EST_SRC_COLS]

# Índices en destino
dest_nie_col  = get_col_index(ws_est_dest, "NIE")
dest_est_cols = [get_col_index(ws_est_dest, h) for h in EST_SRC_COLS]

# Verificar que todas las columnas existen
assert src_nie_col,  "NIE no encontrado en fuente Estudiantes"
assert dest_nie_col, "NIE no encontrado en destino Estudiantes"
for h, c in zip(EST_SRC_COLS, src_est_cols):
    assert c, f"Columna '{h}' no encontrada en fuente Estudiantes"
for h, c in zip(EST_SRC_COLS, dest_est_cols):
    assert c, f"Columna '{h}' no encontrada en destino Estudiantes"

# --- Copiar header fill para las columnas de estatus ---
for src_c, dest_c in zip(src_est_cols, dest_est_cols):
    src_hdr  = ws_est_src.cell(1, src_c)
    dest_hdr = ws_est_dest.cell(1, dest_c)
    dest_hdr.fill = copy_fill(src_hdr)

# --- Construir lookup NIE → {src_col: (value, fill)} ---
nie_lookup = build_lookup(ws_est_src, src_nie_col, src_est_cols)

# --- Rellenar destino ---
matched = 0
for r in range(2, ws_est_dest.max_row + 1):
    nie = ws_est_dest.cell(r, dest_nie_col).value
    if nie in nie_lookup:
        matched += 1
        for src_c, dest_c in zip(src_est_cols, dest_est_cols):
            val, fill = nie_lookup[nie][src_c]
            cell = ws_est_dest.cell(r, dest_c)
            cell.value = val
            cell.fill  = fill

print(f"Estudiantes: {matched} NIEs actualizados de {ws_est_dest.max_row - 1} filas de datos.")

# ══════════════════════════════════════════════════════════════════════════════
# PARTE 2 – HOJA CENTROS ESCOLARES
# Insertar 13 columnas después de "Tipo_de_Centro" (col 3)
# Join por Código
# ══════════════════════════════════════════════════════════════════════════════

ws_ce_dest = wb_dest["Centros Escolares"]
ws_ce_src  = wb_src["Centros Escolares"]

# Índices en fuente
src_cod_col  = get_col_index(ws_ce_src, "Código")
src_ce_cols  = [get_col_index(ws_ce_src, h) for h in CE_SRC_COLS]

assert src_cod_col, "Código no encontrado en fuente Centros Escolares"
for h, c in zip(CE_SRC_COLS, src_ce_cols):
    assert c, f"Columna '{h}' no encontrada en fuente Centros Escolares"

# Construir lookup Código → {src_col: (value, fill)}
cod_lookup = build_lookup(ws_ce_src, src_cod_col, src_ce_cols)

# ── Insertar N columnas justo después de col 3 (Tipo_de_Centro) ──────────────
N = len(CE_SRC_COLS)          # 13
INSERT_AFTER = 3              # después de Tipo_de_Centro
INSERT_AT    = INSERT_AFTER + 1  # = 4

ws_ce_dest.insert_cols(INSERT_AT, N)

# ── Escribir encabezados con fill del fuente ──────────────────────────────────
for i, (h, src_c) in enumerate(zip(CE_SRC_COLS, src_ce_cols)):
    dest_c = INSERT_AT + i
    hdr_cell = ws_ce_dest.cell(1, dest_c)
    hdr_cell.value = h
    # Copiar fill desde el header de la fuente
    hdr_cell.fill = copy_fill(ws_ce_src.cell(1, src_c))

# ── Rellenar datos fila a fila ────────────────────────────────────────────────
# Después de insertar, la columna Código sigue siendo la 1
dest_cod_col = 1

matched_ce = 0
for r in range(2, ws_ce_dest.max_row + 1):
    codigo = ws_ce_dest.cell(r, dest_cod_col).value
    if codigo in cod_lookup:
        matched_ce += 1
        for i, src_c in enumerate(src_ce_cols):
            dest_c = INSERT_AT + i
            val, fill = cod_lookup[codigo][src_c]
            cell = ws_ce_dest.cell(r, dest_c)
            cell.value = val
            cell.fill  = fill

print(f"Centros Escolares: {matched_ce} códigos actualizados de {ws_ce_dest.max_row - 1} filas de datos.")

# ── GUARDAR ───────────────────────────────────────────────────────────────────

wb_dest.save(OUTPUT_FILE)
print(f"\n✔ Archivo guardado en:\n  {OUTPUT_FILE}")