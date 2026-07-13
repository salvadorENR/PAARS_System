import os
import glob
import sys
import re
import pandas as pd
import numpy as np
import warnings
import json
from pandas.errors import PerformanceWarning

# Silenciar warnings no críticos de Pandas
warnings.filterwarnings('ignore', category=PerformanceWarning)

# ==========================================
# 1. CONFIGURACIÓN COMÚN
# ==========================================
ROOT_DIR = r"H:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse"
YEAR_DIR = os.path.join(ROOT_DIR, "2026")
PATH_METADATA = os.path.join(ROOT_DIR, "00_Metadata")

# ==========================================
# 2. CLAVE DE RESPUESTAS OFICIAL
# ==========================================
CLAVE_RESPUESTAS = {
    # ------------------ ÍTEMS DE LENGUA (68) ------------------
    'L_item_1': 'a', 'L_item_2': 'e', 'L_item_3': 'i', 'L_item_4': 'o', 'L_item_5': 'u',
    'L_item_6': 'm', 'L_item_7': 'p', 'L_item_8': 's', 'L_item_9': 'l', 'L_item_10': 't',
    'L_item_11': 'n', 'L_item_12': 'd', 'L_item_13': 'r', 'L_item_14': 'b', 'L_item_15': 'c',
    'L_item_16': 'f', 'L_item_17': 'g', 'L_item_18': 'h', 'L_item_19': 'j', 'L_item_20': 'k',
    'L_item_21': 'ñ', 'L_item_22': 'q', 'L_item_23': 'v', 'L_item_24': 'w', 'L_item_25': 'x',
    'L_item_26': 'y', 'L_item_27': 'z',
    'L_item_28': 'a', 'L_item_29': 'e', 'L_item_30': 'i', 'L_item_31': 'o', 'L_item_32': 'u',
    'L_item_33': 'm', 'L_item_34': 'p', 'L_item_35': 's', 'L_item_36': 'l', 'L_item_37': 't',
    'L_item_38': 'n', 'L_item_39': 'd', 'L_item_40': 'rr', 'L_item_41': 'r', 'L_item_42': 'g',
    'L_item_43': 'f', 'L_item_44': 'b', 'L_item_45': 'k',
    'L_item_46': '2)', 'L_item_47': '1)', 'L_item_48': '1) s - o - l', 'L_item_49': '3) pan', 'L_item_50': '2) l',
    'L_item_51': '1) pa', 'L_item_52': '3) sol', 'L_item_53': '1) mufe', 'L_item_54': '2) palo',
    'L_item_55': '3)', 'L_item_56': '1) camino', 'L_item_57': '2)', 'L_item_58': '1)',
    'L_item_59': '1)', 'L_item_60': '2)', 'L_item_61': '3)',
    'L_item_62': '2)', 'L_item_63': '1)', 'L_item_64': '3)',
    'L_item_65': '2) Nube', 'L_item_66': '2) Es un perro', 
    'L_item_67': '1) Está lloviendo', 'L_item_68': '3) Actuó de forma honesta',

    # ------------------ ÍTEMS DE MATEMÁTICA (18) ------------------
    'M_item_1': '3) 9', 'M_item_2': '2) 15', 'M_item_3': '2) 18 es mayor que 15',
    'M_item_4': '2) 25', 'M_item_5': '1) 4 decenas y 8 unidades', 'M_item_6': '1) 34',
    'M_item_7': '1) 67 < 76', 'M_item_8': '2) 25, 40 y 52', 'M_item_9': '2) Decenas',
    'M_item_10': '3) 205', 'M_item_11': '2) 7', 'M_item_12': '1) 4', 'M_item_13': '3) 57',
    'M_item_14': '2) 33', 'M_item_15': '3) 23', 'M_item_16': '2) 19', 'M_item_17': '3) 16',
    'M_item_18': '2) 12'
}

# ==========================================
# 3. MAPEOS PEDAGÓGICOS
# ==========================================
L_SECCIONES = {
    'L_Bloque0_Reconocimiento_del_Alfabeto': [f'L_item_{i}' for i in range(1, 28)],
    'L_Bloque2_Asociacion_Sonido_Letra': [f'L_item_{i}' for i in range(28, 46)],
    'L_Bloque3_Conciencia_Fonologica_y_Principio_Alfabetico': [f'L_item_{i}' for i in range(46, 51)],
    'L_Bloque2b_Decodificacion_y_Fluidez_Basica': [f'L_item_{i}' for i in range(51, 59)],
    'L_Bloque3b_Comprension_Lectora_Literal_e_Inferencial': [f'L_item_{i}' for i in range(59, 62)],
    'L_Bloque4_Comprension_Auditiva': [f'L_item_{i}' for i in range(62, 65)],
    'L_Bloque5_Comprension_Lectora_Escrita': [f'L_item_{i}' for i in range(65, 69)]
}
M_SECCIONES = {
    'M_Bloque1_Conteo_y_Sentido_Numerico': [f'M_item_{i}' for i in range(1, 5)],
    'M_Bloque2_Valor_Posicional_y_Sistema_Decimal': [f'M_item_{i}' for i in range(5, 11)],
    'M_Bloque3_Operaciones_Aditivas_y_Sustractivas': [f'M_item_{i}' for i in range(11, 19)]
}
TODAS_SECCIONES = {**L_SECCIONES, **M_SECCIONES}

L_PROCESS_MAP = {
    'L_item_1':  'Reconocimiento de Vocales', 'L_item_6':  'Consonantes de Alta Frecuencia',
    'L_item_14': 'Consonantes de Confusión Visual y Baja Frecuencia',
    'L_item_28': 'Fonemas Vocálicos', 'L_item_33': 'Fonemas de Alta Frecuencia',
    'L_item_40': 'Fonemas de Menor Frecuencia o Mayor Dificultad',
    'L_item_46': 'Sonido Inicial', 'L_item_47': 'Rima', 'L_item_48': 'Segmentación de Sonidos',
    'L_item_49': 'Fusión de Sonidos', 'L_item_50': 'Identificación de Sonido Final',
    'L_item_51': 'Decodificación de Sílaba Directa (CV)', 'L_item_52': 'Decodificación de Palabra (CVC)',
    'L_item_53': 'Pseudopalabra Simple', 'L_item_54': 'Identificación de Palabra entre Opciones',
    'L_item_55': 'Fluidez: Reconocer Palabras por Sonido', 'L_item_56': 'Fluidez: Lectura Funcional de Palabras',
    'L_item_57': 'Comprensión Literal (Personaje)', 'L_item_58': 'Comprensión Literal (Acción)',
    'L_item_59': 'Comprensión Literal (Personaje)', 'L_item_60': 'Comprensión Literal (Acción)',
    'L_item_61': 'Comprensión Inferencial: Causa Simple', 'L_item_62': 'Comprensión Inferencial: Relación con Experiencia',
    'L_item_63': 'Comprensión Auditiva: Secuencia Básica', 'L_item_64': 'Comprensión Auditiva: Información Explícita',
    'L_item_65': 'Comprensión Lectora Escrita: Nivel Literal', 'L_item_66': 'Comprensión Lectora Escrita: Nivel Literal',
    'L_item_67': 'Comprensión Lectora Escrita: Nivel Inferencial', 'L_item_68': 'Comprensión Lectora Escrita: Nivel Crítico'
}

M_PROCESS_MAP = {
    'M_item_1': 'Conteo visual', 'M_item_2': 'Secuencia numérica', 'M_item_3': 'Comparación numérica', 'M_item_4': 'Patrón numérico',
    'M_item_5': 'Descomposición DU', 'M_item_6': 'Convertir DU a número', 'M_item_7': 'Comparación con signos', 'M_item_8': 'Ordenar números',
    'M_item_9': 'Valor posicional', 'M_item_10': 'Construcción C-D-U', 'M_item_11': 'Suma básica', 'M_item_12': 'Resta básica',
    'M_item_13': 'Suma de dos cifras', 'M_item_14': 'Resta sin préstamo', 'M_item_15': 'Resta con préstamo', 'M_item_16': 'Problema aditivo',
    'M_item_17': 'Problema sustractivo', 'M_item_18': 'Multiplicación inicial'
}
TODO_PROCESS_MAP = {**L_PROCESS_MAP, **M_PROCESS_MAP}


# ==========================================
# 4. FUNCIÓN DE PUNTUACIÓN PORCENTUAL
# ==========================================
def calcular_puntaje_porcentual(suma_correcta_serie, total_items):
    if total_items == 0:
        return pd.Series([0.0] * len(suma_correcta_serie), index=suma_correcta_serie.index)
    return (suma_correcta_serie / total_items * 100).clip(0, 100)


# ==========================================
# 5. GENERADOR DE EXCEL CON DISTRACTORES
# ==========================================
def exportar_distractores(dict_dfs, ruta_salida, clave_respuestas):
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_fill = PatternFill(start_color="033b6d", end_color="033b6d", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    border = Border(left=Side(style='thin', color='CCCCCC'), 
                    right=Side(style='thin', color='CCCCCC'), 
                    top=Side(style='thin', color='CCCCCC'), 
                    bottom=Side(style='thin', color='CCCCCC'))
    
    for sheet_name, df in dict_dfs.items():
        ws = wb.create_sheet(title=sheet_name)
        
        # Escribir encabezados
        for col_num, col_name in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=col_num, value=col_name)
            cell.fill = header_fill
            cell.font = white_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Ancho de columnas: un poco más angosto para Ítem/Asignatura, más ancho para las opciones
            if col_name in ['Asignatura', 'Ítem']:
                ws.column_dimensions[get_column_letter(col_num)].width = 15
            else:
                ws.column_dimensions[get_column_letter(col_num)].width = 30
        
        # Escribir datos
        for row_num, row_data in enumerate(df.to_dict('records'), 2):
            item_name = row_data['Ítem']
            correcta = clave_respuestas.get(item_name, "").lower().strip()
            
            for col_num, col_name in enumerate(df.columns, 1):
                val = row_data.get(col_name)
                cell = ws.cell(row=row_num, column=col_num, value=val)
                cell.border = border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
                # Identificar si es la opción correcta para rellenar de amarillo
                if col_name.startswith('Opción') and pd.notna(val) and isinstance(val, str):
                    if ' (' in val:
                        opt_text = val.rsplit(' (', 1)[0].strip().lower()
                        if opt_text == correcta and correcta != "":
                            cell.fill = yellow_fill
        
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}1"
        
    wb.save(ruta_salida)

# ==========================================
# 6. GENERADOR DE HTML INTERACTIVO 
# ==========================================
LENGUA_JERARQUIA = [
    {
        'bloque_col':  'L_Bloque0_Reconocimiento_del_Alfabeto',
        'titulo':      'Bloque 0 — Reconocimiento del Alfabeto (Grafemas)',
        'grupos': [
            {'label': 'Grupo 1: Vocales', 'items': [f'L_item_{i}' for i in range(1, 6)]},
            {'label': 'Grupo 2: Consonantes de Alta Frecuencia', 'items': [f'L_item_{i}' for i in range(6, 14)]},
            {'label': 'Grupo 3: Consonantes de Confusión Visual y Baja Frecuencia', 'items': [f'L_item_{i}' for i in range(14, 28)]},
        ]
    },
    {
        'bloque_col':  'L_Bloque2_Asociacion_Sonido_Letra',
        'titulo':      'Bloque 2 — Asociación Sonido–Letra (18 Fonemas)',
        'grupos': [
            {'label': 'Subgrupo 1: Vocales (Inicio/Final de palabra)', 'items': [f'L_item_{i}' for i in range(28, 33)]},
            {'label': 'Subgrupo 2: Consonantes de Alta Frecuencia', 'items': [f'L_item_{i}' for i in range(33, 40)]},
            {'label': 'Subgrupo 3: Consonantes de Menor Frecuencia o Mayor Dificultad', 'items': [f'L_item_{i}' for i in range(40, 46)]},
        ]
    },
    {
        'bloque_col':  'L_Bloque3_Conciencia_Fonologica_y_Principio_Alfabetico',
        'titulo':      'Bloque 3 — Conciencia Fonológica y Principio Alfabético',
        'grupos': [{'label': 'Ítems 1–5 (audio-only)', 'items': [f'L_item_{i}' for i in range(46, 51)]}]
    },
    {
        'bloque_col':  'L_Bloque2b_Decodificacion_y_Fluidez_Basica',
        'titulo':      'Bloque 2 — Decodificación y Fluidez Básica',
        'grupos': [
            {'label': 'Decodificación', 'items': [f'L_item_{i}' for i in range(51, 55)]},
            {'label': 'Fluidez', 'items': [f'L_item_{i}' for i in range(55, 57)]},
            {'label': 'Comprensión básica', 'items': [f'L_item_{i}' for i in range(57, 59)]},
        ]
    },
    {
        'bloque_col':  'L_Bloque3b_Comprension_Lectora_Literal_e_Inferencial',
        'titulo':      'Bloque 3 — Comprensión Lectora Literal e Inferencial',
        'grupos': [
            {'label': 'C1 — Recuperar información explícita', 'items': ['L_item_59', 'L_item_60']},
            {'label': 'C2 — Hacer inferencias directas', 'items': ['L_item_61']},
        ]
    },
    {
        'bloque_col':  'L_Bloque4_Comprension_Auditiva',
        'titulo':      'Bloque 4 — Comprensión Auditiva',
        'grupos': [
            {'label': 'C1 — Escuchar textos orales y secuencias', 'items': ['L_item_62', 'L_item_63']},
            {'label': 'C2 — Identificar información explícita', 'items': ['L_item_64']},
        ]
    },
    {
        'bloque_col':  'L_Bloque5_Comprension_Lectora_Escrita',
        'titulo':      'Ítems de Comprensión Lectora — Prueba de Fundamentos',
        'grupos': [
            {'label': 'Nivel LITERAL', 'items': ['L_item_65', 'L_item_66']},
            {'label': 'Nivel INFERENCIAL / CRÍTICO', 'items': ['L_item_67', 'L_item_68']},
        ]
    },
]

MATEMATICA_JERARQUIA = [
    {
        'bloque_col':  'M_Bloque1_Conteo_y_Sentido_Numerico',
        'titulo':      'Bloque 1 — Conteo y Sentido Numérico (4 Ítems)',
        'grupos': [{'label': 'Ítems 1–4', 'items': [f'M_item_{i}' for i in range(1, 5)]}]
    },
    {
        'bloque_col':  'M_Bloque2_Valor_Posicional_y_Sistema_Decimal',
        'titulo':      'Bloque 2 — Valor Posicional y Sistema Decimal (6 Ítems)',
        'grupos': [
            {'label': 'Decenas y Unidades (D-U)', 'items': [f'M_item_{i}' for i in range(5, 8)]},
            {'label': 'Centenas, Decenas y Unidades (C-D-U)', 'items': [f'M_item_{i}' for i in range(8, 11)]},
        ]
    },
    {
        'bloque_col':  'M_Bloque3_Operaciones_Aditivas_y_Sustractivas',
        'titulo':      'Bloque 3 — Operaciones Aditivas y Sustractivas (8 Ítems)',
        'grupos': [
            {'label': 'Operaciones básicas (una cifra)', 'items': ['M_item_11', 'M_item_12']},
            {'label': 'Operaciones de dos cifras', 'items': ['M_item_13', 'M_item_14', 'M_item_15']},
            {'label': 'Problemas verbales', 'items': ['M_item_16', 'M_item_17', 'M_item_18']},
        ]
    },
]

ITEM_LABELS = {
    'L_item_1': 'Letra A', 'L_item_2': 'Letra E', 'L_item_3': 'Letra I', 'L_item_4': 'Letra O', 'L_item_5': 'Letra U',
    'L_item_6': 'Letra M', 'L_item_7': 'Letra P', 'L_item_8': 'Letra S', 'L_item_9': 'Letra L', 'L_item_10': 'Letra T',
    'L_item_11': 'Letra N', 'L_item_12': 'Letra D', 'L_item_13': 'Letra R', 'L_item_14': 'Letra B', 'L_item_15': 'Letra C',
    'L_item_16': 'Letra F', 'L_item_17': 'Letra G', 'L_item_18': 'Letra H', 'L_item_19': 'Letra J', 'L_item_20': 'Letra K',
    'L_item_21': 'Letra Ñ', 'L_item_22': 'Letra Q', 'L_item_23': 'Letra V', 'L_item_24': 'Letra W', 'L_item_25': 'Letra X',
    'L_item_26': 'Letra Y', 'L_item_27': 'Letra Z',
    'L_item_28': 'Fonema /a/', 'L_item_29': 'Fonema /e/', 'L_item_30': 'Fonema /i/', 'L_item_31': 'Fonema /o/', 'L_item_32': 'Fonema /u/',
    'L_item_33': 'Fonema /m/', 'L_item_34': 'Fonema /p/', 'L_item_35': 'Fonema /s/', 'L_item_36': 'Fonema /l/', 'L_item_37': 'Fonema /t/',
    'L_item_38': 'Fonema /n/', 'L_item_39': 'Fonema /d/', 'L_item_40': 'Fonema /rr/', 'L_item_41': 'Fonema /r/',
    'L_item_42': 'Fonema /g/', 'L_item_43': 'Fonema /f/', 'L_item_44': 'Fonema /b/', 'L_item_45': 'Fonema /k/',
    'L_item_46': 'Ítem 1 — Sonido inicial', 'L_item_47': 'Ítem 2 — Rima', 'L_item_48': 'Ítem 3 — Segmentación',
    'L_item_49': 'Ítem 4 — Fusión', 'L_item_50': 'Ítem 5 — Sonido final',
    'L_item_51': 'Ítem 6 — Sílaba CV', 'L_item_52': 'Ítem 7 — Palabra CVC', 'L_item_53': 'Ítem 8 — Pseudopalabra',
    'L_item_54': 'Ítem 9 — Identificación de palabra', 'L_item_55': 'Ítem 10 — Fluidez', 'L_item_56': 'Ítem 11 — Lectura funcional',
    'L_item_57': 'Ítem 12 — Comprensión literal personaje', 'L_item_58': 'Ítem 13 — Comprensión literal acción',
    'L_item_59': 'Ítem 14 — Literal', 'L_item_60': 'Ítem 15 — Literal',
    'L_item_61': 'Ítem 16 — Inferencial', 'L_item_62': 'Ítem 17 — Inferencial',
    'L_item_63': 'Ítem 18 — Secuencia básica', 'L_item_64': 'Ítem 19 — Info. explícita', 
    'L_item_65': 'Ítem 20 — Nivel Literal (Gato)', 'L_item_66': 'Ítem 21 — Nivel Literal (Perro)',
    'L_item_67': 'Ítem 22 — Nivel Inferencial (Clima)', 'L_item_68': 'Ítem 23 — Nivel Crítico (Juan)',
    'M_item_1':  'Ítem 1 — Conteo visual', 'M_item_2':  'Ítem 2 — Secuencia numérica', 'M_item_3':  'Ítem 3 — Comparación',
    'M_item_4':  'Ítem 4 — Patrón numérico', 'M_item_5':  'Ítem 5 — Descomposición DU', 'M_item_6':  'Ítem 6 — Convertir DU a número',
    'M_item_7':  'Ítem 7 — Comparación signos', 'M_item_8':  'Ítem 8 — Ordenar números', 'M_item_9':  'Ítem 9 — Valor posicional',
    'M_item_10': 'Ítem 10 — Construcción C-D-U', 'M_item_11': 'Ítem 11 — Suma básica', 'M_item_12': 'Ítem 12 — Resta básica',
    'M_item_13': 'Ítem 13 — Suma 2 cifras', 'M_item_14': 'Ítem 14 — Resta sin préstamo', 'M_item_15': 'Ítem 15 — Resta con préstamo',
    'M_item_16': 'Ítem 16 — Problema aditivo', 'M_item_17': 'Ítem 17 — Problema sustractivo', 'M_item_18': 'Ítem 18 — Multiplicación inicial',
}

def generar_html_estatico(df_scored, df_mapping, ruta_salida):
    try:
        import plotly.express as px
        import plotly.graph_objects as go
    except ImportError:
        print("    [!] Para generar el HTML necesitas Plotly (pip install plotly).")
        return

    C_DEEP_NAVY  = '#033b6d'
    C_ROYAL_BLUE = '#094b93'
    C_OCEAN_BLUE = '#3077b9'
    C_SKY_BLUE   = '#a7d2f2'
    C_STEEL_GREY = '#666766'

    def _lm(labels, px_per_char=8, minimum=180):
        return max(max((len(str(l)) for l in labels), default=0) * px_per_char, minimum)

    def _bloque_box(bloque_col, btitle, color, include_plotlyjs=False):
        col = f'{bloque_col}_puntaje_0_100'
        if col not in df_scored.columns: return ''
        data = pd.to_numeric(df_scored[col], errors='coerce').dropna()
        fig = go.Figure()
        fig.add_trace(go.Box(
            y=data, boxpoints='all', jitter=0.3, pointpos=-1.8,
            fillcolor=f'rgba({int(color[1:3],16)},{int(color[3:5],16)},{int(color[5:7],16)},0.3)',
            line_color=color, marker=dict(color=color, size=5, opacity=0.6), name='Distribución'
        ))
        fig.update_layout(
            title=f'Distribución del puntaje — {btitle}',
            yaxis=dict(title='Puntaje (0–100)', range=[-5, 105], gridcolor=C_SKY_BLUE, tickfont=dict(color=C_DEEP_NAVY)),
            xaxis=dict(showticklabels=False), plot_bgcolor='white', paper_bgcolor='white', font_color=C_DEEP_NAVY, height=320,
            margin=dict(l=60, r=60, t=55, b=20), showlegend=False
        )
        fig.add_hline(y=data.mean(), line_dash='dot', line_color=C_DEEP_NAVY, opacity=0.6,
                      annotation_text=f'Media: {data.mean():.1f}', annotation_position='top right')
        return fig.to_html(full_html=False, include_plotlyjs=False)

    def _bar_items(items, labels, vals, title, color):
        num_items = len(items)
        height = max(400, num_items * 38 + 120) 
        
        fig = px.bar(x=vals, y=labels, orientation='h', title=title, text=vals, color_discrete_sequence=[color])
        fig.update_traces(texttemplate='%{text:.1f}%', textposition='outside', marker_color=color)
        fig.update_layout(
            xaxis_range=[0, 118], height=height, 
            margin=dict(l=_lm(labels), r=80, t=50, b=30),
            plot_bgcolor='white', paper_bgcolor='white', font_color=C_DEEP_NAVY,
            yaxis=dict(tickfont=dict(size=11, color=C_DEEP_NAVY), autorange='reversed', fixedrange=True, tickmode='linear', dtick=1),
            xaxis=dict(title='% de Aciertos', gridcolor=C_SKY_BLUE, ticksuffix='%')
        )
        return fig.to_html(full_html=False, include_plotlyjs=False)

    ORDEN_GRADOS_MAESTRO = ['Primer Grado', 'Segundo Grado', 'Tercer Grado', 'Cuarto Grado', 'Quinto Grado', 
                            'Sexto Grado', 'Séptimo Grado', 'Octavo Grado', 'Noveno Grado', '1er Año', '2do Año']

    df_validos = df_scored[df_scored['Nombre del centro'] != 'SIN REGISTRO'].copy()
    
    # Tabla Cruzada: LENGUA
    df_pivot_l = df_validos.pivot_table(index='Nombre del centro', columns='Grado', values='L_total_puntaje_0_100', aggfunc='mean')
    cols_l_order = [g for g in ORDEN_GRADOS_MAESTRO if g in df_pivot_l.columns]
    df_pivot_l = df_pivot_l[cols_l_order].reset_index().round(1).fillna('-')
    df_pivot_l.rename(columns={'Nombre del centro': 'Centro Educativo'}, inplace=True)
    tabla_pivot_l_html = df_pivot_l.to_html(classes='display wrap', table_id='tablaPivotL', index=False)

    # Tabla Cruzada: MATEMÁTICA
    df_pivot_m = df_validos.pivot_table(index='Nombre del centro', columns='Grado', values='M_total_puntaje_0_100', aggfunc='mean')
    cols_m_order = [g for g in ORDEN_GRADOS_MAESTRO if g in df_pivot_m.columns]
    df_pivot_m = df_pivot_m[cols_m_order].reset_index().round(1).fillna('-')
    df_pivot_m.rename(columns={'Nombre del centro': 'Centro Educativo'}, inplace=True)
    tabla_pivot_m_html = df_pivot_m.to_html(classes='display wrap', table_id='tablaPivotM', index=False)

    # ==============================================================
    # PREPARACIÓN DE DATOS JSON PARA EL EXPLORADOR DINÁMICO
    # ==============================================================
    all_l_items = [f'L_item_{i}' for i in range(1, 69) if f'L_item_{i}' in df_scored.columns]
    all_m_items = [f'M_item_{i}' for i in range(1, 19) if f'M_item_{i}' in df_scored.columns]
    all_items = all_l_items + all_m_items
    
    df_item_means = df_validos.groupby(['Nombre del centro', 'Grado'])[all_items].mean() * 100
    df_item_means = df_item_means.round(1)
    df_student_counts = df_validos.groupby(['Nombre del centro', 'Grado']).size()
    
    dynamic_data = {}
    
    # -------------------------------------------------------------
    # OPCIÓN "TODAS LAS ESCUELAS" (Promedio General por Grado)
    # -------------------------------------------------------------
    dynamic_data["-- Todas las escuelas --"] = {}
    df_global_means = df_validos.groupby('Grado')[all_items].mean() * 100
    df_global_counts = df_validos.groupby('Grado').size()
    
    for grade, row in df_global_means.iterrows():
        item_data_dict = row.round(1).to_dict()
        item_data_dict['N'] = int(df_global_counts.loc[grade])
        dynamic_data["-- Todas las escuelas --"][grade] = item_data_dict

    # Agregando opciones por escuela individual
    for (school, grade), row in df_item_means.iterrows():
        if school not in dynamic_data:
            dynamic_data[school] = {}
        
        try:
            total_students_n = int(df_student_counts.loc[school, grade])
        except KeyError:
            total_students_n = 0
            
        item_data_dict = row.to_dict()
        item_data_dict['N'] = total_students_n
        dynamic_data[school][grade] = item_data_dict
        
    json_data = json.dumps(dynamic_data)
    json_labels = json.dumps(ITEM_LABELS)
    json_items_l = json.dumps(all_l_items)
    json_items_m = json.dumps(all_m_items)

    df_tabla = df_scored.copy()
    
    rename_map = {
        'Nombre del centro': 'Centro Educativo',
        'L_total_puntaje_0_100': 'L: Total',
        'M_total_puntaje_0_100': 'M: Total',
        'L_Bloque0_Reconocimiento_del_Alfabeto_puntaje_0_100': 'L-B0: Alfabeto',
        'L_Bloque2_Asociacion_Sonido_Letra_puntaje_0_100': 'L-B2: Sonido/Letra',
        'L_Bloque3_Conciencia_Fonologica_y_Principio_Alfabetico_puntaje_0_100': 'L-B3: Conc. Fonológica',
        'L_Bloque2b_Decodificacion_y_Fluidez_Basica_puntaje_0_100': 'L-B2b: Decodificación',
        'L_Bloque3b_Comprension_Lectora_Literal_e_Inferencial_puntaje_0_100': 'L-B3b: Comp. Lectora',
        'L_Bloque4_Comprension_Auditiva_puntaje_0_100': 'L-B4: Comp. Auditiva',
        'L_Bloque5_Comprension_Lectora_Escrita_puntaje_0_100': 'L-B5: Comp. Escrita',
        'M_Bloque1_Conteo_y_Sentido_Numerico_puntaje_0_100': 'M-B1: Conteo',
        'M_Bloque2_Valor_Posicional_y_Sistema_Decimal_puntaje_0_100': 'M-B2: Valor Posicional',
        'M_Bloque3_Operaciones_Aditivas_y_Sustractivas_puntaje_0_100': 'M-B3: Operaciones'
    }

    cols_show = ['NIE', 'Nombre del centro', 'Grado', 'Grupo', 'L_total_puntaje_0_100', 'M_total_puntaje_0_100']
    cols_show += [c for c in df_tabla.columns if 'puntaje_0_100' in c and 'total' not in c]
    df_tabla = df_tabla[[c for c in cols_show if c in df_tabla.columns]]
    for c in df_tabla.select_dtypes(include=['float64']).columns: df_tabla[c] = df_tabla[c].round(1)
    
    df_tabla = df_tabla.rename(columns=rename_map)
    tabla_html = df_tabla.to_html(classes='display wrap', table_id='tablaEstudiantes', index=False)

    # ENSAMBLADO HTML
    html = f"""<!DOCTYPE html><html><head><meta charset='utf-8'><title>Reporte de Fundamentos</title>
<link rel='stylesheet' href='https://cdn.datatables.net/1.13.6/css/jquery.dataTables.css'>
<script src='https://cdn.plot.ly/plotly-2.24.1.min.js'></script>
<style>
  body{{font-family:'Segoe UI',sans-serif;margin:0;padding:20px;background:#f0f4f8;}}
  h1{{color:{C_DEEP_NAVY};text-align:center;border-bottom:3px solid {C_ROYAL_BLUE};padding-bottom:10px;margin-bottom:20px;font-size:22px;}}
  h2{{color:{C_ROYAL_BLUE};margin:40px 0 16px;background:{C_SKY_BLUE}30;border-left:5px solid {C_OCEAN_BLUE};padding:10px 15px;border-radius:4px;font-size:17px;}}
  
  /* ESTILOS PARA LAS PESTAÑAS (TABS) */
  .tab {{ overflow: hidden; background-color: #fff; border-radius: 8px 8px 0 0; border: 1px solid #ccc; display: flex; }}
  .tab button {{ background-color: inherit; float: left; border: none; outline: none; cursor: pointer; padding: 14px 24px; transition: 0.3s; font-size: 15px; font-weight: bold; color: {C_STEEL_GREY}; flex: 1; }}
  .tab button:hover {{ background-color: {C_SKY_BLUE}50; }}
  .tab button.active {{ background-color: {C_ROYAL_BLUE}; color: white; }}
  .tabcontent {{ display: none; padding: 25px; border: 1px solid #ccc; border-top: none; background-color: #fff; border-radius: 0 0 8px 8px; }}

  .container{{background:#fff;padding:25px;border-radius:8px;box-shadow:0 4px 12px rgba(3,59,109,.10);margin-bottom:24px;}}
  .metrics{{display:flex;justify-content:space-around;text-align:center;margin-bottom:10px;}}
  .metric-box{{background:{C_SKY_BLUE}25;padding:20px;border-radius:8px;width:30%;border-top:4px solid {C_ROYAL_BLUE};}}
  .metric-box.l{{border-top-color:{C_OCEAN_BLUE};}} .metric-box.m{{border-top-color:{C_DEEP_NAVY};}}
  .metric-title{{font-size:12px;color:{C_STEEL_GREY};text-transform:uppercase;font-weight:700;}}
  .metric-value{{font-size:30px;font-weight:700;color:{C_DEEP_NAVY};margin-top:8px;}}
  .subject-header{{padding:10px 16px;border-radius:6px;color:#fff;font-size:15px;font-weight:700;margin-bottom:16px;}}
  .sh-l{{background:{C_ROYAL_BLUE};}} .sh-m{{background:{C_DEEP_NAVY};}}
  .bloque-card{{border:1px solid {C_SKY_BLUE};border-radius:8px;margin-bottom:24px;overflow:hidden;}}
  .bloque-title{{padding:10px 16px;font-size:13px;font-weight:700;color:{C_DEEP_NAVY};background:{C_SKY_BLUE}40;border-bottom:1px solid {C_SKY_BLUE};}}
  .bloque-body{{padding:16px;}} .grupo-label{{font-size:12px;font-weight:600;color:{C_STEEL_GREY};text-transform:uppercase;margin:14px 0 6px;}}
  
  /* ESTILOS GLOBALES DE TABLAS */
  table.dataTable {{ border-collapse: collapse; width: 100% !important; }}
  table.dataTable th, table.dataTable td {{ text-align: center; vertical-align: middle; font-size: 12px; white-space: nowrap; }}
  table.dataTable th {{ background:{C_DEEP_NAVY}; color:#fff; font-size: 11px; padding: 8px 10px; }}
  table.dataTable tr:hover {{ background:{C_SKY_BLUE}40!important; }}
  
  #tablaEstudiantes th:nth-child(2), #tablaEstudiantes td:nth-child(2) {{ text-align: left; min-width: 250px; white-space: normal; }}
  
  /* FIX TABLAS CRUZADAS: Aseguramos que la columna Centro Educativo sea amplia y el texto baje si es necesario */
  #tablaPivotL th:first-child, #tablaPivotL td:first-child,
  #tablaPivotM th:first-child, #tablaPivotM td:first-child {{ text-align: left; min-width: 250px; white-space: normal; }}
</style></head><body>

<h1>📊 Reporte de resultados de la aplicación de la Prueba de Fundamentos</h1>

<div class="tab">
  <button class="tablinks active" onclick="openTab(event, 'Global')" id="defaultOpen">Reporte General Consolidado</button>
  <button class="tablinks" onclick="openTab(event, 'Explorador')">Explorador Dinámico</button>
</div>

<div id="Global" class="tabcontent" style="display:block;">
    <div class='metrics'>
        <div class='metric-box'><div class='metric-title'>Total de Estudiantes Evaluados</div><div class='metric-value'>{len(df_scored)}</div></div>
        <div class='metric-box l'><div class='metric-title'>Promedio General — Lengua</div><div class='metric-value'>{df_scored.get('L_total_puntaje_0_100', pd.Series([0])).mean():.1f} / 100</div></div>
        <div class='metric-box m'><div class='metric-title'>Promedio General — Matemática</div><div class='metric-value'>{df_scored.get('M_total_puntaje_0_100', pd.Series([0])).mean():.1f} / 100</div></div>
    </div>

    <h2 id='sec1'>Distribución Global de Puntajes por Asignatura y Grado</h2>
"""

    cols_melt = [c for c in ['L_total_puntaje_0_100', 'M_total_puntaje_0_100'] if c in df_scored.columns]
    if cols_melt:
        df_m = df_scored.melt(value_vars=cols_melt, var_name='Asignatura', value_name='Puntaje')
        df_m['Asignatura'] = df_m['Asignatura'].map({'L_total_puntaje_0_100': 'Lengua', 'M_total_puntaje_0_100': 'Matemática'})
        fig_g = px.box(df_m, x='Asignatura', y='Puntaje', color='Asignatura', points='all', color_discrete_map={'Lengua': C_OCEAN_BLUE, 'Matemática': C_DEEP_NAVY})
        fig_g.update_layout(yaxis_range=[-5, 105], plot_bgcolor='white', paper_bgcolor='white', font_color=C_DEEP_NAVY, height=450)
        html += f"<div>{fig_g.to_html(full_html=False, include_plotlyjs='cdn')}</div>"

    for subj_col, subj_name, subj_color in [('L_total_puntaje_0_100', 'Lengua', C_OCEAN_BLUE), ('M_total_puntaje_0_100', 'Matemática', C_DEEP_NAVY)]:
        if subj_col in df_scored.columns and 'Grado' in df_scored.columns:
            df_slice = df_scored.dropna(subset=[subj_col, 'Grado'])
            unique_grades = df_slice['Grado'].unique()
            order_filtered = [g for g in ORDEN_GRADOS_MAESTRO if g in unique_grades]
            if not df_slice.empty:
                fig_grado = px.box(df_slice, x='Grado', y=subj_col, points='all',
                                   title=f'Distribución de Puntajes Globales por Grado — {subj_name}',
                                   labels={subj_col: 'Puntaje (0–100)'}, color_discrete_sequence=[subj_color],
                                   category_orders={"Grado": order_filtered})
                fig_grado.update_layout(yaxis_range=[-5, 105], plot_bgcolor='white', paper_bgcolor='white', font_color=C_DEEP_NAVY, height=450)
                html += f"<div>{fig_grado.to_html(full_html=False, include_plotlyjs=False)}</div>"

    # SECCIONES POR BLOQUE (LENGUA)
    html += "<h2>Análisis Exhaustivo — Lengua</h2><div class='subject-header sh-l'>Diagnóstico Lengua · Nivel 1</div>"
    for bloque in LENGUA_JERARQUIA:
        bcol = bloque['bloque_col']
        pcol = f"{bcol}_puntaje_0_100"
        if pcol in df_scored.columns:
            html += f"<div class='bloque-card'><div class='bloque-title'>{bloque['titulo']}</div><div class='bloque-body'>"
            html += _bloque_box(bcol, bloque['titulo'], C_OCEAN_BLUE)
            for grupo in bloque['grupos']:
                gitems = [it for it in grupo['items'] if it in df_scored.columns]
                if gitems:
                    g_vals = [df_scored[it].mean() * 100 for it in gitems]
                    g_labs = [ITEM_LABELS.get(it, it) for it in gitems]
                    html += f"<div class='grupo-label'>{grupo['label']}</div>"
                    html += _bar_items(gitems, g_labs, g_vals, f"% Acierto — {grupo['label']}", C_ROYAL_BLUE)
            html += "</div></div>"

    # SECCIONES POR BLOQUE (MATEMÁTICA)
    html += "<h2>Análisis Exhaustivo — Matemática</h2><div class='subject-header sh-m'>Prueba de Fundamentos Matemáticas</div>"
    for bloque in MATEMATICA_JERARQUIA:
        bcol = bloque['bloque_col']
        pcol = f"{bcol}_puntaje_0_100"
        if pcol in df_scored.columns:
            html += f"<div class='bloque-card'><div class='bloque-title'>{bloque['titulo']}</div><div class='bloque-body'>"
            html += _bloque_box(bcol, bloque['titulo'], C_DEEP_NAVY)
            for grupo in bloque['grupos']:
                gitems = [it for it in grupo['items'] if it in df_scored.columns]
                if gitems:
                    g_vals = [df_scored[it].mean() * 100 for it in gitems]
                    g_labs = [ITEM_LABELS.get(it, it) for it in gitems]
                    html += f"<div class='grupo-label'>{grupo['label']}</div>"
                    html += _bar_items(gitems, g_labs, g_vals, f"% Acierto — {grupo['label']}", C_STEEL_GREY)
            html += "</div></div>"

    # TABLAS CRUZADAS
    html += f"""
<h2>Comparativo por Escuela y Grado</h2>
<div style='overflow-x:auto; background:#fff; padding:20px; border-radius:8px; border:1px solid {C_SKY_BLUE};'>
  <h3 style='color:{C_OCEAN_BLUE}; font-size:15px; border-bottom:1px solid {C_SKY_BLUE}; padding-bottom:5px;'>Promedio General — Lengua</h3>
  {tabla_pivot_l_html}
  <br><br>
  <h3 style='color:{C_DEEP_NAVY}; font-size:15px; border-bottom:1px solid {C_SKY_BLUE}; padding-bottom:5px;'>Promedio General — Matemática</h3>
  {tabla_pivot_m_html}
</div>
"""

    # TABLA FINAL DE ESTUDIANTES
    html += f"""
<h2>Base de Datos de Estudiantes</h2>
<div style='overflow-x:auto; background:#fff; padding:20px; border-radius:8px;'>
  <p style='color:{C_STEEL_GREY};font-size:13px;margin-bottom:12px;'>
    <i>Escribe en las cajas debajo de cada columna para filtrar la información.</i>
  </p>
  {tabla_html}
</div>
</div> """

    # ==============================================
    # PESTAÑA 2: EXPLORADOR DINÁMICO
    # ==============================================
    html += f"""
<div id="Explorador" class="tabcontent">
    <h2>Explorador Interactivo por Ítem y Grado</h2>
    <p style='color:{C_STEEL_GREY};font-size:14px;margin-bottom:20px;'>
        <i>Seleccione "Todas las escuelas" para ver el desempeño nacional por grado, o seleccione un Centro Educativo para ver sus resultados locales. Posicione el cursor sobre cualquier barra para ver el porcentaje exacto y el N (estudiantes evaluados).</i>
    </p>
    
    <div style='margin-bottom: 30px;'>
        <label style='font-weight:bold; color:{C_DEEP_NAVY}; font-size:16px;'>Seleccione un filtro:</label><br>
        <select id='schoolSelect' style='width:50%; padding:10px; border:2px solid {C_SKY_BLUE}; border-radius:4px; font-size:14px;'></select>
    </div>
    
    <div id='dynamicPlotL' style='width:100%; height: 2500px; margin-bottom:60px;'></div>
    <div id='dynamicPlotM' style='width:100%; height: 1000px;'></div>
</div> <script src='https://code.jquery.com/jquery-3.7.0.min.js'></script>
<script src='https://cdn.datatables.net/1.13.6/js/jquery.dataTables.js'></script>

<script>
// ==========================================
// LÓGICA DE PESTAÑAS (TABS)
// ==========================================
function openTab(evt, tabName) {{
  var i, tabcontent, tablinks;
  tabcontent = document.getElementsByClassName("tabcontent");
  for (i = 0; i < tabcontent.length; i++) {{
    tabcontent[i].style.display = "none";
  }}
  tablinks = document.getElementsByClassName("tablinks");
  for (i = 0; i < tablinks.length; i++) {{
    tablinks[i].className = tablinks[i].className.replace(" active", "");
  }}
  document.getElementById(tabName).style.display = "block";
  evt.currentTarget.className += " active";
  
  $(window).trigger('resize');
  if(tabName === 'Explorador') updateDynamicPlots();
}}

// ==========================================
// LÓGICA DEL EXPLORADOR DINÁMICO (PLOTLY)
// ==========================================
var db = {json_data};
var itemLabels = {json_labels};
var itemsL = {json_items_l};
var itemsM = {json_items_m};
var gradeOrder = ['Primer Grado', 'Segundo Grado', 'Tercer Grado', 'Cuarto Grado', 'Quinto Grado', 'Sexto Grado', 'Séptimo Grado', 'Octavo Grado', 'Noveno Grado', '1er Año', '2do Año'];

var schoolSelect = document.getElementById('schoolSelect');

// Poblar selector asegurando que "-- Todas las escuelas --" esté primero
var keys = Object.keys(db);
keys.sort();
var indexAll = keys.indexOf("-- Todas las escuelas --");
if(indexAll > -1) {{
    keys.splice(indexAll, 1);
    keys.unshift("-- Todas las escuelas --");
}}

keys.forEach(function(school) {{
    var opt = document.createElement('option');
    opt.value = school;
    opt.innerHTML = school;
    schoolSelect.appendChild(opt);
}});

schoolSelect.addEventListener('change', updateDynamicPlots);

function updateDynamicPlots() {{
    var school = schoolSelect.value;
    if(!school || !db[school]) return;
    
    var availableGrades = Object.keys(db[school]);
    availableGrades.sort(function(a, b) {{ return gradeOrder.indexOf(a) - gradeOrder.indexOf(b); }});
    
    var yL = itemsL.map(function(i) {{ return itemLabels[i] || i; }}).reverse();
    var yM = itemsM.map(function(i) {{ return itemLabels[i] || i; }}).reverse();
    
    var tracesL = [];
    var tracesM = [];
    
    var gradeColors = ['#094b93', '#3077b9', '#a7d2f2', '#666766', '#033b6d'];
    
    availableGrades.forEach(function(grade, index) {{
        var rowData = db[school][grade];
        var totalStudentsN = rowData['N'] || '?';
        
        var xL = itemsL.map(function(i) {{ return rowData[i] || 0; }}).reverse();
        var xM = itemsM.map(function(i) {{ return rowData[i] || 0; }}).reverse();
        
        var color = gradeColors[index % gradeColors.length];

        tracesL.push({{
            x: xL, y: yL, name: grade, type: 'bar', orientation: 'h', 
            marker: {{color: color}},
            hovertemplate: '<b>%{{y}}</b><br>Grado: '+grade+'<br>Porcentaje: %{{x:.1f}}%<br>N (Estudiantes): '+totalStudentsN+'<extra></extra>'
        }});
        
        tracesM.push({{
            x: xM, y: yM, name: grade, type: 'bar', orientation: 'h', 
            marker: {{color: color}},
            hovertemplate: '<b>%{{y}}</b><br>Grado: '+grade+'<br>Porcentaje: %{{x:.1f}}%<br>N (Estudiantes): '+totalStudentsN+'<extra></extra>'
        }});
    }});

    var layoutBase = {{
        plot_bgcolor: 'white', paper_bgcolor: 'white',
        font: {{color: '{C_DEEP_NAVY}'}},
        xaxis: {{title: '% de Aciertos', range: [0, 115], gridcolor: '{C_SKY_BLUE}'}},
        yaxis: {{tickfont: {{size: 11}}, tickmode: 'linear', dtick: 1}}, 
        margin: {{l: 250, r: 40, t: 50, b: 40}},
        barmode: 'group',
        legend: {{orientation: 'h', y: 1.02, x: 0}}
    }};

    var layoutL = Object.assign({{title: 'Rendimiento por Ítem - <b>LENGUA</b>'}}, layoutBase);
    layoutL.height = 2500; 
    
    var layoutM = Object.assign({{title: 'Rendimiento por Ítem - <b>MATEMÁTICA</b>'}}, layoutBase);
    layoutM.height = 1000; 

    Plotly.newPlot('dynamicPlotL', tracesL, layoutL, {{responsive: true}});
    Plotly.newPlot('dynamicPlotM', tracesM, layoutM, {{responsive: true}});
}}

// ==========================================
// LÓGICA DE TABLAS (DATATABLES)
// ==========================================
$(document).ready(function(){{
    $('#tablaPivotL, #tablaPivotM').DataTable({{
        language: {{url: '//cdn.datatables.net/plug-ins/1.13.6/i18n/es-ES.json'}},
        paging: false, 
        info: false, 
        searching: false, 
        scrollX: false, 
        autoWidth: false
    }});

    $('#tablaEstudiantes thead tr').clone(true).appendTo('#tablaEstudiantes thead');
    $('#tablaEstudiantes thead tr:eq(1) th').each(function (i) {{
        $(this).removeClass('sorting sorting_asc sorting_desc');
        $(this).html('<input type="text" placeholder="Filtrar" style="width:100%; font-size:11px; padding:4px;" />');
        $('input', this).on('keyup change', function () {{
            if (table.column(i).search() !== this.value) {{
                table.column(i).search(this.value).draw();
            }}
        }});
    }});

    var table = $('#tablaEstudiantes').DataTable({{
        orderCellsTop: true,
        language: {{url: '//cdn.datatables.net/plug-ins/1.13.6/i18n/es-ES.json'}},
        pageLength: 15, scrollX: true, autoWidth: true,
        initComplete: function() {{
            setTimeout(function() {{ table.columns.adjust().draw(); }}, 50);
        }}
    }});
    
    $(window).on('resize', function () {{
        try {{
            Plotly.Plots.resize('dynamicPlotL');
            Plotly.Plots.resize('dynamicPlotM');
        }} catch(e) {{}}
    }});
}});
</script>
</body></html>"""
    with open(ruta_salida, 'w', encoding='utf-8') as f: f.write(html)


# ==========================================
# 7. NÚCLEO ETL (LÓGICA DE MATRÍCULA Y QC)
# ==========================================
def _ejecutar_etl(current_month_path):
    PATH_FORMULARIOS = os.path.join(current_month_path, "01_Formularios_Crudos")
    PATH_DATASETS    = os.path.join(current_month_path, "02_Datasets_Procesados")
    os.makedirs(PATH_FORMULARIOS, exist_ok=True)
    os.makedirs(PATH_DATASETS, exist_ok=True)

    archivos_form = glob.glob(os.path.join(PATH_FORMULARIOS, "*.csv"))
    if not archivos_form: return None, None, None, None, None, None, None
    archivo_reciente_form = max(archivos_form, key=os.path.getmtime)
    
    with open(archivo_reciente_form, 'r', encoding='utf-8', errors='ignore') as f:
        sep = ';' if ';' in f.readline() else ','

    df_form = pd.read_csv(archivo_reciente_form, dtype=str, encoding_errors='ignore', sep=sep)
    df_form.columns = df_form.columns.str.strip()
    col_email = next((c for c in df_form.columns if 'correo' in c.lower() or 'email' in c.lower()), None)
    df_form = df_form[df_form[col_email].astype(str).str.endswith('@clases.edu.sv', na=False)].copy()
    df_form['NIE'] = df_form[col_email].astype(str).apply(lambda x: x.split('@')[0].strip())

    columnas_items = df_form.columns[3:89]
    df_form = df_form.rename(columns={col: (f'L_item_{i+1}' if i < 68 else f'M_item_{i - 68 + 1}') for i, col in enumerate(columnas_items)})

    archivos_mat = glob.glob(os.path.join(PATH_METADATA, "MatriculaProgresoMes*.csv"))
    def ext_num(r): m = re.search(r'\d+', os.path.basename(r)); return int(m.group()) if m else 0
    archivo_mat_reciente = max(archivos_mat, key=ext_num)
    df_mat = pd.read_csv(archivo_mat_reciente, sep=';', dtype=str, encoding_errors='ignore')
    
    col_nie = next(c for c in df_mat.columns if 'nie' in c.lower())
    col_cod = next((c for c in df_mat.columns if 'codigo' in c.lower() or 'código' in c.lower()), None)
    col_nom = next(c for c in df_mat.columns if 'nombre' in c.lower() and 'secc' not in c.lower())
    col_grado = next(c for c in df_mat.columns if 'grado' in c.lower())
    col_grupo = next(c for c in df_mat.columns if 'nombre_secc' in c.lower())

    df_meta = df_mat[[col_nie, col_cod, col_nom, col_grado, col_grupo]].copy()
    df_meta.columns = ['NIE', 'Código de infraestructura', 'Nombre del centro', 'Grado', 'Grupo']
    df_meta['NIE'] = df_meta['NIE'].astype(str).str.replace('.0', '', regex=False).str.strip()

    df_cruzado = pd.merge(df_meta, df_form, on='NIE', how='right')
    df_cruzado.fillna({'Código de infraestructura':'SIN REGISTRO', 'Nombre del centro':'SIN REGISTRO', 'Grado':'SIN REGISTRO', 'Grupo':'SIN REGISTRO'}, inplace=True)
    
    col_puntuacion = next((c for c in df_cruzado.columns if 'puntuaci' in c.lower() or 'score' in c.lower()), None)
    if col_puntuacion:
        df_cruzado['Puntaje_Global_Formulario'] = df_cruzado[col_puntuacion].astype(str).str.split('/').str[0].str.strip()
    else:
        df_cruzado['Puntaje_Global_Formulario'] = "0"

    todas_preguntas = ([f'L_item_{i}' for i in range(1, 69)] + [f'M_item_{i}' for i in range(1, 19)])
    columnas_orden = (['Código de infraestructura', 'Nombre del centro', 'Grado', 'Grupo', 'NIE', col_email, 'Puntaje_Global_Formulario'] + todas_preguntas)
    columnas_orden = [c for c in columnas_orden if c in df_cruzado.columns]
    
    df_raw    = df_cruzado[columnas_orden].copy()
    df_scored = df_cruzado.copy()
    
    # -------------------------------------------------------------
    # ANÁLISIS DE DISTRACTORES USANDO df_raw (EN MEMORIA)
    # -------------------------------------------------------------
    distractores_dict = {}
    val_df_raw = df_raw[df_raw['Nombre del centro'] != 'SIN REGISTRO'].copy()
    
    def generar_df_distractores(df_subset):
        rows = []
        for item in todas_preguntas:
            if item not in df_subset.columns: continue
            asig = 'Lengua' if item.startswith('L_') else 'Matemática'
            
            s = df_subset[item].fillna('SIN RESPONDER').astype(str).str.strip()
            if len(s) == 0: continue
            
            vc = s.value_counts(normalize=True) * 100
            
            row_dict = {'Asignatura': asig, 'Ítem': item}
            # Organizar en columnas "Opción 1", "Opción 2", etc.
            for i, (opt, pct) in enumerate(vc.items(), 1):
                row_dict[f'Opción {i}'] = f"{opt} ({pct:.1f}%)"
            rows.append(row_dict)
        return pd.DataFrame(rows)

    distractores_dict['Todas las escuelas'] = generar_df_distractores(val_df_raw)
    for grado, sheet_name in [('Segundo Grado', '2do'), ('Tercer Grado', '3er'), ('Cuarto Grado', '4to')]:
        df_g = val_df_raw[val_df_raw['Grado'] == grado]
        if not df_g.empty:
            distractores_dict[sheet_name] = generar_df_distractores(df_g)

    # -------------------------------------------------------------
    # CALIFICAR df_scored (0 y 1)
    # -------------------------------------------------------------
    for item, correcta in CLAVE_RESPUESTAS.items():
        if item in df_scored.columns:
            df_scored[item] = np.where(df_scored[item].astype(str).str.strip().str.lower() == correcta.lower(), 1, 0)

    for seccion, items in TODAS_SECCIONES.items():
        df_scored[f'{seccion}_puntaje_0_100'] = calcular_puntaje_porcentual(df_scored[[it for it in items if it in df_scored.columns]].sum(axis=1), len(items))

    items_L = [c for c in df_scored.columns if c.startswith('L_item_')]
    items_M = [c for c in df_scored.columns if c.startswith('M_item_')]

    df_scored['L_total_suma_correcta'] = df_scored[items_L].sum(axis=1) if items_L else 0
    df_scored['M_total_suma_correcta'] = df_scored[items_M].sum(axis=1) if items_M else 0
    df_scored['L_total_puntaje_0_100'] = calcular_puntaje_porcentual(df_scored['L_total_suma_correcta'], len(items_L) if items_L else 1)
    df_scored['M_total_puntaje_0_100'] = calcular_puntaje_porcentual(df_scored['M_total_suma_correcta'], len(items_M) if items_M else 1)

    df_qc = df_scored[['NIE', 'Nombre del centro', 'Grado', 'Grupo', 'Puntaje_Global_Formulario']].copy()
    df_qc['Total_Aciertos_Calculados'] = df_scored['L_total_suma_correcta'] + df_scored['M_total_suma_correcta']
    df_qc['Total_Items_Prueba'] = len(items_L) + len(items_M)
    df_qc['Puntaje_Formulario_Num'] = pd.to_numeric(df_qc['Puntaje_Global_Formulario'], errors='coerce').fillna(0)
    df_qc['¿Hay Discrepancia?'] = np.where(df_qc['Puntaje_Formulario_Num'] != df_qc['Total_Aciertos_Calculados'], 'SÍ', 'NO')
    df_qc = df_qc.drop(columns=['Puntaje_Formulario_Num'])

    df_validos = df_scored[df_scored['Nombre del centro'] != 'SIN REGISTRO'].copy()
    evaluados = df_validos.groupby('Nombre del centro').size().reset_index(name='Estudiantes Evaluados')
    cols_puntajes = [c for c in df_scored.columns if 'puntaje_0_100' in c]
    promedios = df_validos.groupby('Nombre del centro')[cols_puntajes].mean().reset_index()
    grados_evaluados = df_validos['Grado'].dropna().unique()
    df_meta_unique = df_meta.drop_duplicates(subset=['NIE'])
    df_meta_filtrado = df_meta_unique[df_meta_unique['Grado'].isin(grados_evaluados)]
    esperados = df_meta_filtrado.groupby('Nombre del centro').size().reset_index(name='Estudiantes Esperados')
    
    df_escuelas = pd.merge(evaluados, promedios, on='Nombre del centro', how='left')
    df_escuelas = pd.merge(df_escuelas, esperados, on='Nombre del centro', how='left')
    codigos_escuelas = df_validos[['Nombre del centro', 'Código de infraestructura']].drop_duplicates(subset=['Nombre del centro'])
    df_escuelas = pd.merge(df_escuelas, codigos_escuelas, on='Nombre del centro', how='left')
    df_escuelas['Estudiantes Esperados'] = df_escuelas['Estudiantes Esperados'].fillna(df_escuelas['Estudiantes Evaluados'])
    df_escuelas['Cobertura (%)'] = np.where(df_escuelas['Estudiantes Esperados'] > 0, (df_escuelas['Estudiantes Evaluados'] / df_escuelas['Estudiantes Esperados'] * 100), 0).round(1)
    df_escuelas['Cobertura (%)'] = df_escuelas['Cobertura (%)'].clip(upper=100.0)

    for col in cols_puntajes: df_escuelas[col] = df_escuelas[col].round(1)
    cols_order = ['Código de infraestructura', 'Nombre del centro', 'Estudiantes Esperados', 'Estudiantes Evaluados', 'Cobertura (%)'] + cols_puntajes
    df_escuelas = df_escuelas[cols_order]

    rename_map = {
        'Nombre del centro': 'Centro Educativo',
        'L_total_puntaje_0_100': 'L: Total',
        'M_total_puntaje_0_100': 'M: Total',
        'L_Bloque0_Reconocimiento_del_Alfabeto_puntaje_0_100': 'L-B0: Alfabeto',
        'L_Bloque2_Asociacion_Sonido_Letra_puntaje_0_100': 'L-B2: Sonido/Letra',
        'L_Bloque3_Conciencia_Fonologica_y_Principio_Alfabetico_puntaje_0_100': 'L-B3: Conc. Fonológica',
        'L_Bloque2b_Decodificacion_y_Fluidez_Basica_puntaje_0_100': 'L-B2b: Decodificación',
        'L_Bloque3b_Comprension_Lectora_Literal_e_Inferencial_puntaje_0_100': 'L-B3b: Comp. Lectora',
        'L_Bloque4_Comprension_Auditiva_puntaje_0_100': 'L-B4: Comp. Auditiva',
        'L_Bloque5_Comprension_Lectora_Escrita_puntaje_0_100': 'L-B5: Comp. Escrita',
        'M_Bloque1_Conteo_y_Sentido_Numerico_puntaje_0_100': 'M-B1: Conteo',
        'M_Bloque2_Valor_Posicional_y_Sistema_Decimal_puntaje_0_100': 'M-B2: Valor Posicional',
        'M_Bloque3_Operaciones_Aditivas_y_Sustractivas_puntaje_0_100': 'M-B3: Operaciones'
    }

    df_estudiantes = df_scored.copy()
    cols_show = ['NIE', 'Nombre del centro', 'Grado', 'Grupo', 'L_total_puntaje_0_100', 'M_total_puntaje_0_100']
    cols_show += [c for c in df_estudiantes.columns if 'puntaje_0_100' in c and 'total' not in c]
    df_estudiantes = df_estudiantes[[c for c in cols_show if c in df_estudiantes.columns]]
    for c in df_estudiantes.select_dtypes(include=['float64']).columns: 
        df_estudiantes[c] = df_estudiantes[c].round(1)
    df_estudiantes = df_estudiantes.rename(columns=rename_map)

    df_mapping = pd.DataFrame([{'Item': k, 'Proceso': v} for k, v in TODO_PROCESS_MAP.items()])
    paths = {
        'raw_df': df_raw, 'scored': df_scored, 'mapping': df_mapping, 'escuelas': df_escuelas, 'qc': df_qc,
        'html': os.path.join(PATH_DATASETS, "4_Reporte_Grafico_Fundamentos.html"),
        'xl_escuelas': os.path.join(PATH_DATASETS, "5_Resultados_Escuelas_Fundamentos.xlsx"),
        'xl_raw': os.path.join(PATH_DATASETS, "1_Resultados_Fundamento_Respuestas_Crudas.xlsx"),
        'xl_scored': os.path.join(PATH_DATASETS, "2_Resultados_Fundamento_Dicotomico_0_1.xlsx"),
        'xl_mapping': os.path.join(PATH_DATASETS, "3_Item_Process_Mapping.xlsx"),
        'xl_qc': os.path.join(PATH_DATASETS, "6_QC_Auditoria_Puntajes.xlsx"),
        'xl_estudiantes': os.path.join(PATH_DATASETS, "7_Base_Datos_Estudiantes.xlsx"),
        'xl_distractores': os.path.join(PATH_DATASETS, "8_Analisis_Distractores_Por_Item.xlsx")
    }
    return df_scored, df_mapping, paths, df_escuelas, df_qc, df_estudiantes, distractores_dict

def procesar_y_generar_excel(current_month_path):
    print("\n=======================================================")
    print(f"  PROCESANDO: {os.path.basename(current_month_path)}")
    print("=======================================================\n")
    df_scored, df_mapping, paths, df_escuelas, df_qc, df_estudiantes, distractores_dict = _ejecutar_etl(current_month_path)
    if df_scored is None: return False
    
    paths['raw_df'].to_excel(paths['xl_raw'], index=False)
    df_scored.to_excel(paths['xl_scored'], index=False)
    df_mapping.to_excel(paths['xl_mapping'], index=False)
    df_escuelas.to_excel(paths['xl_escuelas'], index=False)
    df_qc.to_excel(paths['xl_qc'], index=False)
    df_estudiantes.to_excel(paths['xl_estudiantes'], index=False)
    
    # GUARDAR EL NUEVO ANÁLISIS DE DISTRACTORES
    exportar_distractores(distractores_dict, paths['xl_distractores'], CLAVE_RESPUESTAS)
    
    print("[*] Generando Dashboard HTML Interactivo...")
    generar_html_estatico(df_scored, df_mapping, paths['html'])
    print(f"[OK] Todos los reportes generados exitosamente en:\n     {os.path.dirname(paths['html'])}")
    print("  - 8_Analisis_Distractores_Por_Item.xlsx (NUEVO)")
    return True