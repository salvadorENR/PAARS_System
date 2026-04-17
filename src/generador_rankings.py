# src/generador_rankings.py
import os
import re
import glob
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

# ==========================================
# CONFIGURACIÓN DE RUTAS Y CONSTANTES
# ==========================================
ROOT_DIR = r"G:\Mi unidad\Modernización_Educativa\Gerencia de Evaluación_Proyectos_Análisis\PAARS_Warehouse\2026"

CLASIFICACIONES_INFO = [
    {"clasif": "Crítica", "rango": "(0 a 35)"},
    {"clasif": "Baja", "rango": "(>35 a 45)"},
    {"clasif": "Media", "rango": "(>45 a 55)"},
    {"clasif": "Buena", "rango": "(>55 a 65)"},
    {"clasif": "Excelente", "rango": "(>65 a 100)"},
]

FILLS = {
    "Crítica": PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid"),
    "Baja": PatternFill(start_color="FBEEE6", end_color="FBEEE6", fill_type="solid"),
    "Media": PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid"),
    "Buena": PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid"),
    "Excelente": PatternFill(start_color="D4EFDF", end_color="D4EFDF", fill_type="solid")
}

def clasificar_puntaje(media):
    if pd.isna(media) or str(media).strip() == "": return ""
    try:
        media = float(media)
        if 0 <= media <= 35: return "Crítica"
        elif 35 < media <= 45: return "Baja"
        elif 45 < media <= 55: return "Media"
        elif 55 < media <= 65: return "Buena"
        elif 65 < media <= 100: return "Excelente"
        else: return "Fuera de rango"
    except:
        return ""

def obtener_medias_escuela(progreso_folder):
    """Lee las medias directamente de los CSV Interim (¡Más seguro y rápido!)"""
    interim_dir = os.path.join(progreso_folder, "Interim_CSVs")
    csv_files = glob.glob(os.path.join(interim_dir, "*.csv"))
    
    if not csv_files: 
        return pd.DataFrame()

    resultados_dfs = []
    for f in csv_files:
        try:
            df_res = pd.read_csv(f, dtype=str, encoding_errors='ignore')
            
            # Detectar las columnas de manera flexible por si cambia una mayúscula
            col_cod = next((c for c in df_res.columns if 'nro de centro' in str(c).lower() or 'infra' in str(c).lower() or 'código' in str(c).lower()), None)
            
            # Buscar la columna 'Centro', asegurando que no sea la misma que 'col_cod'
            col_centro = None
            posibles_centro = [c for c in df_res.columns if 'centro' in str(c).lower() or 'institu' in str(c).lower() or 'escuela' in str(c).lower()]
            for p in posibles_centro:
                if p != col_cod:
                    col_centro = p
                    break

            col_theta = next((c for c in df_res.columns if 'escala 0-100' in str(c).lower()), 'theta.global (escala 0-100)')

            # Si el archivo tiene la escuela y las notas, lo guardamos
            if col_cod and col_centro and col_theta in df_res.columns:
                temp_df = df_res[[col_cod, col_centro, col_theta]].copy()
                temp_df.rename(columns={col_cod: 'Nro de centro', col_centro: 'Centro', col_theta: 'Media'}, inplace=True)
                resultados_dfs.append(temp_df)

        except Exception as e:
            pass

    if not resultados_dfs: 
        return pd.DataFrame()
    
    res_df = pd.concat(resultados_dfs, ignore_index=True)
    res_df['Nro de centro'] = res_df['Nro de centro'].astype(str).str.strip().str.replace('.0', '', regex=False)
    res_df['Centro'] = res_df['Centro'].astype(str).str.strip().str.upper()
    res_df['Media'] = pd.to_numeric(res_df['Media'].astype(str).str.replace(',', '.'), errors='coerce')
    res_df = res_df.dropna(subset=['Media', 'Nro de centro'])

    # Agrupamos por escuela y calculamos el promedio global del mes
    grouped = res_df.groupby(['Nro de centro', 'Centro'])['Media'].mean().reset_index()
    grouped['Media'] = grouped['Media'].round(1)

    return grouped

def aplicar_estilos_excel(filepath, df_ranking, df_summary_all, df_summary_sub):
    wb = load_workbook(filepath)
    
    ws1 = wb["Ranking General"]
    ws1.column_dimensions['A'].width = 15
    ws1.column_dimensions['B'].width = 50
    ws1.column_dimensions['C'].width = 12
    ws1.column_dimensions['D'].width = 18

    for row in ws1.iter_rows(min_row=1, max_row=len(df_ranking)+1, min_col=1, max_col=4):
        if row[0].row == 1:
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            continue
        
        clasificacion = row[3].value
        if clasificacion in FILLS:
            for cell in row: cell.fill = FILLS[clasificacion]
        row[2].alignment = Alignment(horizontal="center")
        row[3].alignment = Alignment(horizontal="center")

    ws2 = wb["Resumen General"]
    ws2.column_dimensions['A'].width = 30 
    ws2.column_dimensions['B'].width = 25
    ws2.column_dimensions['C'].width = 25

    for row in ws2.iter_rows(min_row=1, max_row=len(df_summary_all)+1, min_col=1, max_col=3):
        if row[0].row == 1:
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            continue
        
        clasificacion_str = str(row[0].value)
        for base_class in FILLS.keys():
            if base_class in clasificacion_str:
                for cell in row: cell.fill = FILLS[base_class]
                break
                
        row[1].alignment = Alignment(horizontal="center")
        row[2].alignment = Alignment(horizontal="center")

    start_row_sub = len(df_summary_all) + 3
    min_row_sub = start_row_sub + 1
    max_row_sub = min_row_sub + len(df_summary_sub)
    
    for row in ws2.iter_rows(min_row=start_row_sub, max_row=max_row_sub, min_col=1, max_col=3):
        if row[0].row == start_row_sub:
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            continue
        
        for cell in row: cell.fill = FILLS["Media"]
        row[1].alignment = Alignment(horizontal="center")
        row[2].alignment = Alignment(horizontal="center")

    wb.save(filepath)

def aplicar_estilos_historico(filepath, max_cols, max_rows):
    wb = load_workbook(filepath)
    ws = wb["Evolución Histórica"]
    
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 50
    
    from openpyxl.utils import get_column_letter
    for col_idx in range(3, max_cols + 1):
        col_name = str(ws.cell(row=1, column=col_idx).value)
        if "Media" in col_name:
            ws.column_dimensions[get_column_letter(col_idx)].width = 10
        else:
            ws.column_dimensions[get_column_letter(col_idx)].width = 15

    for row in ws.iter_rows(min_row=1, max_row=max_rows, min_col=1, max_col=max_cols):
        if row[0].row == 1:
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            continue
        
        for cell in row[2:]:
            cell.alignment = Alignment(horizontal="center")
            # LA MAGIA DEL COLOR: Pinta la celda de nivel según la palabra clave
            val = str(cell.value)
            if val in FILLS:
                cell.fill = FILLS[val]

    wb.save(filepath)

def generador_rankings_masivo(df_master=None):
    # Ya no necesita el df_master porque extrae todo directo del CSV!
    print("======================================================")
    print("   GENERADOR DE RANKINGS Y EVOLUCIÓN HISTÓRICA")
    print("======================================================\n")

    # 1. ENCONTRAR TODOS LOS INTERIM_CSVS
    patron_progreso = os.path.join(ROOT_DIR, "*PROGRESO*", "Interim_CSVs")
    carpetas_interim = glob.glob(patron_progreso)
    
    examenes = {} 
    for c in carpetas_interim:
        folder_progreso = os.path.dirname(c)
        match = re.search(r'^(\d+)_PROGRESO', os.path.basename(folder_progreso), re.IGNORECASE)
        if match:
            numero = int(match.group(1))
            examenes[numero] = folder_progreso

    if not examenes:
        print("[!] No se encontraron carpetas Interim_CSVs de Pruebas de Progreso.")
        return

    historial_medias = {}

    # 2. PROCESAR ORDENADAMENTE DEL MES 1 AL N
    for numero_examen in sorted(examenes.keys()):
        progreso_folder = examenes[numero_examen]
        nombre_mes = os.path.basename(progreso_folder).split('_')[-1]
        print(f"[*] Procesando Prueba de Progreso {numero_examen} ({nombre_mes})...")

        df_medias = obtener_medias_escuela(progreso_folder)
        
        if not df_medias.empty:
            historial_medias[numero_examen] = df_medias.copy()
        else:
            print(f"  [!] No hay datos en Interim_CSVs para Progreso {numero_examen}.")

        out_dir = os.path.join(progreso_folder, "Final_Reports", "Ranking_Escuelas")
        os.makedirs(out_dir, exist_ok=True)

        # ==========================================
        # ARCHIVO 1: RANKING GENERAL ACTUAL
        # ==========================================
        if not df_medias.empty:
            df_ranking = df_medias.copy()
            df_ranking['Clasificación'] = df_ranking['Media'].apply(clasificar_puntaje)
            df_ranking = df_ranking.sort_values(by='Media', ascending=True)

            counts_all = df_ranking['Clasificación'].value_counts()
            total_all = len(df_ranking)
            summary_all_data = []
            for item in CLASIFICACIONES_INFO:
                c_name = item["clasif"]
                cnt = counts_all.get(c_name, 0)
                pct = (cnt / total_all) * 100 if total_all > 0 else 0
                summary_all_data.append({
                    'Clasificación': f"{c_name} {item['rango']}", 
                    'Cantidad de Escuelas': cnt, 
                    'Porcentaje del Total (%)': round(pct, 1)
                })
            df_summary_all = pd.DataFrame(summary_all_data)

            rango_45_50 = len(df_ranking[(df_ranking['Media'] > 45) & (df_ranking['Media'] <= 50)])
            rango_50_55 = len(df_ranking[(df_ranking['Media'] > 50) & (df_ranking['Media'] <= 55)])
            df_summary_sub = pd.DataFrame([
                {'Rango de Puntaje': '(>45 a 50)', 'Cantidad de Escuelas': rango_45_50, 'Porcentaje del Total (%)': round((rango_45_50 / total_all) * 100, 1) if total_all > 0 else 0},
                {'Rango de Puntaje': '(>50 a 55)', 'Cantidad de Escuelas': rango_50_55, 'Porcentaje del Total (%)': round((rango_50_55 / total_all) * 100, 1) if total_all > 0 else 0}
            ])

            file1_path = os.path.join(out_dir, f"1_Ranking_General_Progreso_{numero_examen}.xlsx")
            try:
                with pd.ExcelWriter(file1_path, engine='openpyxl') as writer:
                    df_ranking.to_excel(writer, sheet_name="Ranking General", index=False)
                    df_summary_all.to_excel(writer, sheet_name="Resumen General", index=False)
                    df_summary_sub.to_excel(writer, sheet_name="Resumen General", startrow=len(df_summary_all)+3, index=False)
                
                aplicar_estilos_excel(file1_path, df_ranking, df_summary_all, df_summary_sub)
                print(f"  [OK] Creado: 1_Ranking_General_Progreso_{numero_examen}.xlsx")
            except PermissionError:
                pass

        # ==========================================
        # ARCHIVO 2: EVOLUCIÓN HISTÓRICA
        # ==========================================
        df_historico = None
        nombres_centros = {}
        
        for i in range(1, numero_examen + 1):
            if i in historial_medias:
                df_temp = historial_medias[i][['Nro de centro', 'Centro', 'Media']].copy()
                
                for _, row in df_temp.iterrows():
                    nombres_centros[row['Nro de centro']] = row['Centro']
                
                col_media_nueva = f'Media {i}'
                col_nivel_nueva = f'Nivel {i}'
                
                df_temp[col_nivel_nueva] = df_temp['Media'].apply(clasificar_puntaje)
                df_temp = df_temp[['Nro de centro', 'Media', col_nivel_nueva]].rename(columns={'Media': col_media_nueva})
                
                if df_historico is None:
                    df_historico = df_temp
                else:
                    # Cruzar exclusivamente por el número de centro
                    df_historico = pd.merge(df_historico, df_temp, on=['Nro de centro'], how='outer')

        if df_historico is not None:
            cols_ordenadas = ['Nro de centro', 'Centro']
            
            for i in range(1, numero_examen + 1):
                col_m = f'Media {i}'
                col_n = f'Nivel {i}'
                if col_m not in df_historico.columns:
                    df_historico[col_m] = ""
                if col_n not in df_historico.columns:
                    df_historico[col_n] = ""
                cols_ordenadas.extend([col_m, col_n])
            
            df_historico.insert(1, 'Centro', df_historico['Nro de centro'].map(nombres_centros))
            df_historico = df_historico[cols_ordenadas]
            df_historico = df_historico.sort_values(by='Nro de centro')
            df_historico = df_historico.fillna("")

            file2_path = os.path.join(out_dir, f"2_Evolucion_Medias_Progreso_{numero_examen}.xlsx")
            try:
                df_historico.to_excel(file2_path, sheet_name="Evolución Histórica", index=False)
                aplicar_estilos_historico(file2_path, max_cols=len(df_historico.columns), max_rows=len(df_historico)+1)
                print(f"  [OK] Creado: 2_Evolucion_Medias_Progreso_{numero_examen}.xlsx")
            except PermissionError:
                pass

    print("\n¡Generación de Rankings y Evolución Histórica finalizada con éxito!")